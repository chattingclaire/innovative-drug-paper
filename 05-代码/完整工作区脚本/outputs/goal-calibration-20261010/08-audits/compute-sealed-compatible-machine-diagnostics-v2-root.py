"""Diagnostics bound to a completed explicitly attributed rule composite and one sealed whole-batch LLM snapshot.

Reference77 is disclosed development. Two machines are not gold standards for each other.
"""
import argparse, collections, csv, datetime, hashlib, json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DIMS = ('S', 'E', 'Ec', 'I', 'F')
OLD = ROOT/'06-coding/rule-full4481-final-v5-receipt-recovery-root/recovered-predictions.jsonl'
EXPECTED_SHA = {'A':'2ad1cad77f599bd25dbcb3002e09994ad895ecc85b2c6f3019600d96f10c5cb4',
                'B':'6329f4d94a1c667464811326d3a7916be6409a2861da676d69b916aee8a20f4f'}

def load(p): return json.loads(Path(p).read_bytes())
def bind(p):
    p=Path(p).resolve();b=p.read_bytes()
    return {'path':str(p),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()}
def exact(b):
    current=bind(b['path']);assert all(current[k]==b[k] for k in ('path','bytes','sha256'))
    return current
def save(p,x):
    with Path(p).open('x',encoding='utf8') as f:
        json.dump(x,f,ensure_ascii=False,indent=2,allow_nan=False);f.write('\n')
def jsonlines(p):
    result={}
    with Path(p).open() as f:
        for line in f:
            x=json.loads(line);assert x['id'] not in result
            assert set(x['labels'])==set(DIMS)
            assert all(v is None or type(v) is int and v in (0,1) for v in x['labels'].values())
            result[x['id']]=x
    return result
def verify_composite_seal(path, seen=None):
    seen=set() if seen is None else seen
    path=Path(path).resolve()
    if path in seen:return
    seen.add(path);seal=load(path)
    for key in ('files','output_files','implementation_files','material_bindings','control_inputs'):
        for item in seal.get(key,[]):exact(item)
    for ancestor in seal.get('ancestor_seals',[]):
        exact(ancestor);verify_composite_seal(ancestor['path'],seen)

def verify_prediction_origins(directory):
    lineage=[json.loads(line) for line in (directory/'execution-origin-lineage.jsonl').read_bytes().splitlines()]
    records=(directory/'predictions.jsonl').read_bytes().splitlines(keepends=True)
    assert len(lineage)==len(records)==4481
    ids=[]
    for ordinal,(origin,raw) in enumerate(zip(lineage,records),1):
        record=json.loads(raw);ids.append(record['id'])
        assert origin['corpus_ordinal_1based']==ordinal and origin['id']==record['id']
        assert origin['prediction_line_bytes']==len(raw)
        assert origin['actual_prediction_line_sha256']==hashlib.sha256(raw).hexdigest()
        assert not record['implementation_errors'] and set(record['decisions'])==set(DIMS)
        assert all(record['labels'][d]==record['decisions'][d]['value'] for d in DIMS)
    assert len(set(ids))==4481

def ratio(a,b):return None if b==0 else a/b
def metrics(pairs):
    a=sum(x==y==1 for x,y in pairs);b=sum(x==0 and y==1 for x,y in pairs)
    c=sum(x==1 and y==0 for x,y in pairs);d=sum(x==y==0 for x,y in pairs);n=a+b+c+d
    po=ratio(a+d,n);pe=ratio((a+b)*(a+c)+(d+b)*(d+c),n*n)
    return {'N':n,'TP':a,'FP':b,'FN':c,'TN':d,'reference_positive_N':a+c,
            'reference_negative_N':b+d,'prediction_positive_N':a+b,'prediction_negative_N':c+d,
            'Po':po,'Pe':pe,'kappa':None if not n else ratio(po-pe,1-pe),
            'Precision':ratio(a,a+b),'Recall':ratio(a,a+c),'F1':ratio(2*a,2*a+b+c),
            'positive_agreement':ratio(2*a,2*a+b+c),'negative_agreement':ratio(2*d,2*d+b+c),
            'specificity':ratio(d,d+b),'NPV':ratio(d,d+c)}
def reference(side):
    p=ROOT/f'09-delivery/agent-review-20261010/真人4项复核后-v2/编码复核表-{side}.xlsx'
    b=bind(p);assert b['sha256']==EXPECTED_SHA[side]
    wb=load_workbook(p,read_only=True,data_only=True);vals={}
    for row in wb['编码表'].iter_rows(min_row=7,values_only=True):
        if row[0] is None:continue
        sid=str(row[0]).strip();sid=sid if sid.startswith('RECODE25-') else 'RECODE25-'+sid.zfill(5)
        dim=str(row[2]).split('（')[0].split('(')[0].strip().split()[0]
        assert dim in DIMS and type(row[3]) in (int,float) and row[3] in (0,1)
        assert (sid,dim) not in vals;vals[(sid,dim)]=int(row[3])
    wb.close();assert len(vals)==385
    return vals,b
def directed(pred,ref):
    result={}
    for dim in DIMS:
        allocated=[(sid,v) for (sid,d),v in ref.items() if d==dim];pairs=[];missing=[];null=[]
        for sid,v in allocated:
            if sid not in pred:missing.append({'id':sid,'reference':v})
            elif pred[sid]['labels'][dim] is None:null.append({'id':sid,'reference':v})
            else:pairs.append((v,pred[sid]['labels'][dim]))
        m=metrics(pairs);m.update(allocated_N=len(allocated),allocated_reference_positive_N=sum(v for _,v in allocated),
                               allocated_reference_negative_N=sum(v==0 for _,v in allocated),
                               missing_N=len(missing),lawful_null_N=len(null),binary_coverage=ratio(len(pairs),len(allocated)),
                               excluded_positive_N=sum(x['reference'] for x in missing+null),
                               excluded_negative_N=sum(x['reference']==0 for x in missing+null),
                               missing=missing,lawful_null=null,
                               conditional_core_targets_passed=all(m[k] is not None and m[k]>=.8 for k in ('kappa','Precision','Recall','F1')),
                               formal_validation_complete=False)
        result[dim]=m
    return result
def render(v):return 'NA' if v is None else f'{v:.3f}'

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--native-snapshot',required=True)
    parser.add_argument('--rule-composite',required=True)
    parser.add_argument('--output',required=True);args=parser.parse_args()
    native_dir=Path(args.native_snapshot).resolve();out=Path(args.output).resolve()
    RULE=Path(args.rule_composite).resolve()
    # No output or statistical claim is created while the rule run is incomplete.
    execution=load(RULE/'composite-status.json');seal=load(RULE/'seal.json')
    assert (execution['status'].startswith('complete_') and execution['status'].endswith('_compatible_composite_with_explicit_execution_origins')) or execution['status']=='complete_compatible_rule_composite_with_explicit_execution_origins'
    assert execution['actual_record_N']==4481 and execution['actual_dimension_N']==22405
    assert execution['technical_failure_dimension_N']==0
    verify_composite_seal(RULE/'seal.json')
    for key in ('impact_closure_plan','adoption','original_input_lineage','merger'):exact(execution[key])
    verify_prediction_origins(RULE)
    for b in load(native_dir/'seal.json')['files']:exact(b)
    coverage=load(native_dir/'coverage.json');assert coverage['thin35_included'] is False
    assert coverage['pointer_quote_issue_N']==0
    rule=jsonlines(RULE/'predictions.jsonl');native=jsonlines(native_dir/'selected-records.jsonl');old=jsonlines(OLD)
    assert len(rule)==len(old)==4481 and set(rule)==set(old) and set(native)<=set(rule)
    assert len(native)==coverage['selected_record_N']
    assert sum(x['labels'][d] is None for x in rule.values() for d in DIMS)==execution['lawful_null_N']
    ra,ab=reference('A');rb,bb=reference('B');assert ra==rb
    assert metrics([(0,1),(1,0)])['F1']==0
    assert metrics([(0,0)])['kappa'] is None and metrics([])['Po'] is None
    out.mkdir(parents=True,exist_ok=False)
    rule_metrics=directed(rule,ra);native_metrics=directed(native,ra)
    matched_metrics={}
    for dim in DIMS:
        keys=[(sid,d) for sid,d in ra if d==dim and sid in native and rule[sid]['labels'][d] is not None and native[sid]['labels'][d] is not None]
        matched_metrics[dim]={'matched_binary_N':len(keys),
                             'rules':metrics([(ra[k],rule[k[0]]['labels'][dim]) for k in keys]),
                             'native':metrics([(ra[k],native[k[0]]['labels'][dim]) for k in keys])}
    transitions={d:collections.Counter() for d in DIMS};changed=[];disagreements=[];agreement={}
    for sid,x in rule.items():
        for dim in DIMS:
            ov=old[sid]['labels'][dim];nv=x['labels'][dim]
            transitions[dim][f'{ov}_to_{nv}']+=1
            if ov!=nv:changed.append({'id':sid,'dimension':dim,'old':ov,'new':nv,
                                      'new_decision':x['decisions'][dim],'old_decision':old[sid].get('decisions',{}).get(dim)})
    for dim in DIMS:
        pairs=[(rule[s]['labels'][dim],native[s]['labels'][dim]) for s in native
               if rule[s]['labels'][dim] is not None and native[s]['labels'][dim] is not None]
        m=metrics(pairs)
        agreement[dim]={'binary_pair_N':m['N'],'native_output_N':len(native),'native_missing_N':4481-len(native),
                        'both_1':m['TP'],'rule0_native1':m['FP'],'rule1_native0':m['FN'],'both_0':m['TN'],
                        'Po':m['Po'],'Pe':m['Pe'],'kappa':m['kappa'],'positive_agreement':m['positive_agreement'],
                        'negative_agreement':m['negative_agreement'],'is_accuracy':False}
        for sid in native:
            rv=rule[sid]['labels'][dim];nv=native[sid]['labels'][dim]
            if rv!=nv:
                disagreements.append({'id':sid,'dimension':dim,'rule':rv,'native':nv,
                                      'rule_decision':rule[sid]['decisions'][dim],
                                      'native_decision':native[sid]['dimensions'][dim],
                                      'native_first':native[sid]['first'],'packet':native[sid]['packet'],
                                      'not_adjudicated_truth':True})
    counts={d:{'positive_N':sum(r['labels'][d]==1 for r in rule.values()),
               'negative_N':sum(r['labels'][d]==0 for r in rule.values()),
               'lawful_null_N':sum(r['labels'][d] is None for r in rule.values())} for d in DIMS}
    report={'at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'identity':'complete4481 compatible rule composite with sealed individual execution origins; LLM snapshot may be partial; reference77 is disclosed development',
            'rules_full_corpus_completed':True,'native_full_corpus_completed':coverage['full_corpus_complete'],
            'native_coverage':coverage,'rules_allocated_N':4481,'rules_counts':counts,
            'rules_technical_failure_N':0,'rules_lawful_null_label_N':execution['lawful_null_N'],
            'development77_rule_metrics':rule_metrics,'development77_native_metrics':native_metrics,
            'development77_same_binary_slots':matched_metrics,'machine_pair_agreement_not_accuracy':agreement,
            'old_v5_to_current_composite_changed_label_N':len(changed),'old_v5_to_current_composite_transitions':{d:dict(v) for d,v in transitions.items()},
            'machine_disagreement_including_null_N':len(disagreements),
            'formal_validation_complete':False,'population_CI_computed_from_development':False,'paper_writing_gate_open':False,
            'reference_bindings':[ab,bb],'rule_execution':bind(RULE/'composite-status.json'),'rule_seal':bind(RULE/'seal.json'),
            'native_snapshot_seal':bind(native_dir/'seal.json'),'old_v5_predictions':bind(OLD),'calculator':bind(__file__)}
    save(out/'metrics.json',report)
    for name,rows in [('old-v5-to-current-composite-label-changes.jsonl',changed),('machine-pair-disagreements.jsonl',disagreements)]:
        with (out/name).open('x',encoding='utf8') as f:
            for row in rows:f.write(json.dumps(row,ensure_ascii=False,allow_nan=False)+'\n')
    with (out/'development77-comparison.csv').open('x',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['id','dimension','reference_A','reference_B','rules_current_compatible_composite','native_original_whole_batch','native_state'])
        for sid,d in sorted(ra):
            nv=native[sid]['labels'][d] if sid in native else None
            w.writerow([sid,d,ra[(sid,d)],rb[(sid,d)],rule[sid]['labels'][d],nv,
                        'missing_record' if sid not in native else 'lawful_null' if nv is None else 'binary'])
    lines=['规则结果完整覆盖4481条、22405项，技术失败0。逐条实际执行来源、复用依据及版本绑定见结果目录；复用结果不声称为新版本下重新运行。逐维阳性和待定数量见下表。LLM采用本次封存的原完整整批快照，覆盖'+str(len(native))+'/4481条。',
           '', '|维度|规则阳性|规则阴性|规则待定|', '|---|---:|---:|---:|']
    for d in DIMS:lines.append(f"|{d}|{counts[d]['positive_N']}|{counts[d]['negative_N']}|{counts[d]['lawful_null_N']}|")
    lines+=['','现有77条开发材料的核对结果如下，待定/缺失另列；所有指标仅在对应可判分母上计算。',
            '', '|方法|维度|可判N/77|待定|缺失|κ|精确率|召回率|F1|', '|---|---|---:|---:|---:|---:|---:|---:|---:|']
    for label,ms in [('规则',rule_metrics),('LLM完整整批快照',native_metrics)]:
        for d,m in ms.items():
            lines.append(f"|{label}|{d}|{m['N']}/77|{m['lawful_null_N']}|{m['missing_N']}|{render(m['kappa'])}|{render(m['Precision'])}|{render(m['Recall'])}|{render(m['F1'])}|")
    lines+=['','两机器相互一致衡量方法之间的相似程度，不能作为正确率；完整四格、正负类一致率、同一可判集合比较和具体分歧见metrics.json及配套记录。',
            '', '77条曾用于规则开发，因此该表不能替代独立检验，也不据此计算总体置信区间。尚未接收新的真人独立提交，正式验收与完整论文写作仍待后续步骤。', '']
    with (out/'结果说明.md').open('x',encoding='utf8') as f:f.write('\n'.join(lines))
    save(out/'seal.json',{'files':[bind(p) for p in sorted(out.iterdir()) if p.is_file()],'self_excluded':True})
    print(json.dumps({'rules_record_N':4481,'native_selected_N':len(native),'rule_changed_label_N':len(changed),
                      'rule_lawful_null_label_N':execution['lawful_null_N'],'formal_validation_complete':False,
                      'output':str(out)},ensure_ascii=False))

if __name__=='__main__':main()
