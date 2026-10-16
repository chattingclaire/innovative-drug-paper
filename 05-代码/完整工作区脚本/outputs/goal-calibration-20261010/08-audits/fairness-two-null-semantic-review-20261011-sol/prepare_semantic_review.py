#!/usr/bin/env python3
"""Read-only bounded semantic review; no classifier execution or source edits."""
import datetime
import hashlib
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
BASE = HERE.parent.parent
PACKETS = BASE / '06-coding/full4481-common-v15-production-preparation-sol/packets'
V7 = BASE / '06-coding/rule-five-dimension-joint-root/versions/joint-v7-relation-content-sol'
MANUAL = BASE / '03-manual/full4481-common-release-20261010-root/manual.txt'
BOOKS = BASE / '09-delivery/agent-review-20261010/真人4项复核后-v2'
INPUTS = {}

def bind(path):
    path = Path(path).resolve()
    raw = path.read_bytes()
    return {'path': str(path), 'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest()}

def reg(path):
    value = bind(path)
    INPUTS.setdefault(value['path'], value)
    return value

def load(path):
    reg(path)
    return json.loads(Path(path).read_bytes())

def save(name, value):
    with (HERE / name).open('x', encoding='utf8') as f:
        json.dump(value, f, ensure_ascii=False, indent=2)
        f.write('\n')

def exact_quote(path, text, label, source_class):
    reg(path)
    content = Path(path).read_text()
    start = content.index(text)
    return {'label': label, 'source_class': source_class, 'binding': bind(path),
            'line': content.count('\n', 0, start) + 1, 'start': start, 'end': start + len(text),
            'range_kind': 'zero_based_half_open_unicode', 'quote': text, 'exact': True}

manual_text = MANUAL.read_text()
reg(MANUAL)
manual_evidence = [exact_quote(MANUAL, text, label, 'existing_frozen_common_manual') for label, text in [
    ('general_F_definition', '患者按治疗需要获得或持续接受所需药品治疗的合理机会与负担，以及可避免、不合理的获取差异或保障缺口的关注。公平性与一般可及性不同。'),
    ('need_gap_link', '明确药物保障缺口与按需要的补偿、优先保障或资源分配相连；不同治疗需要/持续治疗负担与适当保障安排相连；'),
    ('three_separate_examples', '若明确的未满足需要表现为无有效治疗选项、境内无获批对应治疗选项或现实供应不足，并与针对性申请机会或资源安排直接相连，可按按需机会配置纳入；这三类缺口必须区分，不能统称短缺，单有无有效药也不自动增加F。'),
    ('qualitative_identifiable_needed_product_gap', '具体路径仍须能从当前及有限语境写出可辨认的治疗需要差别、所需药品保障缺口、获取障碍或患者负担关系，并与评价/保障安排直接相连。'),
    ('single_word_insufficient', '急需、未上市或短缺词本身不足，须原文表达保障缺口与按需获取安排的关系。'),
    ('upstream_conjunction', '上游F是本研究明确采用的按需资源机会口径：须同时保存可辨认的目标治疗需要/缺口，以及以该缺口为条件的优先、申请、支持资源机会或其限制；不得从一般创新身份、普通申报或流程位置推出。'),
    ('no_universal_option_inference', '无有效防治手段、境内未获批目标适应症选项、现实供药不足分别留痕；三者不互相推定。'),
    ('pure_administration_excluded', '纯收件、通知、签章、一般截止日期不因对象为短缺/罕见病就记F；'),
]]
records = load(V7 / 'final-development/records77.json')
selected = {r['id']: r for r in records if r['id'] in ('RECODE25-00110', 'RECODE25-00118')}
assert len(selected) == 2
reg(V7 / 'dimensions.py')
reg(V7 / 'final-development/code-snapshot/dimensions.py')
assert bind(V7 / 'dimensions.py')['sha256'] == bind(V7 / 'final-development/code-snapshot/dimensions.py')['sha256']
packet_evidence = []
packet_data = {}
for sid in selected:
    packet = load(PACKETS / (sid + '.json'))
    packet_data[sid] = packet
    source = packet['source_bundle']['source_record']
    assert selected[sid]['labels']['F'] is None and selected[sid]['decisions']['F']['value'] is None
    assert selected[sid]['text'] == source['text']
    packet_evidence.append({'id': sid, 'packet': bind(PACKETS / (sid + '.json')),
        'pointer': '/source_bundle/source_record/text', 'quote': source['text'],
        'same_current_proposition_components': {
            'official_clinically_urgent_list': '列入国家药品监督管理局《临床急需境外新药名单》' in source['text'],
            'clinically_urgent_rare_disease_product': '临床急需' in source['text'] and '罕见病药品' in source['text'],
            'same_product_foreign_marketed_domestic_unmarketed': '境外已上市境内未上市' in source['text'],
            'priority_application_or_direct_admission': ('提出优先审评审批申请' in source['text']) or ('直接纳入优先审评审批程序' in source['text'])},
        'inference_excluded': ['No effective alternative anywhere', 'No approved drug for the entire indication',
            'Physical supply shortage', 'All domestic patients have no exceptional access', 'Priority outcome already realized']})
documents = packet_data['RECODE25-00110']['source_bundle']['source_documents']
source_path = Path(documents[0]['snapshot'])
source_text = source_path.read_text()
reg(source_path)
archive_path = Path(documents[0]['original_path'])
reg(archive_path)
assert bind(source_path)['sha256'] == documents[0]['source_sha256'] == bind(archive_path)['sha256']
for packet in packet_data.values():
    chunks = packet['source_bundle']['source_full_text_resources']
    assert ''.join(x['text'] for x in chunks) == source_text
    assert all(x['source_sha256'] == documents[0]['source_sha256'] for x in chunks)
source_evidence = []
for label, prefix in [
    ('current_priority_application_same_version', '其中，对于列入'),
    ('current_direct_admission_same_version', '对于列入国家药品监督管理局《临床急需境外新药名单》的临床急需境外已上市境内未上市的罕见病药品，药审中心受理后直接纳入'),
    ('priority_resources_and_relative_time_same_version', '（六）技术审评。'),
    ('clinical_value_allows_advantage_over_existing_treatment', '（一）符合优先审评审批范围的药品上市许可申请，应具有明显临床价值'),
    ('alternatives_may_exist', '2.对于尚无有效防治手段的，该药物可以提供有效防治手段；或者与现有治疗手段相比'),
]:
    start = source_text.index(prefix)
    end = source_text.find('\n', start)
    source_evidence.append(exact_quote(source_path, source_text[start:end], label,
                                      'same_frozen_policy_version_already_inside_both_packets'))
code_evidence = [exact_quote(V7 / 'dimensions.py',
    "if re.search('境外已上市[^。；;]{0,8}境内未上市', text) and opportunity and (not any((re.search(pattern, text) for pattern in GAPS.values()))):\n        acc.boundary('F', unit, '单品种未上市与按需优先机会', '该品种境内未上市不证明对应治疗选项缺口；临床急需/罕见病及差别时限是否独立构成按需机会关系须按冻结手册统一，不能沿用错误等同。')",
    'unconditional_boundary_without_three_regex_gap_matches', 'frozen_v7_implementation')]
human = []
for name in ('A', 'B'):
    path = BOOKS / f'编码复核表-{name}.xlsx'
    reg(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb['编码表']
    for row in ws.iter_rows(min_row=7):
        values = [c.value for c in row]
        if values[0] is None:
            continue
        sid = str(values[0]).strip()
        sid = sid if sid.startswith('RECODE25-') else 'RECODE25-' + sid.zfill(5)
        dimension = str(values[2]).split('（')[0].split('(')[0].strip().split()[0]
        if sid in selected and dimension == 'F':
            assert values[3] == 1
            human.append({'existing_workbook': name, 'binding': bind(path), 'source_id': sid,
                'sheet': ws.title, 'value_cell': row[3].coordinate, 'value': values[3],
                'reason_cell': row[4].coordinate, 'reason': values[4],
                'existing_record_only_not_new_human_review': True})
            if name == 'A':
                assert selected[sid]['reference_reasons']['F'] == values[4]
    wb.close()
assert len(human) == 4
controls = [
    {'name': 'single_unmarketed_only', 'text': '境外已上市境内未上市的新药，可以申请普通审评。', 'semantic_expectation': 0,
     'boundary': 'No identified clinically required product and no need-conditioned priority arrangement'},
    {'name': 'urgent_priority_without_identified_gap', 'text': '对临床急需的罕见病药品提供优先审评，现有产品在境内已上市且供应充足。', 'semantic_expectation': 0,
     'boundary': 'Need/group and priority alone do not establish an identified access gap'},
    {'name': 'administrative_receipt_only', 'text': '对临床急需境外已上市境内未上市的罕见病药品，收到申请后发送受理通知并办理签章。', 'semantic_expectation': 0,
     'boundary': 'Same target/gap, but no substantive priority application/admission/resource opportunity'},
    {'name': 'sibling_identity_no_inheritance', 'text': '某同版名单其他条目为临床急需境外已上市境内未上市药品；本条仅登记另一产品的名单身份。', 'semantic_expectation': 0,
     'boundary': 'Need, gap and arrangement cannot be joined across other sibling subjects or bare identity'},
    {'name': 'qualified_application_relation', 'text': '对于临床急需的境外已上市境内未上市罕见病药品，申请人可以提出优先审评申请。', 'semantic_expectation': 1,
     'boundary': 'All three relation components refer to the same identified needed product'},
    {'name': 'qualified_direct_admission_relation', 'text': '对于临床急需的境外已上市境内未上市罕见病药品，受理后直接纳入优先审评程序。', 'semantic_expectation': 1,
     'boundary': 'Substantive need-conditioned admission, not receipt alone'},
]
save('input-seal.json', {'created_at_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
                        'files': list(INPUTS.values()), 'input_N': len(INPUTS)})
save('evidence.json', {'manual': manual_evidence, 'current_packets': packet_evidence,
    'same_version_policy': source_evidence, 'implementation': code_evidence,
    'existing_human_reference_cells': human,
    'version_note': 'The decision uses only the policy text already frozen inside both packets. Filename year is not separately treated as verified publication date. No later explanation or external new material was used.',
    'packet_record_punctuation_and_primary_verbatim_source_preserved_separately': True})
save('counterexample-boundaries.json', {'cases': controls, 'execution_status': 'Semantic expectations fixed before any candidate code; not executed by this semantic review',
                                       'examples_are_constructed_not_new_policy_or_human_labels': True})
save('report.json', {'created_at_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
    'conclusion': 'Existing general F definition supports a bounded operationalization of identified needed-product domestic marketing availability gap linked to priority opportunity; v7 three-regex gate is narrower than that relation.',
    'requires_substantive_definition_revision_first': False,
    'interpretation_not_claimed_as_only_possible_literal_reading': True,
    'supporting_reading': ['manual general needed-medicine opportunity and identifiable safeguard gap',
        'same-proposition recognized clinical need + specific product domestic unmarketed + priority application/admission',
        'same-version explicit resource priority and 130/70-day comparison; no guessed time advantage'],
    'opposing_reading': ['manual examples use corresponding/target-indication treatment options',
        'urgency, unmarketed status and rare-disease identity alone are excluded',
        'same-version priority eligibility allows clinical advantage over available alternatives'],
    'resolution': 'Keep the product availability gap separate from absence of all effective or indication-approved alternatives. The opposing evidence prohibits universal-alternative inference but does not defeat the fully specified needed-product gap/opportunity relation.',
    'recommendation': 'Preserve v7 nulls and all production inputs. In a separate candidate, replace only the disputed boundary with a generic three-part same-subject relation and retain targeted counterexamples; record the interpretation transparently before finite replay.',
    'reference_contact': 'Existing A/B values/reasons known and inspected; disclosed finite semantic review, not blind independent human recoding.',
    'new_human_review_performed_or_claimed': False, 'existing_labels_changed': False,
    'code_or_manual_or_LLM_inputs_changed_by_semantic_review': False,
    'classifier_executions': 0, 'model_calls': 0, 'external_browsing_or_new_material_used': False,
    'kappa_target_used_as_decision_rule': False, 'input_N': len(INPUTS)})
changed = [x for x in INPUTS.values() if bind(x['path']) != x]
assert not changed
save('input-integrity-readback.json', {'input_N': len(INPUTS), 'mismatch_N': 0,
                                     'checked_at_utc': datetime.datetime.now(datetime.timezone.utc).isoformat()})
print(json.dumps({'input_N': len(INPUTS), 'input_mismatch_N': 0,
                  'existing_A_B_F_cells_checked_N': len(human), 'counterexample_and_positive_control_N': len(controls)}))
