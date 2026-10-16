"""Finite, disclosed development replay. Expectations are never coder inputs."""
import copy, csv, datetime, hashlib, importlib.util, json, sys
from pathlib import Path
sys.dont_write_bytecode=True
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'06-coding/rule-five-dimension-joint-root/versions'
NAMES=('evidence','source_packet','source_receipts','relations','dimensions')
DIMS=('S','E','Ec','I','F')
CASES=[
 ('payment downward','药品支付标准下调10个百分点。','Ec',1),
 ('payment upward','药品支付标准的下调比例增加6个百分点。','Ec',1),
 ('negotiated decline','药品续约可以申请通过重新谈判确定下调比例。','Ec',1),
 ('supply cost','优化药品配送，提高效率，降低供应成本。','Ec',1),
 ('drug selection fund constraint','以医保基金安全为底线，对进入目录的药品控制锦上添花。','Ec',1),
 ('general fund IT security','药品管理部门应保证医保基金信息系统安全。','Ec',0),
 ('administrative cost','药品管理办公室应降低办公成本。','Ec',0),
 ('bare decline','药品文件的字体下调10个百分点。','Ec',0),
 ('SE content S','药品申报材料内容涉及有效性信息、安全性信息，摘要应与申报材料一致。','S',1),
 ('SE content E','药品申报材料内容涉及有效性信息、安全性信息，摘要应与申报材料一致。','E',1),
 ('S field','药品申报表中安全性字段应填满并签字。','S',0),
 ('E formatting','药品申报表中有效性栏目仅检查格式。','E',0),
 ('risk actual guidance','为临床试验各方将全生命周期风险管理理念贯穿药物研发全过程提供指导。','S',1),
 ('risk title only','药品文件名为《研发风险管理技术指导原则》。','S',0),
 ('risk dataset','药物研究数据丢失风险应加强管理。','S',0),
 ('purpose only','为保障用药安全，制定本规范。','S',0),
 ('general health purpose','药品指南旨在提升老年患者整体健康水平，关注老年需求和偏好。','E',0),
 ('route identity','药品企业主动申请调整医保支付范围。','Ec',0),
]
def bind(p):
 p=Path(p).resolve();b=p.read_bytes();return {'path':str(p),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()}
def save(p,x):
 with Path(p).open('x',encoding='utf8') as f:json.dump(x,f,ensure_ascii=False,indent=2);f.write('\n')
def modules(code):
 got={}
 for n in NAMES:
  s=importlib.util.spec_from_file_location(n,code/(n+'.py'));m=importlib.util.module_from_spec(s);sys.modules[n]=m;s.loader.exec_module(m);got[n]=m
 return got
def main():
 version,suffix=sys.argv[1:3];code=BASE/version;out=code/suffix;out.mkdir(exist_ok=False)
 snapshot=out/'code-snapshot';snapshot.mkdir()
 for n in NAMES:(snapshot/(n+'.py')).write_bytes((code/(n+'.py')).read_bytes())
 save(out/'pre-execution.json',{'at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'runner':bind(__file__),'code':[bind(code/(n+'.py')) for n in NAMES],'cases':CASES,'identity':'finite disclosed development; not independent validation'})
 m=modules(snapshot);results=[]
 for name,text,d,expected in CASES:
  u={'text':text,'contexts':[],'source_ref':{'sha256':'a'*64,'source_resource_id':'SOURCE-'+'a'*64,'locator':{}}}
  try:
   v=m['relations'].RelationView(u);ds=m['dimensions'].judge(v);m['evidence'].validate_decisions(u,ds)
   results.append({'case':name,'text':text,'dimension':d,'expected':expected,'actual':ds[d]['value'],'passed':ds[d]['value']==expected,'decisions':ds})
  except Exception as e:results.append({'case':name,'passed':False,'error':repr(e)})
 save(out/'counterpairs.json',results)
 # Only the loader sees workbook values. They never enter semantic_unit/judge.
 from openpyxl import load_workbook
 review=ROOT/'09-delivery/agent-review-20261010/真人4项复核后-v2/编码复核表-A.xlsx'
 wb=load_workbook(review,read_only=True,data_only=True);refs={};reasons={}
 for r in wb['编码表'].iter_rows(min_row=7,values_only=True):
  if r[0] is None:continue
  raw_id=str(r[0]).strip();sid=raw_id if raw_id.startswith('RECODE25-') else 'RECODE25-'+raw_id.zfill(5)
  d=str(r[2]).split('（')[0].split('(')[0].strip().split()[0]
  assert d in DIMS
  refs[(sid,d)]=int(r[3]);reasons[(sid,d)]=r[4]
 wb.close();records=[]
 packets=ROOT/'06-coding/full4481-common-v15-production-preparation-sol/packets'
 save(out/'review-input-bindings.json',{'reference':bind(review),'packets':[bind(packets/(sid+'.json')) for sid in sorted({k[0] for k in refs})],'disclosed_reference_not_classifier_input':True})
 for sid in sorted({k[0] for k in refs}):
  p=packets/(sid+'.json');raw=json.loads(p.read_bytes());u=m['source_packet'].semantic_unit(raw['source_bundle']);before=copy.deepcopy(u)
  v=m['relations'].RelationView(u);ds=m['dimensions'].judge(v);m['evidence'].validate_decisions(u,ds);assert u==before
  records.append({'id':sid,'labels':{d:ds[d]['value'] for d in DIMS},'text':u['text'],'decisions':ds,'source_requests':v.source_requests,'material_requests':v.material_requests,'reference':{d:refs[(sid,d)] for d in DIMS},'reference_reasons':{d:reasons[(sid,d)] for d in DIMS},'packet':bind(p)})
 save(out/'records77.json',records)
 def metric(rows,d):
  vals=[(r['reference'][d],r['labels'][d]) for r in rows if r['labels'][d] in (0,1)];tp=sum(a==b==1 for a,b in vals);fp=sum(a==0 and b==1 for a,b in vals);fn=sum(a==1 and b==0 for a,b in vals);tn=sum(a==b==0 for a,b in vals);n=len(vals);po=(tp+tn)/n;pe=((tp+fn)*(tp+fp)+(tn+fp)*(tn+fn))/n**2
  div=lambda a,b:None if b==0 else a/b
  return {'N':n,'null':len(rows)-n,'TP':tp,'FP':fp,'FN':fn,'TN':tn,'kappa':div(po-pe,1-pe),'P':div(tp,tp+fp),'R':div(tp,tp+fn),'F1':div(2*tp,2*tp+fp+fn),'Po':po,'specificity':div(tn,tn+fp),'NPV':div(tn,tn+fn),'positive_agreement':div(2*tp,2*tp+fp+fn),'negative_agreement':div(2*tn,2*tn+fp+fn)}
 stats={d:metric(records,d) for d in DIMS};mismatch=[{'id':r['id'],'dimension':d,'reference':r['reference'][d],'prediction':r['labels'][d],'text':r['text'],'reference_reason':r['reference_reasons'][d]} for r in records for d in DIMS if r['labels'][d]!=r['reference'][d]]
 summary={'identity':'development only; no formal population inference','counterpair_N':len(results),'counterpairs_passed':sum(r['passed'] for r in results),'counterpairs_failed':[r['case'] for r in results if not r['passed']],'metrics':stats,'mismatches_including_null':mismatch,'reference_unchanged':True,'formal_validation_complete':False}
 save(out/'summary.json',summary);save(out/'seal.json',{'files':[bind(p) for p in sorted(out.rglob('*')) if p.is_file()],'self_excluded':True})
 print(json.dumps({k:v for k,v in summary.items() if k!='mismatches_including_null'},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
