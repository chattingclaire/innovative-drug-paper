from pathlib import Path
import json,hashlib,re
from docx import Document
from pypdf import PdfReader

ROOT=Path('/Users/rainbow/Desktop/战略院/药企项目-创新药')
OUT=ROOT/'outputs/goal-calibration-20261010/parallel-manuscript'
ARC=Path('/Users/rainbow/Desktop/战略院/创新药论文-完整归档-20260914')
paths=list((ROOT/'outputs/final-three-department-20261008-human-validation-template').glob('*.docx'))+list((ROOT/'outputs/final-three-department-20261008-adopted-results').glob('*.docx'))
paths+=list((ARC/'12-投稿定稿工作区').glob('*.docx'))+list((ARC/'12-投稿定稿工作区/期刊论述修订-20260927').glob('*.docx'))
paths+=[ARC/'12-投稿定稿工作区/原文补全/科学决策投稿要求-20260925/科学决策稿件模板及规范要求-2026年9月修订.docx']
paths+=list((ROOT/'outputs/编码修订核查-20261008/本轮结果').glob('*.docx'))
rows=[]
for i,p in enumerate(paths,1):
    d=Document(p); paras=[x.text for x in d.paragraphs]; tables=[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    refstart=next((j for j,t in enumerate(paras) if t.strip() in ['参考文献','参考文献：','参考文献:']),len(paras))
    refs=[]
    for j,t in enumerate(paras[refstart+1:],refstart+1):
        m=re.match(r'^\[(\d+)\](.*)',t)
        if m: refs.append({'number':int(m[1]),'text':m[2],'paragraph_1based':j+1})
    body='\n'.join(paras[:refstart])+'\n'+'\n'.join(c for tb in tables for r in tb for c in r)
    cite={}; first=[]
    for j,t in enumerate(paras[:refstart],1):
        for m in re.finditer(r'\[(\d+(?:\s*[-–,，]\s*\d+)*)\]',t):
            ids=[]
            for part in re.split('[,，]',m[1]):
                ab=re.split('[-–]',part)
                ids+=list(range(int(ab[0]),int(ab[1])+1)) if len(ab)==2 else [int(ab[0])]
            for n in ids:
                cite.setdefault(n,[]).append({'paragraph_1based':j,'text':t})
                if n not in first: first.append(n)
    for k,tb in enumerate(tables,1):
        for r in tb:
            for t in r:
                for m in re.finditer(r'\[(\d+)\]',t):
                    cite.setdefault(int(m[1]),[]).append({'table_1based':k,'text':t})
    rid=f'D{i:02}'
    row={'id':rid,'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'paragraphs':len(paras),'tables':len(tables),'inline_shapes':len(d.inline_shapes),'title':paras[0] if paras else '', 'reference_count':len(refs),'references':refs,'citations':cite,'uncited':[x['number'] for x in refs if x['number'] not in cite],'undefined':sorted(set(cite)-{x['number'] for x in refs}),'first_citation_order':first,'han_before_references':len(re.findall('[\u4e00-\u9fff]',body)),'nonwhite_before_references':len(re.sub(r'\s','',body)),'headings':[{'p':j+1,'text':t} for j,t in enumerate(paras[:refstart]) if re.match(r'^\d+(\.\d+)*\s',t) or re.match(r'^附录',t)]}
    rows.append(row)
    text='\n'.join(f'P{j+1:03} {t}' for j,t in enumerate(paras))+'\n\n'+'\n\n'.join(f'TABLE {k+1}\n'+'\n'.join(' | '.join(r) for r in tb) for k,tb in enumerate(tables))
    (OUT/'extracts'/f'{rid}.txt').write_text(text)
    (OUT/'extracts'/f'{rid}.json').write_text(json.dumps({'paragraphs':paras,'tables':tables},ensure_ascii=False,indent=2))
(OUT/'document_inventory.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
for x in rows: print(json.dumps({k:x[k] for k in ['id','path','sha256','paragraphs','reference_count','uncited','undefined','headings']},ensure_ascii=False))
pdfrows=[]
for p in (ARC/'12-投稿定稿工作区/期刊论述修订-20260927/样文研读').glob('*.pdf'):
    r=PdfReader(p); pages=[x.extract_text() for x in r.pages]; txt='\n'.join(f'PAGE {j+1}\n{t}' for j,t in enumerate(pages))
    (OUT/'extracts'/f'{p.stem}.txt').write_text(txt)
    pdfrows.append({'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'pages':len(pages),'han':len(re.findall('[\u4e00-\u9fff]',txt)),'extract':str(OUT/'extracts'/f'{p.stem}.txt')})
(OUT/'sample_pdf_inventory.json').write_text(json.dumps(pdfrows,ensure_ascii=False,indent=2))
