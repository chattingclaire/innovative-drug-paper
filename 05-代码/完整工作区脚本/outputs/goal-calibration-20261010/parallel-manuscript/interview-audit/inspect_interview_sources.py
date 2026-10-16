from pathlib import Path
from docx import Document
import hashlib,json,shutil,datetime,re

OUT=Path(__file__).resolve().parent
PROJECT=OUT.parents[3]
ARCHIVE=Path('/Users/rainbow/Desktop/战略院/创新药论文-完整归档-20260914')
RAW=ARCHIVE/'04-访谈与研讨会原始资料'
WORK=ARCHIVE/'12-投稿定稿工作区'
BASE=PROJECT/'outputs/编码修订核查-20261008/本轮结果/创新药评价的注意力配置差异与协调路径.docx'
REPORT=PROJECT/'2025-09-24 研究报告合稿-打印（检查版）(4).docx'
SPEC=[
 ('M00',BASE,'current_adopted_manuscript','当前采用稿，仅核查不修改'),
 ('I01',RAW/'史录文教授访谈会议纪要-0605.docx','expert_interview_minutes','当前稿两份专家纪要之一'),
 ('I02',RAW/'医保局专家调研-0616.docx','expert_research_minutes','当前稿两份专家纪要之一'),
 ('W01',RAW/'创新药政策（百济神州）会议纪要-250212.docx','meeting_minutes','三场会议之一'),
 ('W02T',RAW/'创新药专家咨询会-0319.docx','timestamped_meeting_transcription_text','三场会议之一，同场转录文本'),
 ('W02S',RAW/'创新药专家研讨会-0319.docx','meeting_summary_minutes','3月19日同场整理纪要'),
 ('W03A',RAW/'0514创新药研讨会议程.docx','planned_meeting_agenda','三场会议之一，议程仅证计划'),
 ('W03M',RAW/'创新药研讨论会议纪要-0514.docx','meeting_minutes','5月14日实际内容纪要'),
 ('R01',REPORT,'research_report_with_interview_appendix','附件重复存档与作者汇总；不新增独立访谈'),
 ('A01',WORK/'访谈执行与作者信息补充核查.json','prior_audit_record','指定接收元数据'),
 ('A02',WORK/'用户补充研究报告-访谈依据接收-20260927.json','prior_receipt_record','指定接收元数据'),
 ('A03',WORK/'核实访谈来源.py','prior_source_check_script','只读源代码，禁止执行其修改旧文件的逻辑'),
 ('A04',WORK/'访谈匿名公开使用-作者确认.json','recorded_author_fact_confirmation','负责人同意匿名公开使用的既有确认记录'),
 ('A05',WORK/'研究报告正文附件-论文论证对应.json','prior_report_mapping','既有报告附件映射，仅作检索线索'),
 ('A06',PROJECT/'outputs/final-three-department-20261007/prose_evidence_QA_refs61-63_interviews.json','prior_interview_audit','三场会议绑定与旧正文审查，不替代本次原文检查'),
 ('A07',PROJECT/'outputs/final-three-department-20261007/论证证据索引.json','prior_argument_index','旧版论证使用位置，不直接冒用当前段号'),
]
for folder in ('snapshots','extracts'):(OUT/folder).mkdir(parents=True,exist_ok=True)
created=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat()
sources=[]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def extract(p,sid):
 d=Document(p)
 paragraphs=[{'paragraph_1based':i,'text':x.text} for i,x in enumerate(d.paragraphs,1)]
 tables=[{'table_1based':i,'rows':[[c.text for c in row.cells] for row in t.rows]} for i,t in enumerate(d.tables,1)]
 payload={'source_id':sid,'source_path':str(p),'source_sha256':sha(p),'locator_convention':'P为python-docx直接正文段落1基编号（含空段）；T为表1基编号；不是页码或逐字发言序号','paragraphs':paragraphs,'tables':tables}
 (OUT/'extracts'/f'{sid}.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2))
 (OUT/'extracts'/f'{sid}.txt').write_text('\n'.join(f"P{x['paragraph_1based']:03d} {x['text']}" for x in paragraphs)+'\n'+''.join(f"\nT{x['table_1based']:03d} "+json.dumps(x['rows'],ensure_ascii=False) for x in tables))
 return {'paragraph_count':len(paragraphs),'table_count':len(tables),'text_cjk_chars':len(re.findall(r'[\u4e00-\u9fff]',''.join(x['text'] for x in paragraphs))),'extract_json_path':str(OUT/'extracts'/f'{sid}.json'),'extract_text_path':str(OUT/'extracts'/f'{sid}.txt')}
for sid,p,kind,role in SPEC:
 if not p.exists():
  sources.append({'source_id':sid,'source_path':str(p),'status':'missing','source_kind':kind,'use_role':role})
  continue
 snapshot=OUT/'snapshots'/f'{sid}__{p.name}'
 if snapshot.exists():assert sha(snapshot)==sha(p),(sid,'existing snapshot differs')
 else:shutil.copyfile(p,snapshot);snapshot.chmod(0o444)
 record={'source_id':sid,'source_path':str(p),'snapshot_path':str(snapshot),'sha256':sha(p),'snapshot_sha256':sha(snapshot),'size_bytes':p.stat().st_size,'source_mtime':datetime.datetime.fromtimestamp(p.stat().st_mtime,datetime.timezone(datetime.timedelta(hours=8))).isoformat(),'source_kind':kind,'use_role':role,'status':'received_and_hash_verified'}
 if p.suffix=='.docx':record.update(extract(snapshot,sid))
 sources.append(record)
known={Path(x['source_path']).name:x['source_id'] for x in sources}
context=[]
for i,p in enumerate(sorted(RAW.glob('*.docx')),1):
 if p.name in known:continue
 sid=f'C{i:02d}'
 record={'source_id':sid,'source_path':str(p),'sha256':sha(p),'source_kind':'context_document_not_bound_to_current_three_meetings','use_role':'仅用于排除把目录全部材料计作当前已采用资料','status':'context_read_only_not_counted_as_adopted'}
 record.update(extract(p,sid));context.append(record)
manifest={'created_at':created,'audit_scope':'当前采用稿访谈与研讨证据接收核查；无正文改写','manuscript_sha256_expected':'2335ad123737735ea5ae32234de323d3c7ae1ac0c23136d4d2eb7a2742b15db6','source_locator_convention':'P1基直接段落含空段；T1基表格；不是页码','source_files_modified':False,'previous_sealed_outputs_modified':False,'sources':sources,'context_inventory':context,'adopted_event_groups':{'interview_records':['I01','I02'],'meeting_events':[{'event_id':'EV01','date':'2025-02-12','records':['W01']},{'event_id':'EV02','date':'2025-03-19','records':['W02T','W02S']},{'event_id':'EV03','date':'2025-05-14','records':['W03A','W03M']}],'report_repeats':['R01']} }
assert sources[0]['sha256']==manifest['manuscript_sha256_expected']
(OUT/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'received':len(sources),'context':len(context),'all_sources_found':all(x['status']=='received_and_hash_verified' for x in sources),'current_manuscript_hash':sources[0]['sha256']},ensure_ascii=False))
