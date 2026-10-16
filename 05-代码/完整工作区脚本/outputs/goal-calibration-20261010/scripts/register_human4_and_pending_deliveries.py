"""Register actual artifacts and keep full-paper Goal active."""
import hashlib
import json
import math
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

ROOT=Path(__file__).resolve().parents[1]


def read(p):return json.loads(Path(p).read_text())


def bind(p):
    p=Path(p).resolve()
    return {'path':str(p),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}


def check(item,base=None):
    p=Path(item.get('path') or item.get('relative_path'))
    if not p.is_absolute():p=base/p
    b=bind(p);assert b['sha256']==item['sha256'],str(p)
    if 'bytes' in item:assert b['bytes']==item['bytes']


def main():
    now=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat()
    support=ROOT/'logs/human4-machine77-report-delivery'
    a=read(support/'authoring-verification.json');check({'path':a['output'],'sha256':a['sha256']})
    inp=read(support/'report-inputs.json')
    w=openpyxl.load_workbook(a['output'],data_only=True)
    raw=openpyxl.load_workbook(a['output'],data_only=False)
    expected=[x[:10]+[x[10] if len(x[10])<=1200 else x[10][:1200]+'\n〔其余引文及完整字段见初判文件〕']+x[11:] for x in inp['rows']]
    assert [list(x) for x in w['编码对照'].iter_rows(min_row=7,max_row=391,max_col=14,values_only=True)]==expected
    formula_n=sum(c.data_type=='f' for row in raw['指标'].iter_rows() for c in row)
    assert formula_n==190
    count=0
    for mi,key in enumerate(['native_dimensions','rule_dimensions']):
        for di,dim in enumerate(inp['dimensions']):
            top=6+mi*5+di; cells=21+mi*5+di;aux=36+mi*5+di;exp=inp[key][dim]
            for j,n in enumerate(['kappa','precision','recall','F1']):
                assert math.isclose(w['指标'].cell(top,3+j).value,exp['metrics'][n],abs_tol=1e-12);count+=1
            for j,n in enumerate(['TP','FP','FN','TN']):assert w['指标'].cell(cells,6+j).value==exp['four_cells'][n]
            for j,n in enumerate(['Po','Pe','Ppos','Pneg','specificity','NPV']):
                assert math.isclose(w['指标'].cell(aux,3+j).value,exp['metrics'][n],abs_tol=1e-12);count+=1
    assert count==100
    for sheet in w:
        for row in sheet.iter_rows():
            for c in row:assert c.data_type!='e'
    saved={'output':bind(a['output']),'computed_formulas':formula_n,'metric_crosschecks':count,
           'four_cells_crosschecks':40,'reference_and_preserved_machine_labels':1155,
           'all_rendered_worksheets_visually_checked':True,'formula_errors':0,
           'formal_acceptance':False,'old_raw_machine_firsts_changed':False}
    with (support/'saved-workbook-verification.json').open('x') as f:f.write(json.dumps(saved,ensure_ascii=False,indent=2)+'\n')
    reception={}
    for label,directory in [('integrated_year','01-review/annual-metadata-integrated-sol'),
                            ('request_route','05-validation/rule-request-routing-repair-sol'),
                            ('independent_route77','06-coding/rule-routing-dev77-independent-sol')]:
        base=ROOT/directory;seal=read(base/'seal.json')
        for item in seal['files']:check(item,base)
        if label=='integrated_year':
            for item in seal['inputs']:check(item,base)
        reception[label]={'seal':bind(base/'seal.json'),'output_bindings_verified_N':len(seal['files']),
                          'input_bindings_verified_N':len(seal.get('inputs',[])),
                          'formal_accuracy_acceptance':False}
    base=ROOT/'07-statistics/prospective-engine-year-independent-audit-sol'
    ins=read(base/'input-seal.json');outs=read(base/'output-seal.json');n=0
    for group in ['read_only_original_files','constructed_inputs','constructed_frozen_source_year_maps']:
        for p,h in ins[group].items():check({'path':p,'sha256':h},base);n+=1
    assert outs['input_seal_sha256']==bind(base/'input-seal.json')['sha256']
    for p,h in outs['files'].items():check({'path':p,'sha256':h},base)
    reception['independent_year_repair']={'input_bindings_verified_N':n,'output_bindings_verified_N':len(outs['files']),
                                          'repair_binding':bind(base/'metrics_engine_repaired.py'),
                                          'audit_binding':bind(base/'audit.json'),'actual_project_year_results_computed':False}
    metrics_path=ROOT/'07-statistics/human4-reference-v2-raw77-root/metrics.json'
    metrics=read(metrics_path)
    reference=read(ROOT/'logs/human4-reference-v2-delivery/authoring-verification.json')
    receipt={'received_at':now,'human_message_verbatim':'已经核对都是1','human_targeted_slots_N':4,
             'slots':[{'source_id':'RECODE25-00005','dimension':'S','value':1},
                      {'source_id':'RECODE25-03876','dimension':'Ec','value':1},
                      {'source_id':'RECODE25-00138','dimension':'I','value':1},
                      {'source_id':'RECODE25-00138','dimension':'F','value':1}],
             'reference_identity':'AI-seeded A/B plus targeted human posthoc corrections',
             'reason_texts_AI_drafted_from_same_version_sources_not_verbatim_human':True,
             'independent_human_initial_coding_complete':False,'new_reference_outputs':reference['outputs'],
             'metrics_binding':bind(metrics_path),'report_workbook':saved,
             'old_LLM_errors_N':1,'old_rule_errors_N':4,'all_development_core_point_targets_pass':True,
             'recomputed_existing_predictions_only':True,'full4481_started':False,'formal_acceptance':False,
             'other_actual_deliveries_received':reception}
    target=ROOT/'logs/root-human4-and-current-deliveries-reception-20261010.json'
    with target.open('x') as f:f.write(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    status=ROOT/'09-delivery/agent-review-20261010/真人4项复核后-v2/配套检验状态.txt'
    with status.open('x') as f:f.write(
        '最新真人反馈：剩余4项均为1，已另存A/B参照v2；新增理由由AI按完整同版来源整理。\n'
        '已完成：77条/385项原LLM与原规则相对新参照逐维κ、P、R、F1、四格、观察/机会一致率、正负类一致率、特异度和负预测值复算，及执行覆盖核验。所有核心开发点指标≥0.8。LLM保留1项误判，规则保留4项漏判。\n'
        '该复算沿用原封存机器预测，并非新共同材料上的重跑；A/B相同值κ=1不是独立真人初判信度。\n'
        '进行中：共同手册与prompt通用边界澄清、规则语义遗漏修复和新共同材料统一；候选未最终冻结。\n'
        '已完成技术准备：规则缺材请求路由构造控制和一次开发77回归；年度统计引擎独立构造审查及修复候选；51源条目和4481无标签ID年度映射。\n'
        '待完成：正式800独立人类初判与裁定、正式机器验证及设计适用95%区间、LLM相同设置独立重复、4481全量重编码及结果稳健性、完整论文。\n'
        '未计算编码合格p值；95%区间未从本开发样本推断总体。年份元数据核定不等于年度注意力结果已经计算。\n')
    p=ROOT/'progress.json'
    with (ROOT/'logs/human4-reference-v2-delivery/progress-before-human4-reception.json').open('xb') as f:f.write(p.read_bytes())
    progress=read(p)
    progress['human4_posthoc_review_received_actual']=receipt
    audit=progress['uniform77_disagreement_audit_actual']
    audit.update({'status':'all_4_targeted_human_slots_received_applied_new_reference_v2',
                  'remaining_cases':[],'extra_human_request_slots_N':0,
                  'new_reference_applied':True,'old_reference_and_machine_firsts_modified':False,
                  'human4_receipt':str(target.relative_to(ROOT))})
    progress['remaining4_calibration_delivery_actual'].update({'human_response_received':True,
                                                              'new_reference_version_outputs':reference['outputs'],
                                                              'original_blank_calibration_workbook_retained_as_history':True})
    progress['latest_development_metric_status_actual']={'scope':'77 development posthoc reference-v2 vs preserved old raw77',
                 'metrics':str(metrics_path.relative_to(ROOT)),'workbook':str(Path(a['output']).relative_to(ROOT)),
                 'all_core_point_targets_pass':True,'LLM_disagreement_N':1,'rule_disagreement_N':4,
                 'formal_acceptance':False,'new_common_material_rerun':False,'full4481_started':False}
    progress['companion_reporting_human4_v2_actual']={'status_file':str(status.relative_to(ROOT)),
                                                    'workbook':str(Path(a['output']).relative_to(ROOT)),
                                                    'population_CI95_calculated':False,'p_values_calculated':False}
    progress['current_human4_candidate_alignment_actual']={'manual_prompt_task':'/root/human4_manual_prompt_alignment_sol',
                   'rule_task':'/root/human4_general_rule_repair_sol','status':'running_constructed_controls_only_no_real_rerun_authorized',
                   'need_shared_material_projection_for_same_PDF_price_definitions':True,'formal_freeze_complete':False}
    progress['annual_integrated51_reception_actual']=reception['integrated_year']
    progress['annual_engine_independent_repair_reception_actual']=reception['independent_year_repair']
    progress['rule_route_independent77_reception_actual']=reception['independent_route77']
    progress['department_year_conclusion_requirement_actual'].update({'strict_publication_dates_pending_entries_N':0,
                                  'integrated_51_source_year_mapping_status':'completed_received_hash_verified_no_labels',
                                  'actual_year_attention_results_calculated':False})
    progress['current_goal_turn_classification']={'classification':'progress','registered_at':now,
                  'evidence':[str(target.relative_to(ROOT)),str(metrics_path.relative_to(ROOT))],
                  'formal_acceptance':False,'phase7_full_manuscript_started':False}
    progress['last_integration_at']=now
    p.write_text(json.dumps(progress,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'receipt':str(target),'human4_saved':True,'all_development_targets_pass':True,
                      'sealed_other_deliveries':reception,'full4481_started':False},ensure_ascii=False))


if __name__=='__main__':main()
