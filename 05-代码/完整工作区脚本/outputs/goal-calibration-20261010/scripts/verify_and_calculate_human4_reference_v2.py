"""Verify narrow workbook edits, preserved firsts, and new posthoc arithmetic."""
import hashlib
import importlib.util
import json
import math
from pathlib import Path

import openpyxl
from sklearn.metrics import cohen_kappa_score, confusion_matrix, precision_recall_fscore_support

ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / "logs/human4-reference-v2-delivery"
CALC = ROOT / "07-statistics/human4-reference-v2-raw77-root"
REPORT = ROOT / "logs/human4-machine77-report-delivery"
DIMS = ("S", "E", "Ec", "I", "F")


def read(p):
    return json.loads(Path(p).read_text())


def bind(p):
    p = Path(p).resolve()
    return {"path":str(p),"bytes":p.stat().st_size,
            "sha256":hashlib.sha256(p.read_bytes()).hexdigest()}


def main():
    edits = read(SUPPORT / "edit-inputs.json")
    authored = read(SUPPORT / "authoring-verification.json")
    refs, comparisons = [], []
    errors = {"#REF!","#DIV/0!","#VALUE!","#NAME?","#NUM!","#N/A","#NULL!","#SPILL!","#CALC!"}
    for job, saved in zip(edits["jobs"],authored["outputs"]):
        assert bind(job["source"]["path"]) == job["source"]
        assert bind(job["target"]) == saved
        old = openpyxl.load_workbook(job["source"]["path"],data_only=False)
        new = openpyxl.load_workbook(job["target"],data_only=False)
        assert old.sheetnames == new.sheetnames == ["编码表","原文语境"]
        changed = []
        allowed = {("编码表","A4"),("原文语境","A4")}
        allowed |= {("编码表",f"{c}{x['row']}") for x in job["changes"] for c in "DEH"}
        for name in old.sheetnames:
            for row in old[name].iter_rows():
                for cell in row:
                    value=new[name][cell.coordinate].value
                    if cell.value != value:
                        assert (name,cell.coordinate) in allowed,(name,cell.coordinate)
                        changed.append([name,cell.coordinate])
            for row in new[name].iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value,str) and cell.value in errors)
        assert len(changed)==14
        current={}
        for row in new["编码表"].iter_rows(min_row=7,max_row=391,values_only=True):
            sid,dept,dim,value = row[:4]; dim=dim.split()[0]
            assert type(value) is int and value in (0,1)
            current.setdefault(sid,{"department":dept,"labels":{}})["labels"][dim]=value
        assert len(current)==77 and all(set(x["labels"])==set(DIMS) for x in current.values())
        for change in job["changes"]:
            row=change["row"]
            assert new["编码表"].cell(row,4).value==1
            assert new["编码表"].cell(row,5).value==change["reason"]
            assert new["编码表"].cell(row,8).value==change["context"]
        appendix=list(new["原文语境"].iter_rows(min_row=851,max_row=893,max_col=6,values_only=True))
        assert [list(x) for x in appendix]==edits["appendix_rows"]
        refs.append(current)
        comparisons.append({"output":saved,"original_binding":job["source"],
                            "changed_existing_cells":changed,"all_other_existing_cell_values_preserved":True,
                            "appendix_rows_verified":43,"binary_labels_N":385})
    assert refs[0]==refs[1]
    formula_path=ROOT/'scripts/verify_native_v16_campaign.py'
    spec=importlib.util.spec_from_file_location('native_math',formula_path)
    m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
    old_native=read(ROOT/'07-statistics/native-v16-current-posthoc-reference-root/verification-and-metrics.json')
    old_rule=read(ROOT/'07-statistics/rule-v18-current-posthoc-reference-root/metrics.json')
    for b in [old_native['predictions_binding'],old_rule['prediction_binding'],old_native['run_manifest_binding']]:
        m.check_binding(b)
    manifest=read(old_native['run_manifest_binding']['path']);native=[];batches=[]
    for batch in manifest['batches']:
        verified,records=m.verify_batch(ROOT/'06-coding/native-llm77-v16-sol',batch)
        assert verified['state']=='sealed_verified'
        batches.append(verified);native.extend(records)
    stored_native=[json.loads(x) for x in Path(old_native['predictions_binding']['path']).read_text().splitlines()]
    assert native==stored_native
    rule=[json.loads(x) for x in Path(old_rule['prediction_binding']['path']).read_text().splitlines()]
    assert [x['id'] for x in rule]==[x['id'] for x in native]==manifest['global_order']
    dimensions={};differences={};crosschecks=[]
    for name,rows in [('native',native),('rule',rule)]:
        dimensions[name],differences[name]=m.calculate(rows,refs[0])
        for dim,result in dimensions[name].items():
            truth=[refs[0][x['id']]['labels'][dim] for x in rows]
            pred=[x['labels'][dim] for x in rows]
            k=cohen_kappa_score(truth,pred)
            p,r,f,_=precision_recall_fscore_support(truth,pred,average='binary',zero_division=0)
            for key,v in {'kappa':k,'precision':p,'recall':r,'F1':f}.items():
                assert math.isclose(v,result['metrics'][key],abs_tol=1e-12)
            tn,fp,fn,tp=confusion_matrix(truth,pred,labels=[0,1]).ravel()
            assert result['four_cells']=={'TP':int(tp),'FP':int(fp),'FN':int(fn),'TN':int(tn)}
            crosschecks.append({'method':name,'dimension':dim,'four_cells':result['four_cells'],'all_core_metrics_agreed':True})
    assert len(differences['native'])==1 and len(differences['rule'])==4
    metrics={"trusted_human_message_verbatim":"已经核对都是1",
             "scope":"77 development records / 385 labels; saved human4 posthoc reference v2 vs unchanged existing machine firsts",
             "reference_identity":edits['reference_identity'],"reference_bindings":authored['outputs'],
             "reference_receipt_binding":bind(SUPPORT/'edit-inputs.json'),
             "prediction_bindings":{"native":old_native['predictions_binding'],'rule':old_rule['prediction_binding']},
             "native_batches_verified":batches,"dimensions":dimensions,'binary_disagreements':differences,
             "sklearn_crosschecks":crosschecks,"all_core_point_targets_pass":all(all(x['point_targets'].values()) for ys in dimensions.values() for x in ys.values()),
             "same_new_common_material_rerun":False,"new_human_review_context_bound":True,
             "machine_firsts_modified":False,"machine_coding_executed_again":False,
             "independent_human_initial_coding":False,"formal_acceptance":False,
             "population_CI95_calculated":False,"p_values_calculated":False}
    CALC.mkdir(exist_ok=False)
    with (CALC/'metrics.json').open('x') as f:f.write(json.dumps(metrics,ensure_ascii=False,indent=2)+'\n')
    with (SUPPORT/'saved-workbook-verification.json').open('x') as f:f.write(json.dumps({'tables':comparisons,'A_B_labels_equal':True,'all_labels_binary':True},ensure_ascii=False,indent=2)+'\n')
    prior_inputs=ROOT/'logs/uniform-machine77-report-delivery/report-inputs.json'
    report=read(prior_inputs)
    report['native_dimensions']=dimensions['native'];report['rule_dimensions']=dimensions['rule']
    report['native_disagreements']=differences['native'];report['rule_disagreements']=differences['rule']
    ref_book=openpyxl.load_workbook(edits['jobs'][0]['target'],read_only=True,data_only=True)
    latest={(x[0],x[2].split()[0]):x for x in ref_book['编码表'].iter_rows(min_row=7,max_row=391,values_only=True)}
    raw_maps={name:{x['id']:x['labels'] for x in rows} for name,rows in [('native',native),('rule',rule)]}
    for row in report['rows']:
        dim=row[2].split()[0];source=latest[(row[0],dim)]
        assert row[4]==raw_maps['native'][row[0]][dim] and row[5]==raw_maps['rule'][row[0]][dim]
        row[3]=source[3];row[6]=source[4]
    report['bindings']=[bind(CALC/'metrics.json'),bind(prior_inputs),*authored['outputs'],old_native['predictions_binding'],old_rule['prediction_binding']]
    REPORT.mkdir(exist_ok=False)
    with (REPORT/'report-inputs.json').open('x') as f:f.write(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    with (CALC/'seal.json').open('x') as f:f.write(json.dumps({'files':[bind(CALC/'metrics.json')],'formula_binding':bind(formula_path),'script_binding':bind(__file__)},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'saved_labels_each_table':385,'4_slots_updated_each':True,'native_errors':differences['native'],'rule_errors':differences['rule'],
                      'all_core_targets_pass':metrics['all_core_point_targets_pass'],'kappa':{k:{d:v['metrics']['kappa'] for d,v in ds.items()} for k,ds in dimensions.items()}},ensure_ascii=False))


if __name__=='__main__':main()
