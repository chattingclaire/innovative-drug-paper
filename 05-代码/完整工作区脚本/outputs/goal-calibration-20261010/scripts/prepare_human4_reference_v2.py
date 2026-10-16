"""Prepare trusted targeted human labels and source-grounded AI-drafted reasons.

Reads workbooks only; authors no XLSX and modifies no raw machine judgment.
"""
import hashlib
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / "logs/human4-reference-v2-delivery"
OUT = ROOT / "09-delivery/agent-review-20261010/真人4项复核后-v2"


def bind(p):
    p = Path(p).resolve()
    return {"path": str(p), "bytes": p.stat().st_size,
            "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}


def main():
    SUPPORT.mkdir(exist_ok=False)
    OUT.mkdir(exist_ok=False)
    packet = ROOT / "logs/remaining4-calibration-delivery/report-inputs.json"
    materials = json.loads(packet.read_text())
    for b in materials["bindings"]:
        assert bind(b["path"])["sha256"] == b["sha256"]
    reasons = [
        "当前以具有足够证据的明显临床优势作为申请适用条件，承接同版完整定义中的显著安全性优势分支，构成操作性安全标准，S=1。疾病严重程度本身不作为安全性依据。",
        "当前流程明确采用价格协商；同版PDF物理第8—9页说明商保专家提出协商价格评估意见，并按评估确定的价格开展协商，构成实质定价安排，Ec=1。商保目录推荐参考不等于已纳入基本医保支付或已实际承保。",
        "当前对已纳入品种实施技术审评沟通；同版程序明确创新药、改良型新药等适用对象，并规定审评中沟通与补充技术资料的具体安排，形成创新对象与可用技术支持的直接联系，I=1。依据纳入后安排，不以创新药名称或申报前沟通单独支持。",
        "当前纳入后沟通与同版审评中按具体情况优先安排沟通属于同一程序安排；适用范围及条件明确短缺、罕见病和儿童用药等需求缺口对象，形成按需要配置优先沟通机会的联系，F=1。不将所有优先程序品种一概视为有同一种缺口。",
    ]
    definitions = [(7,"RECODE25-00005","S"),(179,"RECODE25-03876","Ec"),
                   (215,"RECODE25-00138","I"),(216,"RECODE25-00138","F")]
    appendix_start = 851
    spans = {name: [a + appendix_start - 7, b + appendix_start - 7]
             for name, (a,b) in materials["material_spans"].items()}
    ids = {"00005": "RECODE25-00005", "00138": "RECODE25-00138", "03876": "RECODE25-03876"}
    appendix = []
    for block, title, pointer, text, path in materials["materials"]:
        sid = next((v for k,v in ids.items() if block.startswith(k)), "共同手册")
        appendix.append([sid, block, "本轮真人补核所用完整材料", pointer,
                         text.replace("\f", "␌"), title + "\n" + path])
    assert len(appendix) == 43
    contexts = [
        f"原文语境E851—E854；共同语境E878—E880；安全性E881",
        f"原文语境E865—E877；流程原图见此前剩余4项单独校准表；经济性E882—E883",
        f"原文语境E855—E864；共同语境E878—E880；创新性E884—E888",
        f"原文语境E855—E864；共同语境E878—E880；公平性E889—E893",
    ]
    jobs = []
    old_labels = []
    for side in "AB":
        p = ROOT / f"09-delivery/agent-review-20261010/编码复核表-{side}.xlsx"
        w = openpyxl.load_workbook(p, read_only=True, data_only=False)
        rows = list(w["编码表"].iter_rows(min_row=7,max_row=391,values_only=True))
        old_labels.append({(x[0],x[2].split()[0]):x[3] for x in rows})
        changes = []
        for (row,sid,dim), reason, context in zip(definitions,reasons,contexts):
            assert w["编码表"].cell(row,1).value == sid
            assert w["编码表"].cell(row,3).value.split()[0] == dim
            assert w["编码表"].cell(row,4).value == 0
            changes.append({"row":row,"source_id":sid,"dimension":dim,
                            "old_label":0,"value":1,"reason":reason,"context":context})
        jobs.append({"source":bind(p),"target":str(OUT / p.name),"changes":changes})
    assert old_labels[0] == old_labels[1] and len(old_labels[0]) == 385
    data = {"created_at":datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
            "trusted_human_message_verbatim":"已经核对都是1",
            "message_interpretation":"Four slots in previously supplied remaining-four calibration table; not all 385 labels.",
            "human_targeted_labels_received_N":4,"reference_identity":"AI-seeded A/B plus targeted human posthoc corrections",
            "reason_origin":"AI-drafted from supplied same-version evidence; user provided label values only, not these reason texts",
            "independent_human_initial_coding":False,"formal_acceptance":False,
            "original_files_modified":False,"raw_machine_firsts_modified":False,
            "appendix_start_row":appendix_start,"appendix_rows":appendix,
            "display_formfeed":"Only U+000C displayed reversibly as U+240C; no source file modified",
            "material_spans":spans,"source_packet_binding":bind(packet),
            "source_bindings":materials["bindings"],"jobs":jobs}
    with (SUPPORT / "edit-inputs.json").open("x") as f:
        f.write(json.dumps(data,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"jobs":len(jobs),"labels_received":4,"appendix_rows":43,"output":str(OUT)},ensure_ascii=False))


if __name__ == "__main__":
    main()
