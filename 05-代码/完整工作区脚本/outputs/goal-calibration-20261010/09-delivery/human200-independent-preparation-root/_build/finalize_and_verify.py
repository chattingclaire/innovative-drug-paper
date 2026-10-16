"""Finalize native material links, verify saved blank forms, and package by coder.

Workbook cells/styles are authored by Artifact Tool. ZIP/XML here supplies its
missing native local-hyperlink feature; openpyxl is used for readback only.
"""
import hashlib
import html
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parents[1]
MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('',MAIN)
ET.register_namespace('r',REL)


def binding(p):
    raw=p.read_bytes()
    return {'path':str(p),'bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest()}


def json_write(p,d):
    with p.open('x') as f:json.dump(d,f,ensure_ascii=False,indent=2);f.write('\n')


def check(condition,message):
    if not condition:raise ValueError(message)


def main():
    check(not (OUT/'seal.json').exists(),'Do not modify a sealed human preparation')
    source=json.loads((OUT/'_build/workbook-input.json').read_text())
    draw=json.loads((OUT/'coordinator-private/draw-and-mapping.json').read_text())
    pages=json.loads((OUT/'material-pages-index.json').read_text())['pages']
    checks=[]
    workbook_manifest=[]
    for side in ['A','B']:
        records=source['sides'][side]
        target=OUT/('coder-'+side)/('独立编码表-'+side+'-200条.xlsx')
        original=binding(target)
        with zipfile.ZipFile(target) as z:
            items={name:z.read(name) for name in z.namelist()}
        sheet_path='xl/worksheets/sheet1.xml'
        rel_path='xl/worksheets/_rels/sheet1.xml.rels'
        sheet=ET.fromstring(items[sheet_path])
        check(sheet.find('{'+MAIN+'}hyperlinks') is None,'Unexpected pre-existing hyperlinks')
        relationships=ET.fromstring(items[rel_path]) if rel_path in items else ET.Element('{'+PKG+'}Relationships')
        existing={r.attrib['Id'] for r in relationships}
        hyperlinks=ET.Element('{'+MAIN+'}hyperlinks')
        for n,record in enumerate(records):
            rid='rIdMaterial'+str(n+1)
            check(rid not in existing,'Relationship ID conflict')
            ET.SubElement(relationships,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink',
                'Target':'materials/'+record['blind_id']+'.html','TargetMode':'External'})
            for dim in range(5):
                ET.SubElement(hyperlinks,'{'+MAIN+'}hyperlink',{'ref':'A'+str(2+n*5+dim),'{'+REL+'}id':rid})
        following={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
        at=next((i for i,e in enumerate(sheet) if e.tag.rsplit('}',1)[-1] in following),len(sheet))
        sheet.insert(at,hyperlinks)
        validations=sheet.find('{'+MAIN+'}dataValidations')
        check(validations is not None and len(validations)==1,'Expected one coding-value validation')
        validation=validations[0]
        check(validation.attrib['sqref']=='D2:D1001' and validation.attrib['type']=='list','Unexpected value-input range')
        validation.set('allowBlank','1');validation.set('showErrorMessage','1');validation.set('errorStyle','stop')
        validation.set('errorTitle','仅填0或1')
        validation.set('error','值仅能为0或1；证据不足请留空，理由以“待定：”开头。')
        items[sheet_path]=ET.tostring(sheet,encoding='utf-8',xml_declaration=True)
        items[rel_path]=ET.tostring(relationships,encoding='utf-8',xml_declaration=True)
        temp=target.with_suffix('.native-links.tmp')
        with zipfile.ZipFile(temp,'x',compression=zipfile.ZIP_DEFLATED) as z:
            for name,raw in items.items():z.writestr(name,raw)
        temp.replace(target)
        wb=load_workbook(target,data_only=False,read_only=False)
        check(wb.sheetnames==['独立初判'],'Unexpected worksheets or hidden answer sheet')
        ws=wb['独立初判']
        check(ws.max_row==1001 and ws.max_column==5,'Unexpected row/column count')
        check(ws.freeze_panes=='B2','Expected frozen header row and identifier column')
        check(len(ws.data_validations.dataValidation)==1,'Missing coding input validation')
        check(ws.data_validations.dataValidation[0].allowBlank is True,'Lawful NULL must allow blank D')
        for n,record in enumerate(records):
            for d,dim in enumerate(source['dimensions']):
                row=2+n*5+d
                check([ws.cell(row,c).value for c in range(1,6)]==[record['blind_id'],record['body'],dim,None,None],'Saved source or blank fields changed')
                link=ws.cell(row,1).hyperlink
                check(link is not None and link.target=='materials/'+record['blind_id']+'.html','Missing or mismatched source link')
                check((target.parent/link.target).is_file(),'Native link destination missing')
        wb.close()
        for p in [p for p in pages if p['side']==side]:
            raw=Path(p['json']['path']).read_bytes()
            check(hashlib.sha256(raw).hexdigest()==p['source_bundle_sha256'],'Source JSON mutated')
            text=Path(p['html']['path']).read_text()
            for url in re.findall(r'(?:href|src)="([^"]+)"',text):
                url=html.unescape(url)
                parsed=urlsplit(url)
                if parsed.scheme or parsed.netloc or not parsed.path:continue
                dest=(Path(p['html']['path']).parent/unquote(parsed.path)).resolve()
                check(dest.is_relative_to(OUT),'Material link escapes its delivery root')
                check(dest.is_file(),'A selected material-page local link is missing')
        final=binding(target)
        workbook_manifest.append({'side':side,'file':final,'artifact_export_original':original,
            'record_N':200,'data_rows_N':1000,'header_row':1,'first_data_row':2,'last_data_row':1001,
            'worksheet':'独立初判','columns_N':5,'dimension_order':source['dimensions'],
            'native_material_hyperlinks_N':1000,'source_values_preserved':True,'D_E_blank':True,
            'human_results_received':0,'manual':binding(target.parent/'manual.txt'),
            'execution_prompt':binding(target.parent/'execution-prompt.txt'),
            'human_workflow_protocol':binding(target.parent/'common-execution-protocol.txt')})
        checks.append({'side':side,'all1000_saved_rows_verified':True,'all1000_native_links_verified':True,
            'all200_coder_material_json_bindings_verified':True,'selected_material_page_links_resolve':True,
            'D_E_blank':True,'five_columns_no_review_column':True,'visual_review_top_and_longest_completed_by_ROOT':True,
            'native_excel_application_open_not_tested':True,'human_semantic_judgments_not_created':True})
    ledger=json.loads((OUT/'_build/copy-ledger.json').read_text())
    for entry in ledger['copies']:
        check(binding(Path(entry['source']['path']))==entry['source'],'Original copied input changed')
        check(binding(Path(entry['destination']['path']))==entry['destination'],'A copied byte-bound source changed')
    entry=draw['pre_draw_freeze']
    check(binding(Path(entry['path']))==entry,'Pre-draw freeze mutated')
    readback={'status':'blank200_preparation_content_verified','at_utc':datetime.now(timezone.utc).isoformat(),
        'checks':checks,'source_copy_bindings_checked_N':len(ledger['copies']),
        'actual_human_submissions_N':0,'formal_validation_complete':False,
        'qualification_note':'Mechanical content checks do not authenticate two human executions or semantic correctness.',
        'source_asset_limit':'Shared original HTML may retain external navigation/assets not collected; original source bytes and known unresolved locators are preserved. Resolving selected material-page links is not a claim that every original webpage asset was captured.'}
    json_write(OUT/'mechanical-readback.json',readback)
    json_write(OUT/'workbook-manifest.json',{'status':'prepared200_not_distributed_no_human_submissions','workbooks':workbook_manifest,
        'budget_is_assumed_not_user_confirmed':True,'project_thresholds_changed':False,
        'draw':binding(OUT/'coordinator-private/draw-and-mapping.json'),'mechanical_readback':binding(OUT/'mechanical-readback.json')})
    json_write(OUT/'coordinator-private/reception-contract.json',{'expected_records_per_coder':200,'expected_judgment_rows_per_coder':1000,
        'worksheet':'独立初判','protected_columns':['A','B','C'],'editable_columns':['D','E'],
        'binary_input':'0/1 with a nonempty reason; not bool, nonfinite, other values, or conflicting pending reason',
        'lawful_null_input':'D blank and E starts待定：or待定: with a specific nonempty reason',
        'ordinary_blank_or_bare_pending_prefix':'incomplete, never silently converted to0/1',
        'two_nulls_count_as_binary_agreement':False,'original_files_and_initial_judgments_preserved':True,
        'actual_human_identity_and_independence_not_inferred_from_file_content':True,
        'new200_weights':'use the frozen new n_h/N_h and N_h/n_h; never use the800 weights',
        'remaining_domain_N':120,'exposed_domain_N':80,'private_mapping':binding(OUT/'coordinator-private/draw-and-mapping.json')})
    package_records=[]
    for side in ['A','B']:
        filename=OUT/('真人独立判读-'+side+'-200条.zip')
        members=[]
        coder=OUT/('coder-'+side)
        for p in coder.rglob('*'):
            if p.is_file() and not p.name.endswith('.inspect.ndjson'):members.append(p)
        members.extend(p for p in (OUT/'shared-source').rglob('*') if p.is_file())
        with zipfile.ZipFile(filename,'x',compression=zipfile.ZIP_DEFLATED) as z:
            for p in sorted(members):z.write(p,p.relative_to(OUT))
        with zipfile.ZipFile(filename) as z:
            check(z.testzip() is None,'ZIP CRC failure')
            names=z.namelist()
            check(len(names)==len(members),'ZIP membership mismatch')
            check(all(n.startswith('coder-'+side+'/') or n.startswith('shared-source/') for n in names),'Cross-coder or private content leaked')
            for p in members:check(z.read(str(p.relative_to(OUT)))==p.read_bytes(),'ZIP bytes differ from prepared material')
        package_records.append({'side':side,'file':binding(filename),'member_N':len(members),'CRC_and_each_member_bytes_verified':True,'distributed':False})
    json_write(OUT/'distribution-package-manifest.json',{'packages':package_records,'actual_distribution_confirmed':False})
    json_write(OUT/'seal.json',{'artifact_type':'human200_blank_preparation_seal','at_utc':datetime.now(timezone.utc).isoformat(),
        'status':'prepared_for_review_not_a_completed_human_validation','members':[binding(p) for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='seal.json'],
        'original800_files_overwritten':False,'sample_size_is_a_method_minimum':False,'actual_human_results':0})
    print(json.dumps({'status':'two200_forms_and_separate_packages_verified','workbooks_N':2,'data_rows_per_workbook':1000,
        'native_links_per_workbook':1000,'human_results_received':0,'formal_validation_complete':False},ensure_ascii=False))


if __name__=='__main__':main()
