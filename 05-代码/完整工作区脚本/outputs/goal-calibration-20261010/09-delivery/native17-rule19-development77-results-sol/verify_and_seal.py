from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import math
import shutil
import openpyxl

OUT = Path('/Users/rainbow/Desktop/战略院/药企项目-创新药/outputs/goal-calibration-20261010/09-delivery/native17-rule19-development77-results-sol')
ROOT = OUT.parent.parent
final = OUT / '当前77条机器编码与指标.xlsx'
if not final.exists():
    shutil.move(OUT / 'versions/v4/当前77条机器编码与指标.xlsx', final)
shutil.copyfile(OUT / 'build_workbook.mjs', OUT / 'versions/v4/build_workbook.mjs')

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def load(p): return json.loads(Path(p).read_text())
data = load(OUT / 'inputs/mechanical-workbook-input.json')
metrics = load(ROOT / '07-statistics/native-v17-rule-v19-human-confirmed77-root/metrics.json')
source_checks = []
for entry in data['input_manifest']:
    old, copy = Path(entry['original_path']), OUT / entry['archive_relative_path']
    assert sha(old) == sha(copy) == entry['sha256']
    source_checks.append({'path': str(old), 'sha256': sha(old), 'unchanged': True, 'copy_exact': True})

w = openpyxl.load_workbook(final, data_only=True, read_only=False)
assert w.sheetnames == ['说明', '逐维指标', '编码对照', '两处差异']
panes = {'说明': 'A7', '逐维指标': 'C7', '编码对照': 'D7', '两处差异': 'A7'}
assert all(w[k].freeze_panes == v and not w[k].sheet_view.showGridLines for k, v in panes.items())
actual_codes = [list(r) for r in w['编码对照'].iter_rows(min_row=7, max_row=391, max_col=12, values_only=True)]
assert actual_codes == data['code_rows']
keys = [(r[0], r[2].split()[0]) for r in actual_codes]
assert len(keys) == len(set(keys)) == 385
assert all(v in (0, 1) for row in actual_codes for v in row[3:7])

derived = [json.loads(x) for x in (ROOT / '07-statistics/native-v17-rule-v19-human-confirmed77-root/derived-native-predictions.jsonl').read_text().splitlines() if x]
llm_map = {r['id']: r for r in derived}
rule_path = Path(metrics['rule_predictions']['path'])
rule_rows = [json.loads(x) for x in rule_path.read_text().splitlines() if x]
rule_map = {r['id']: (i + 1, r) for i, r in enumerate(rule_rows)}
human_maps = []
for binding in metrics['reference_bindings']:
    a = openpyxl.load_workbook(binding['path'], read_only=True, data_only=True)
    amap = {}
    for i, vals in enumerate(a['编码表'].iter_rows(values_only=True), 1):
        if vals and isinstance(vals[0], str) and vals[0].startswith('RECODE25-'):
            amap[vals[0], vals[2].split()[0]] = (i, vals)
    human_maps.append(amap)
    a.close()

native_cache = {}
for row in actual_codes:
    rid, dim = row[0], row[2].split()[0]
    d = llm_map[rid]
    original = d['original_whole_batch_submission']['path']
    if original not in native_cache: native_cache[original] = load(original)
    rec_index = int(d['record_pointer'].split('/')[-1])
    record = native_cache[original]['records'][rec_index]
    assert record['source_id'] == rid
    assert row[5] == d['labels'][dim] == record['dimensions'][dim]['value']
    assert row[7] == record['dimensions'][dim]['reason']
    line, rule = rule_map[rid]
    assert row[6] == rule['labels'][dim]
    assert row[8] == str(Path(original).relative_to(ROOT)) + '#' + d['record_pointer'] + '/dimensions/' + dim
    assert row[9] == str(rule_path.relative_to(ROOT)) + ':第' + str(line) + '行 /decisions/' + dim
    for i in (0, 1):
        source_row, vals = human_maps[i][rid, dim]
        assert row[3 + i] == vals[3] and row[1] == vals[1]
        assert row[10 + i] == str(Path(metrics['reference_bindings'][i]['path']).relative_to(ROOT)) + '#编码表!D' + str(source_row) + ',E' + str(source_row)

actual_metrics = [list(r) for r in w['逐维指标'].iter_rows(min_row=7, max_row=21, max_col=30, values_only=True)]
assert actual_metrics == data['metric_rows']
metric_keys = ['kappa', 'precision', 'recall', 'F1', 'Po', 'Pe', 'Ppos', 'Pneg', 'specificity', 'NPV']
count_keys = ['assigned_N', 'saved_N', 'joint_binary_N', 'reference_positive_N', 'reference_negative_N', 'prediction_positive_N', 'prediction_negative_N']
for row in actual_metrics:
    pointer = row[29].split('/')
    expected = metrics['comparisons'][pointer[2]]['dimensions'][pointer[4]]
    assert row[2:9] == [expected[k] for k in count_keys]
    assert row[9:13] == [expected['four_cells'][k] for k in ['TP', 'FP', 'FN', 'TN']]
    assert row[13:23] == [expected['metrics'][k] if expected['metrics'][k] is not None else 'NA' for k in metric_keys]
    assert row[23:27] == [('达到' if expected['point_targets'][k] else '未达到') if expected['point_targets'][k] is not None else 'NA' for k in ['kappa', 'precision', 'recall', 'F1']]
    assert row[27:29] == ['NA', 'NA']

found = [(r[0], r[2].split()[0], r[3], r[4], r[5], r[6]) for r in actual_codes if r[5] != r[3]]
assert set(found) == {('RECODE25-04457', 'I', 0, 0, 1, 0), ('RECODE25-00131', 'F', 0, 0, 1, 0)}
for idx, d in enumerate(data['differences']):
    start = 7 + 18 * idx
    values = [w['两处差异'].cell(r, 2).value for r in range(start, start + 16)]
    expected = [d['id'], d['department'], d['dimension'], d['reference_A'], d['reference_B'], d['LLM'], d['rule'], d['relative_direction'], d['LLM_reason'], d['rule_reason'], d['A_reason'], d['B_reason'], d['LLM_source'], d['rule_source'], d['A_source'], d['B_source']]
    assert values == expected
    dim = d['dimension_code']
    assert d['rule_reason'] == rule_map[d['id']][1]['decisions'][dim]['reason']
    assert d['A_reason'] == human_maps[0][d['id'], dim][1][4]
    assert d['B_reason'] == human_maps[1][d['id'], dim][1][4]

verification = {'created_at_utc': datetime.now(timezone.utc).isoformat(), 'final_file': str(final), 'final_sha256': sha(final), 'sheets': w.sheetnames, 'code_rows_N': 385, 'label_columns_each_N': 385, 'all_4_label_columns_match_originals': True, 'LLM_reasons_N': 385, 'all_reasons_exact_match': True, 'metric_rows_N': 15, 'metric_values_N': 150, 'all_metric_values_and_counts_match_metrics_JSON_exactly': True, 'all_source_pointers_mechanically_verified': True, 'differences': found, 'difference_reasons_exact_match_originals': True, 'frozen_panes': panes, 'inputs': source_checks, 'original_A_B_unchanged': True, 'semantic_labels_changed': False, 'p_values_calculated': False, 'CI_calculated': False, 'new_holdout_claimed': False, 'full4481_claimed': False, 'genuine_repeats_claimed': False, 'visual_review': 'four sheets reviewed through eight rendered views; no clipped labels or values observed', 'failure_history': ['v1: log redirection failed before builder started', 'v2: memory readback caught ISO timestamp normalization', 'v3: diagnostic rerun retained the normalization mismatch', 'v4: added source-text prefix to timestamp; all readback checks passed']}
(OUT / 'verification.json').write_text(json.dumps(verification, ensure_ascii=False, indent=2))
(OUT / '交付与核验报告.txt').write_text('当前77条机器编码与指标\n\n交付：当前77条机器编码与指标.xlsx。四页为说明、逐维指标、编码对照、两处差异。\n范围：77条已曝光开发材料，385个二元标签。13批、15次实际尝试、2次技术无效尝试、101条记录尝试。技术重试不是独立重复。\n本表未计算p值或95%区间，未开展真正重复；不声称新holdout验证、全量4481正确或正式验收。\n\n机械核验：\n1. Artifact内存及导出重新导入逐值读回通过。\n2. 独立Python只读核对385行四列标签、385条LLM理由、15行全部指标和计数，与原JSON及原A/B文件完全一致。\n3. 两处差异仅为04457 I和00131 F：A/B=0、LLM=1、规则=0。双方机器及A/B理由均从原件复制，未裁定语义正确性。\n4. 来源指针逐行机械验证。原A/B及所有输入SHA-256未变，输入副本与原件字节一致。\n5. 四页冻结表头，指标页同时冻结两列，编码页冻结三列。八幅预览已目视检查。\n\n真实失败保留在versions：v1日志目录错误；v2/v3读回发现裸ISO时间字符串被工具自动转换。v4使用来源文本前缀保留完整时间文本后通过；未改标签或指标。\n\n最终工作簿SHA-256：' + sha(final) + '\n')

excluded = {'seal.json'}
files=[]
for p in sorted(OUT.rglob('*')):
    if not p.is_file() or 'node_modules' in p.parts or p.name in excluded: continue
    files.append({'relative_path': str(p.relative_to(OUT)), 'bytes': p.stat().st_size, 'sha256': sha(p)})
seal = {'created_at_utc': datetime.now(timezone.utc).isoformat(), 'final_workbook': {'path': str(final), 'bytes': final.stat().st_size, 'sha256': sha(final)}, 'verified_inputs': len(source_checks), 'no_original_modified': True, 'files': files}
(OUT / 'seal.json').write_text(json.dumps(seal, ensure_ascii=False, indent=2))
w.close()
print(json.dumps({'final_file': str(final), 'sha256': sha(final), 'code_rows': 385, 'metric_rows': 15, 'verified_source_pointers': True, 'unchanged_inputs': len(source_checks), 'sealed_files': len(files)}, ensure_ascii=False))
