"""Read-only independent verification of Artifact Tool exports. Never saves a workbook."""
import collections
import hashlib
import html.parser
import json
import pathlib
import sys
import urllib.parse

import openpyxl


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def colour(value):
    if value is None:
        return None
    return (value.type, value.value, value.tint)


def style(cell):
    font, alignment, fill, border = cell.font, cell.alignment, cell.fill, cell.border
    sides = []
    for name in ("left", "right", "top", "bottom", "diagonal", "vertical", "horizontal"):
        side = getattr(border, name)
        sides.append(None if side is None or not side.style else (name, side.style, colour(side.color)))
    return {
        "font": (font.name, font.sz, bool(font.b), bool(font.i), font.u, colour(font.color)),
        "alignment": (alignment.horizontal, alignment.vertical, bool(alignment.wrap_text), bool(alignment.shrink_to_fit), alignment.indent or 0, alignment.textRotation or 0),
        "fill": (fill.patternType, colour(fill.fgColor)),
        "border": sides,
        "number_format": cell.number_format,
        "protection": (bool(cell.protection.locked), bool(cell.protection.hidden)),
    }


def validation(ws):
    return sorted((v.type, str(v.sqref), v.formula1, v.formula2, bool(v.allowBlank), bool(v.showErrorMessage), v.errorStyle, v.errorTitle, v.error) for v in ws.data_validations.dataValidation)


class LocalLinkParser(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.targets = []

    def handle_starttag(self, tag, attrs):
        for attribute, value in attrs:
            if attribute.lower() in ("href", "src") and value is not None:
                self.targets.append(value)


def verify_local_navigation(output_dir, include_entrance=False):
    output_dir = output_dir.resolve()
    material_pages = sorted((output_dir / "材料").glob("*.html"))
    require(len(material_pages) == 200, "Expected 200 material HTML pages")
    pages = material_pages + [output_dir / "材料目录.html"]
    if include_entrance:
        pages.append(output_dir / "AI复核阅读入口.html")
    for required_name in ("共同编码手册.txt", "AI复核说明.txt", "AI交付说明.txt", "填写说明.txt"):
        require((output_dir / required_name).is_file(), f"Missing root navigation dependency: {required_name}")
    checked = 0
    skipped_external = 0
    for page in pages:
        require(page.is_file(), f"Missing navigation page: {page}")
        parser = LocalLinkParser()
        parser.feed(page.read_text(encoding="utf8"))
        for target in parser.targets:
            parsed = urllib.parse.urlsplit(target.strip())
            if parsed.scheme or parsed.netloc:
                skipped_external += 1
                continue
            if not parsed.path:  # same-page fragment or query
                continue
            destination = (page.parent / urllib.parse.unquote(parsed.path)).resolve()
            require(destination.is_relative_to(output_dir), f"Local link leaves the result package: {page.name} -> {target}")
            require(destination.exists(), f"Broken local href/src: {page.relative_to(output_dir)} -> {target}")
            checked += 1
    return {"html_pages_verified": len(pages), "local_href_src_verified": checked, "external_href_src_skipped": skipped_external, "broken_local_links": 0}


if sys.argv[1] == "--links-only":
    print(json.dumps(verify_local_navigation(pathlib.Path(sys.argv[2]), include_entrance=True), ensure_ascii=False))
    sys.exit(0)


if sys.argv[1] == "--source-layout":
    source_path = pathlib.Path(sys.argv[2])
    source = openpyxl.load_workbook(source_path, data_only=False)
    ws = source["独立初判"]
    print(json.dumps({
        "widths": {col: ws.column_dimensions[col].width for col in "ABCDE"},
        "rows": {
            row: {
                "height": ws.row_dimensions[row].height or ws.sheet_format.defaultRowHeight or 15,
                "dFontPoints": ws.cell(row, 4).font.sz or 11,
                "eFontPoints": ws.cell(row, 5).font.sz or 11,
            }
            for row in range(1, 1002)
        },
    }, ensure_ascii=False))
    sys.exit(0)

source_path, result_path, prediction_path, output_dir, layout_path = map(pathlib.Path, sys.argv[1:])
source = openpyxl.load_workbook(source_path, data_only=False)
result = openpyxl.load_workbook(result_path, data_only=False)
predictions = {p["blind_id"]: p["judgments"] for p in json.loads(prediction_path.read_text())}
approved_layout = json.loads(layout_path.read_text())
require(source.sheetnames == result.sheetnames == ["独立初判"], "Unexpected worksheet change")
original, generated = source["独立初判"], result["独立初判"]
require((original.max_row, original.max_column) == (1001, 5), "Unexpected source extent")
require((generated.max_row, generated.max_column) == (1001, 5), "Unexpected output extent")
require(generated["D1"].value == "AI判定（0/1）\n待定留空", "AI column heading missing")
require(generated["E1"].value == "AI理由与原文依据", "AI reason heading missing")
require(original.freeze_panes == generated.freeze_panes, "Freeze pane changed")
require(str(original.merged_cells) == str(generated.merged_cells), "Merged cells changed")
require(original.sheet_view.showGridLines == generated.sheet_view.showGridLines, "Grid visibility changed")
require(validation(original) == validation(generated), "Data validation changed")
require(len(original.conditional_formatting) == len(generated.conditional_formatting), "Conditional formatting changed")
require(len(original._charts) == len(generated._charts) == 0, "Unexpected chart content")
require(len(original._images) == len(generated._images) == 0, "Unexpected image content")

for column in "ABCDE":
    a, b = original.column_dimensions[column], generated.column_dimensions[column]
    require(not a.hidden and not b.hidden, f"Hidden column {column}")
    if column in "ABC":
        require(a.width == b.width, f"Source column {column} width changed")
    else:
        expected_width = approved_layout["widths"][column]
        require(a.width <= expected_width <= max(a.width, 24 if column == "D" else 96), f"Unapproved width range for column {column}")
        require(abs(b.width - expected_width) < 0.01, f"Column {column} differs from recorded readability adjustment")
require(not any(dimension.hidden for dimension in generated.column_dimensions.values()), "Hidden output column")
require(not any(dimension.hidden for dimension in generated.row_dimensions.values()), "Hidden output row")

changed_styles = []
checked_cells = 0
hyperlinks = 0
identity_counts = collections.Counter()
for row in range(1, 1002):
    a, b = original.row_dimensions[row], generated.row_dimensions[row]
    original_height = a.height or original.sheet_format.defaultRowHeight or 15
    output_height = b.height or generated.sheet_format.defaultRowHeight or 15
    expected_height = approved_layout["rowHeights"][str(row)]
    require(not a.hidden and not b.hidden, f"Hidden row {row}")
    require(original_height <= expected_height <= 409, f"Invalid readability height for row {row}")
    require(abs(output_height - expected_height) < 0.01, f"Row {row} differs from recorded readability adjustment")
    for col in range(1, 6):
        old, new = original.cell(row, col), generated.cell(row, col)
        if col <= 3:
            require(old.value == new.value and old.data_type == new.data_type, f"Unexpected source change: {old.coordinate}")
        if style(old) != style(new):
            changed_styles.append(old.coordinate)
        if old.hyperlink is not None:
            require(new.hyperlink is not None and old.hyperlink.target == new.hyperlink.target, f"Hyperlink changed: {old.coordinate}")
            target = urllib.parse.unquote(new.hyperlink.target)
            require(not urllib.parse.urlparse(target).scheme, f"Expected local hyperlink: {old.coordinate}")
            require((output_dir / target).is_file(), f"Missing linked material: {target}")
            hyperlinks += 1
        else:
            require(new.hyperlink is None, f"Unexpected hyperlink: {old.coordinate}")
        require(new.data_type != "f", f"Unexpected formula: {new.coordinate}")
    if row > 1:
        blind_id, dimension = generated.cell(row, 1).value, generated.cell(row, 3).value
        require(blind_id in predictions and dimension in predictions[blind_id], f"Unexpected identity row {row}")
        expected = predictions[blind_id][dimension]
        require(generated.cell(row, 4).value == expected["value"], f"Judgment mismatch row {row}")
        require(generated.cell(row, 5).value == expected["reason"], f"Reason mismatch row {row}")
        for col in (4, 5):
            require(generated.cell(row, col).fill.fgColor.rgb == "FFFFF4D6", f"Input yellow changed row {row}")
        identity_counts[(blind_id, dimension)] += 1
        checked_cells += 2
require(not changed_styles, f"Unexpected cell style changes: {changed_styles[:15]} ({len(changed_styles)} total)")
require(len(predictions) == 200 and len(identity_counts) == 1000 and all(v == 1 for v in identity_counts.values()), "Record identity count mismatch")
require(hyperlinks == 1000, "Missing material hyperlinks")

copied_files = 0
for folder in ("材料", "附件"):
    for src in (source_path.parent / folder).rglob("*"):
        if src.is_file():
            dest = output_dir / src.relative_to(source_path.parent)
            require(dest.is_file(), f"Missing supporting file: {dest}")
            require(hashlib.sha256(src.read_bytes()).digest() == hashlib.sha256(dest.read_bytes()).digest(), f"Supporting file differs: {dest}")
            copied_files += 1
for filename in ("材料目录.html", "共同编码手册.txt"):
    src = source_path.parent / filename
    dest = output_dir / filename
    require(dest.is_file(), f"Missing copied navigation file: {filename}")
    require(hashlib.sha256(src.read_bytes()).digest() == hashlib.sha256(dest.read_bytes()).digest(), f"Copied navigation file differs: {filename}")
    copied_files += 1
navigation = verify_local_navigation(output_dir)

print(json.dumps({"records": 200, "dimension_rows": 1000, "verified_D_E_cells": checked_cells, "material_hyperlinks": hyperlinks, "supporting_files_verified": copied_files, "material_navigation": navigation, "unexpected_content_changes": 0, "unexpected_style_changes": 0, "approved_D_E_widths": approved_layout["widths"], "rows_expanded_for_readability": sum((generated.row_dimensions[r].height or generated.sheet_format.defaultRowHeight or 15) > (original.row_dimensions[r].height or original.sheet_format.defaultRowHeight or 15) for r in range(1, 1002)), "null_judgments": sum(j["value"] is None for p in predictions.values() for j in p.values()), "source_sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(), "output_sha256": hashlib.sha256(result_path.read_bytes()).hexdigest()}, ensure_ascii=False))
