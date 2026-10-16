"""Supply native local links missing from Artifact Tool and read back the export.

All cell content and formatting are authored in build.mjs. This changes only
worksheet hyperlink elements and hyperlink relationships. openpyxl reads only.
"""
import hashlib
import io
import json
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent
MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('', MAIN)
ET.register_namespace('r', REL)


def check(condition, message):
    if not condition:
        raise ValueError(message)


def binding(path):
    raw = path.read_bytes()
    return {'path': str(path), 'bytes': len(raw),
            'sha256': hashlib.sha256(raw).hexdigest()}


def main():
    receipt = OUT / 'saved-readback.json'
    check(not receipt.exists(), 'Already finalized; preserve this snapshot')
    source_path = OUT / 'workbook-input.json'
    source = json.loads(source_path.read_text())
    target = OUT / '创新性与公平性-全部阳性查看表.xlsx'
    before = binding(target)
    with zipfile.ZipFile(target) as z:
        check(z.testzip() is None, 'Export ZIP CRC failure')
        infos = {info.filename: info for info in z.infolist()}
        originals = {name: z.read(name) for name in z.namelist()}
    items = dict(originals)
    modified = set()
    for index, item in enumerate(source['sheets'], 1):
        sheet_path = f'xl/worksheets/sheet{index}.xml'
        rel_path = f'xl/worksheets/_rels/sheet{index}.xml.rels'
        sheet = ET.fromstring(items[sheet_path])
        check(sheet.find(f'{{{MAIN}}}hyperlinks') is None,
              'Unexpected pre-existing worksheet hyperlinks')
        relationships = (ET.fromstring(items[rel_path]) if rel_path in items
                         else ET.Element(f'{{{PKG}}}Relationships'))
        existing = {r.attrib['Id'] for r in relationships}
        hyperlinks = ET.Element(f'{{{MAIN}}}hyperlinks')
        for number, record in enumerate(item['rows'], 1):
            uri = record['link']
            parsed = urlsplit(uri)
            check(parsed.scheme == 'file' and not parsed.netloc,
                  'Expected an absolute local file URI')
            check(Path(unquote(parsed.path)).is_file(),
                  'Material destination does not exist')
            rid = 'rIdMaterial' + str(number)
            check(rid not in existing, 'Relationship ID collision')
            ET.SubElement(relationships, f'{{{PKG}}}Relationship', {
                'Id': rid, 'Type': REL + '/hyperlink',
                'Target': uri, 'TargetMode': 'External'})
            ET.SubElement(hyperlinks, f'{{{MAIN}}}hyperlink', {
                'ref': 'A' + str(number + 4), f'{{{REL}}}id': rid})
        following = {'printOptions', 'pageMargins', 'pageSetup', 'headerFooter',
                     'rowBreaks', 'colBreaks', 'customProperties', 'cellWatches',
                     'ignoredErrors', 'smartTags', 'drawing', 'legacyDrawing',
                     'legacyDrawingHF', 'picture', 'oleObjects', 'controls',
                     'webPublishItems', 'tableParts', 'extLst'}
        at = next((n for n, element in enumerate(sheet)
                   if element.tag.rsplit('}', 1)[-1] in following), len(sheet))
        sheet.insert(at, hyperlinks)
        items[sheet_path] = ET.tostring(sheet, encoding='utf-8',
                                      xml_declaration=True)
        items[rel_path] = ET.tostring(relationships, encoding='utf-8',
                                    xml_declaration=True)
        modified.update([sheet_path, rel_path])
    temp = target.with_suffix('.native-links.tmp')
    with zipfile.ZipFile(temp, 'x', compression=zipfile.ZIP_DEFLATED) as z:
        for name, raw in items.items():
            z.writestr(infos.get(name, name), raw)
    with zipfile.ZipFile(temp) as z:
        check(z.testzip() is None, 'Final ZIP CRC failure')
        check(set(z.namelist()) == set(items), 'Unexpected ZIP members')
        for name, raw in originals.items():
            if name not in modified:
                check(z.read(name) == raw, 'Unrelated workbook part changed')
    wb = load_workbook(io.BytesIO(temp.read_bytes()), read_only=False, data_only=False)
    check(wb.sheetnames == [s['name'] for s in source['sheets']],
          'Unexpected worksheet names')
    checks = []
    for index, item in enumerate(source['sheets']):
        ws = wb[item['name']]
        check(ws.max_row == 4 + len(item['rows']), 'Unexpected row count')
        check(ws.max_column == len(item['headers']), 'Unexpected column count')
        check([ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
              == item['headers'], 'Saved headers changed')
        check(ws.freeze_panes == ('E5' if index == 0 else 'D5'),
              'Expected frozen headings and identifiers')
        for number, record in enumerate(item['rows'], 5):
            cells = [ws.cell(number, c) for c in range(1, ws.max_column + 1)]
            check([cell.value for cell in cells] == record['values'],
                  'Saved row differs from its source matrix')
            for cell, value in zip(cells, record['values']):
                check(cell.data_type != 'f', 'Unexpected generated formula')
                if type(value) in (int, float):
                    check(type(cell.value) in (int, float),
                          'Numeric labels became text')
            check(cells[0].hyperlink is not None and
                  cells[0].hyperlink.target == record['link'],
                  'Saved native material link differs from source')
            check(Path(unquote(urlsplit(cells[0].hyperlink.target).path)).is_file(),
                  'Saved link destination missing')
        checks.append({'sheet': item['name'], 'data_rows': len(item['rows']),
                       'all_row_values_read_back_equal': True,
                       'numeric_label_types_preserved': True,
                       'native_material_links_verified': len(item['rows']),
                       'all_link_destinations_exist': True,
                       'freeze_panes': ws.freeze_panes})
    wb.close()
    temp.replace(target)
    report = {'at_utc': datetime.now(timezone.utc).isoformat(),
              'workbook': binding(target),
              'artifact_tool_export_before_hyperlink_feature': before,
              'source_matrix': binding(source_path),
              'checks': checks,
              'all_unrelated_zip_part_bytes_preserved': True,
              'ZIP_CRC_verified': True,
              'visual_review_opening_and_longest_rows_completed_previously': True,
              'hyperlink_patch_changes_displayed_values_or_styles': False,
              'native_excel_application_behaviour_not_tested': True,
              'original_coding_files_modified': False,
              'scope': '68 auxiliary800 AI I/F positive slots plus 102 full4481 rule I/F positive slots; overlapping sets, not an additive sample',
              'formal_human_reference_or_accuracy_claim': False}
    with receipt.open('x') as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
        f.write('\n')
    print(json.dumps({'status': 'saved_values_and_170_native_links_verified',
                      'data_rows_by_sheet': [68, 102],
                      'workbook': str(target)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
