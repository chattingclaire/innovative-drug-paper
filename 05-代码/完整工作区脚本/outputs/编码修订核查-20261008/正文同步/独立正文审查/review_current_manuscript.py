"""Independent source-bound document/CSV audit; never edits delivery inputs."""
from pathlib import Path
from fractions import Fraction as F
from itertools import combinations
import csv, datetime, hashlib, json, math, re, sys
sys.dont_write_bytecode = True
from docx import Document

ROOT = Path('/PROJECT_ROOT')
BASE = ROOT/'outputs/编码修订核查-20261008'
DEL = BASE/'本轮结果'
STAT = BASE/'统计复算/完整39与LLM03-同步复算01'
OUT = BASE/'正文同步/独立正文审查'
INITIAL = BASE/'正文同步/接入裁定与清除非同源复测前记录'
MAIN = INITIAL/'创新药评价的注意力配置差异与协调路径.docx'
NOTE = INITIAL/'编码复核说明与真人检验事项.docx'
OLD = ROOT/'outputs/final-three-department-20261008-human-validation-template/创新药评价的注意力配置差异与协调路径.docx'
FREEZE = STAT/'output_freeze.json'
RK='rules-v39'; LK='LLM-candidate03-single-unit-reconciled'
DS=('S','E','Ec','I','F')
DN=dict(zip(DS,('安全性','有效性','经济性','创新性','公平性')))
GN={'ALL':'总体','NMPA':'药监','NHSA':'医保','NHC':'卫健'}
MN={RK:'自动规则编码',LK:'LLM编码','A_C':'编码员A—编码员B'}
RN={'A':'编码员A','C':'编码员B','candidate':'复核参照（含AI复核者T判读）','symmetric A/C':'编码员A与编码员B的对称比较'}
CN={'probability_canonical':'概率样本','enriched_challenge':'挑战材料'}
WN={'unweighted':'等额样本描述','eligible_design_weighted':'部门抽样设计加权'}
MET={'kappa':'κ','agreement':'观察一致率','precision':'精确率P','recall':'召回率R','f1':'F1','left_positive_rate':'左侧编码1比例','right_positive_rate':'右侧编码1比例'}
EXPECTED={MAIN:'042cf1f7154fa8b2ad915b6cbdb6948e8583a8f164fccff0f5e168868d151705',NOTE:'ac7f161cd4ea2238c7a259df8314eb1932f14737b9498a751171d958ac76c046',OLD:'4d61c921a8ea9837f9d31eb4cd09ace0af1839b51e2793232ee0f402ede905e0',FREEZE:'f411935421447c6fd39fdf1887078dc6ba297e201ac5bed6498befe2ba753c3a'}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(p):return json.loads(Path(p).read_text())
def canonical(x):return hashlib.sha256(json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
for p,h in EXPECTED.items():assert sha(p)==h,(p,'provided_SHA_mismatch')
freeze=read(FREEZE)
for item in freeze['files']:assert sha(item['path'])==item['sha256'],item['path']
doc=Document(MAIN); note=Document(NOTE); old=Document(OLD)
classification=read(STAT/'machine_reference_classification.json')['comparisons']
ac=read(STAT/'A_C_before_third_agreement.json')['comparisons']
findings=read(STAT/'full4481/完整定义构成核查.json')
details=read(STAT/'full4481/all-pair-details/all60_conditional_department_contrasts.json')['rows']
allrecords=classification+ac
assert len(classification)==360 and len(ac)==60 and len(details)==60
look={(r['component'],r['group'],r['weighting'],r['method'],r['reference'],r['dimension']):r for r in allrecords}
assert len(look)==420
errors=[]; checks=[]
def check(tag,actual,expected,where):
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
        good=(isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12))
    else:good=actual==expected
    if not good:errors.append({'check':tag,'location':where,'actual':actual,'expected':expected})
    checks.append({'check':tag,'location':where,'pass':good})
    return good
def wrow(method,ref,dim):return look['probability_canonical','ALL','eligible_design_weighted',method,ref,dim]
def exact_metrics(row):
    key='symmetric_weighted_cells' if row['method']=='A_C' else 'weighted_confusion'
    a,b,c,d=[F(str(v)).limit_denominator(60) for v in row[key].values()]
    n=a+b+c+d
    ratio=lambda x,y:x/y if y else None
    po=ratio(a+d,n); lp=ratio(a+b,n); rp=ratio(a+c,n)
    pe=lp*rp+(1-lp)*(1-rp) if n else None
    return {'kappa':ratio(po-pe,1-pe) if n else None,'precision':ratio(a,a+b),'recall':ratio(a,a+c),'f1':ratio(2*a,2*a+b+c)}
table2=[]
for ri,row in enumerate(doc.tables[1].rows[1:],1):
    dim=next(d for d,n in DN.items() if n==row.cells[0].text)
    method=next(m for m,n in MN.items() if n==row.cells[1].text)
    for col,metric in enumerate(('kappa','precision','recall','f1'),2):
        actual=row.cells[col].text.split('/')
        assert len(actual)==2
        for side,ref in enumerate(('A','C')):
            source=wrow(method,ref,dim); v=source['metrics'][metric]['estimate']; expect='NA' if v is None else f'{v:.3f}'
            check('table2_display',actual[side],expect,f'table2.r{ri}.c{col}.ref{ref}')
            rational=exact_metrics(source)[metric]
            check('table2_independent_Fraction',v,float(rational) if rational is not None else None,f'{method}/{ref}/{dim}/{metric}')
            table2.append({'row':ri,'column':col,'display_reference':'A' if ref=='A' else 'B','dimension':dim,'method':method,'metric':metric,'actual':actual[side],'expected':expect,'source_key':['probability_canonical','ALL','eligible_design_weighted',method,ref,dim,'metrics',metric,'estimate'],'Fraction_from_weighted_cells':str(rational)})
assert len(table2)==80
table3=[]
for ri,row in enumerate(doc.tables[2].rows[1:],1):
    g=next(k for k,v in GN.items() if v==row.cells[0].text);m=next(k for k,v in MN.items() if v==row.cells[1].text)
    profile=findings['primary_full4481'][m]['departments'][g]['profile']
    totals=[int(profile['coding_count'][d]) for d in DS];den=sum(totals)
    check('five_dimension_total',profile['five_dimension_coding_count'],den,f'{m}/{g}')
    for ci,(d,n) in enumerate(zip(DS,totals),2):
        q=F(n,den);display=f'{float(q)*100:.2f}'
        check('table3_display',row.cells[ci].text,display,f'table3.r{ri}.c{ci}')
        check('table3_source_ratio',profile['attention_share'][d],float(q),f'{m}/{g}/{d}')
        table3.append({'row':ri,'column':ci,'dimension':d,'method':m,'department':g,'actual':row.cells[ci].text,'expected':display,'numerator':n,'five_dimension_denominator':den,'source_key':['primary_full4481',m,'departments',g,'profile','attention_share',d]})
assert len(table3)==30
note25=[]
for ri,row in enumerate(note.tables[0].rows[1:],1):
    dim=next(d for d,n in DN.items() if n==row.cells[0].text)
    for col,(m,ref) in enumerate((('A_C','symmetric A/C'),(RK,'A'),(RK,'C'),(LK,'A'),(LK,'C')),1):
        source=wrow(m,ref,dim);v=source['metrics']['kappa']['estimate'];expect='NA' if v is None else f'{v:.3f}'
        check('explanation25_kappa',row.cells[col].text,expect,f'note.table0.r{ri}.c{col}')
        note25.append({'row':ri,'column':col,'dimension':dim,'method':m,'reference':ref,'actual':row.cells[col].text,'expected':expect})
assert len(note25)==25

# Every delivery CSV cell is compared against independently assembled JSON fields.
csv420=DEL/'编码检验补充统计表.csv'
with csv420.open(encoding='utf-8-sig',newline='') as h:given=list(csv.DictReader(h))
expected420={}
for r in allrecords:
    out={'材料':CN[r['component']],'部门':GN[r['group']],'权重':WN[r['weighting']],'比较方法':MN[r['method']],'参照':RN[r['reference']],'维度':DN[r['dimension']],
         '有效配对数':r['paired_n'],'左侧编码1单元数':r['left_positive_n'],'右侧编码1单元数':r['reference_positive_n']}
    raw=r.get('raw_confusion',r.get('symmetric_raw_cells'));weighted=r.get('weighted_confusion',r.get('symmetric_weighted_cells'))
    for suffix,v,w in zip(('左1右1','左1右0','左0右1','左0右0'),raw.values(),weighted.values()):out['未加权'+suffix]=v;out['设计加权'+suffix]=w
    for metric,cn in MET.items():
        v=r['metrics'].get(metric)
        if v is None:
            for suffix in ('','近似95%区间下限','近似95%区间上限','区间有效重复数','区间未定义重复数'):out[cn+suffix]='不适用'
        else:
            ci=v['ci95_primary'];out[cn]=v['estimate'];out[cn+'近似95%区间下限']=ci['low'];out[cn+'近似95%区间上限']=ci['high'];out[cn+'区间有效重复数']=ci['defined_repetitions'];out[cn+'区间未定义重复数']=ci['undefined_repetitions']
    key=tuple(out[k] for k in ('材料','部门','权重','比较方法','参照','维度'))
    assert key not in expected420;expected420[key]=out
assert len(given)==len(expected420)==420
csv_reviews=[]
def cell_match(actual,expected):
    if expected is None:return actual=='NA'
    if isinstance(expected,(int,float)):
        try:return math.isclose(float(actual),expected,rel_tol=1e-12,abs_tol=1e-12)
        except ValueError:return False
    return actual==expected
def review_csv(rows,expected,keys,name):
    seen=set();n=0;bad=[];audit=[]
    for i,r in enumerate(rows,2):
        k=tuple(r[x] for x in keys);assert k not in seen,(name,'duplicate',k);seen.add(k)
        e=expected[k];assert set(r)==set(e),(name,'header_difference',set(r)^set(e))
        fails=[]
        for col,v in e.items():
            n+=1
            if not cell_match(r[col],v):fails.append({'column':col,'actual':r[col],'expected':v})
        if fails:bad.append({'row':i,'key':k,'mismatches':fails})
        audit.append({'csv_row':i,'source_key':k,'columns_compared':len(r),'pass':not fails,'actual_row_sha256':canonical(r),'expected_json_field_row_sha256':canonical(e)})
    assert seen==set(expected)
    errors.extend({'check':name,**x} for x in bad)
    return {'file':str(DEL/name),'sha256':sha(DEL/name),'rows':len(rows),'columns':len(rows[0]),'cells_compared':n,'mismatch_n':sum(len(x['mismatches']) for x in bad),'complete_key_set_matches':True,'row_checks':audit,'mismatches':bad}
csv_reviews.append(review_csv(given,expected420,('材料','部门','权重','比较方法','参照','维度'),'编码检验补充统计表.csv'))
expected120={}
for r in details:
    for m,v in r['methods'].items():
        out={'编码方法':MN[m],'作为条件的维度':DN[r['condition_dimension']],'共同表达的维度':DN[r['cooccurring_dimension']],'部门一':GN[r['department_one']],'部门二':GN[r['department_two']]}
        for side,cn in [('one','部门一'),('two','部门二')]:
            out[cn+'共同表达单元数']=v[side]['both_n'];out[cn+'条件维度单元数']=v[side]['condition_positive_n'];out[cn+'条件比例%']=None if v[side]['conditional_share'] is None else 100*v[side]['conditional_share']
            # Independent division check; no manuscript-generation helper imported.
            ratio=F(v[side]['both_n'],int(v[side]['condition_positive_n'])) if v[side]['condition_positive_n'] else None
            check('conditional_integer_ratio',v[side]['conditional_share'],float(ratio) if ratio is not None else None,f'{m}/{r["condition_dimension"]}/{r["cooccurring_dimension"]}/{r["department_one"]}/{r["department_two"]}/{side}')
        out['部门一减部门二百分点']=v['difference_percentage_points']
        for src,prefix in [('definition6','定义口径'),('composition56','文件构成口径')]:
            z=v[src]
            for s,c in [('scenario_n','数'),('one_higher_n','部门一较高数'),('two_higher_n','部门二较高数'),('equal_n','相等数'),('undefined_n','未定义数')]:out[prefix+c]=z[s]
            lo,hi=z['difference_percentage_point_range'] or (None,None)
            out[prefix+'差值最小百分点']=lo;out[prefix+'差值最大百分点']=hi
        k=tuple(out[n] for n in ('编码方法','作为条件的维度','共同表达的维度','部门一','部门二'));assert k not in expected120;expected120[k]=out
with (DEL/'维度共同表达补充统计表.csv').open(encoding='utf-8-sig',newline='') as h:given=list(csv.DictReader(h))
assert len(given)==len(expected120)==120
csv_reviews.append(review_csv(given,expected120,('编码方法','作为条件的维度','共同表达的维度','部门一','部门二'),'维度共同表达补充统计表.csv'))

# Only original bibliography and citation tokens are used from the old manuscript.
def ref_index(d):return next(i for i,p in enumerate(d.paragraphs) if p.text.strip()=='参考文献')
ni,oi=ref_index(doc),ref_index(old)
newrefs=[p.text for p in doc.paragraphs[ni+1:] if p.text.strip()]
oldrefs=[p.text for p in old.paragraphs[oi+1:] if p.text.strip()]
check('reference_count',len(newrefs),66,'bibliography');check('all_reference_original_text_and_order',newrefs,oldrefs,'bibliography')
def cite_sequence(d,stop):
    result=[]
    for i,p in enumerate(d.paragraphs[:stop]):
        for m in re.finditer(r'\[[0-9,， \-—–]+\]',p.text):result.append(m.group(0))
    return result
newc,oldc=cite_sequence(doc,ni),cite_sequence(old,oi)
check('complete_body_citation_token_sequence',newc,oldc,'all_body_paragraphs_before_bibliography')
refaudit={'old_baseline':{'path':str(OLD),'sha256':EXPECTED[OLD]},'old_file_read_scope':'mechanical bibliography and body citation tokens only; old conclusions/QA not reused','reference_n':len(newrefs),'reference_text_order_exact':newrefs==oldrefs,'reference_list_sha256':canonical(newrefs),'citation_token_n':len(newc),'citation_sequence_exact':newc==oldc,'citation_sequence_sha256':canonical(newc),'new_reference_sequence':newrefs,'body_citation_token_sequence':newc}

def pair(m,g,a,b):return next(x for x in findings['primary_full4481'][m]['departments'][g]['all10_pairs'] if (x['dimension_a'],x['dimension_b'])==(a,b))
def q(m,g,d):return findings['primary_full4481'][m]['departments'][g]['profile']['attention_share'][d]
def percentage_strings(text):return re.findall(r'(\d+\.\d+)%',text)
def decstrings(text):return re.findall(r'\d+\.\d+',text)
def pcheck(i,expected,percent=True):
    actual=percentage_strings(doc.paragraphs[i].text) if percent else decstrings(doc.paragraphs[i].text)
    check('paragraph_numeric_sequence',actual,expected,f'main.p0={i}/p1={i+1}')
    return {'p0':i,'p1':i+1,'text':doc.paragraphs[i].text,'actual_numbers':actual,'expected_numbers':expected,'pass':actual==expected}
paragraph_reviews=[]
abstract=[f'{q(RK,"NMPA","S")*100:.2f}',f'{q(RK,"NHSA","Ec")*100:.2f}']+[f'{pair(m,g,"Ec","F")["p_b_given_a"]*100:.2f}' for g,m in [('NHC',RK),('NHC',LK),('NHSA',RK),('NHSA',LK)]]
for i in (1,6):paragraph_reviews.append(pcheck(i,abstract))
ab=[wrow('A_C','symmetric A/C',d)['metrics']['kappa']['estimate'] for d in DS]
three={m:[wrow(m,r,d)['metrics']['kappa']['estimate'] for d in ('S','E','Ec') for r in ('A','C')] for m in (RK,LK)}
vals=ab+[min(three[RK]),max(three[RK]),min(three[LK]),max(three[LK])]+[wrow(m,r,'I')['metrics']['kappa']['estimate'] for m in (RK,LK) for r in ('A','C')]+[wrow(m,r,'F')['metrics']['kappa']['estimate'] for m in (RK,LK) for r in ('A','C')]+[wrow(m,r,'F')['metrics']['f1']['estimate'] for m,r in ((LK,'A'),(RK,'A'),(LK,'C'),(RK,'C'))]
paragraph_reviews.append(pcheck(55,[f'{v:.3f}' for v in vals],False))
paragraph_reviews.append(pcheck(66,[f'{q(RK,g,d)*100:.2f}' for g in ('NMPA','NHSA','NHC') for d in DS]))
directions=[]
for one,two in combinations(('NMPA','NHSA','NHC'),2):
    for d in DS:
        ds=[q(m,one,d)-q(m,two,d) for m in (RK,LK)];sign=[1 if v>1e-12 else -1 if v< -1e-12 else 0 for v in ds]
        directions.append({'department_one':one,'department_two':two,'dimension':d,'rule_difference':ds[0],'LLM_difference':ds[1],'same_direction':sign[0]==sign[1]})
check('15_department_dimension_directions_same_count',sum(x['same_direction'] for x in directions),14,'main.p0=71')
check('only_direction_disagreement',[(x['department_one'],x['department_two'],x['dimension']) for x in directions if not x['same_direction']],[('NHSA','NHC','I')],'main.p0=71')
check('p71_states14','14项点估计方向一致' in doc.paragraphs[71].text,True,'main.p0=71')
paragraph_reviews.append(pcheck(71,[f'{q(m,g,"I")*100:.2f}' for m,g in [(RK,'NHSA'),(RK,'NHC'),(LK,'NHSA'),(LK,'NHC')]]))
paragraph_reviews.append(pcheck(72,[f'{q(m,g,d)*100:.2f}' for m,g,d in [(RK,'NMPA','S'),(LK,'NMPA','S'),(RK,'NHSA','Ec'),(LK,'NHSA','Ec')]]))
ratio72=[]
for m,g,d in [(RK,'NMPA','S'),(LK,'NMPA','S'),(RK,'NHSA','Ec'),(LK,'NHSA','Ec')]:
    z=findings['primary_full4481'][m]['departments'][g]['profile'];ratio72.append(f'{int(z["coding_count"][d])}/{int(z["five_dimension_coding_count"])}')
check('p72_integer_denominators',re.findall(r'\d+/\d+',doc.paragraphs[72].text),ratio72,'main.p0=72')
for m in (RK,LK):check('p72_Ec_56_direction',findings['summary'][m]['composition56']['attention_share_Ec']['NHSA_higher_n'],56,m)
series=next(x for x in findings['composition56'] if x['scenario']=='policy_series_equal')
check('p72_F_series_reverse',series['methods'][RK]['contrasts']['attention_share_F']['NHSA_minus_NHC']>0,True,'main.p0=72')
paragraph_reviews.append(pcheck(75,[f'{pair(m,g,"Ec","F")["p_b_given_a"]*100:.2f}' for m,g in [(RK,'NHC'),(RK,'NHSA'),(LK,'NHC'),(LK,'NHSA')]]))
range75=[]
for m in (RK,LK):
    z=findings['summary'][m]['composition56']['P_F_given_Ec']['NHSA_minus_NHC_range'];range75 += [f'{-z[1]:.2f}',f'{-z[0]:.2f}']
check('p75_range_direction_and_rounding',re.findall(r'(\d+\.\d+)—(\d+\.\d+)',doc.paragraphs[75].text),[(range75[0],range75[1]),(range75[2],range75[3])],'main.p0=75')
for pidx,reverse,gorder in [(75,False,('NHC','NHSA')),(76,True,('NHSA','NHC'))]:
    wanted=[]
    for m in (RK,LK):
        for g in gorder:
            z=pair(m,g,'Ec','F');wanted.append(f'{z["both_n"]}/{int(z["b_positive" if reverse else "a_positive"])}')
        for s in ('definition6','composition56'):
            v=findings['summary'][m][s]['P_Ec_given_F' if reverse else 'P_F_given_Ec'];check('p75_76_all_scenarios_direction',v['NHSA_higher_n' if reverse else 'NHC_higher_n'],v['scenario_n'],f'p{pidx}/{m}/{s}')
    check('p75_76_exact_conditional_fractions',re.findall(r'\d+/\d+',doc.paragraphs[pidx].text),wanted,f'main.p0={pidx}')
paragraph_reviews.append(pcheck(76,[f'{pair(m,g,"Ec","F")["p_a_given_b"]*100:.2f}' for m,g in [(RK,'NHSA'),(RK,'NHC'),(LK,'NHSA'),(LK,'NHC')]]))
paragraph_reviews.append(pcheck(82,[f'{pair(m,g,"S","E")["p_b_given_a"]*100:.2f}' for m in (RK,LK) for g in ('NHSA','NMPA','NHC')]))
check('p82_exact_conditional_fractions',re.findall(r'\d+/\d+',doc.paragraphs[82].text),[f'{pair(m,g,"S","E")["both_n"]}/{int(pair(m,g,"S","E")["a_positive"])}' for m in (RK,LK) for g in ('NHSA','NMPA','NHC')],'main.p0=82')
for z in details:
    if z['condition_dimension']=='S' and z['cooccurring_dimension']=='E' and (z['department_one'],z['department_two']) in [('NMPA','NHSA'),('NHSA','NHC')]:
        for m,v in z['methods'].items():
            for sc in ('definition6','composition56'):check('p82_all_scenarios_direction',v[sc]['two_higher_n' if z['department_one']=='NMPA' else 'one_higher_n'],v[sc]['scenario_n'],f'{z["department_one"]}/{z["department_two"]}/{m}/{sc}')

semantic_reviews=[
 {'locations':[25,38,39,47,57,60,70,113],'judgment':'SUPPORTED','basis':'多标签同维至多一次；q以五维句次合计为分母，coverage以单位数为分母，条件比例以给定维阳性单元为分母；180估计与4481固定文本分开。等文件q/条件比例的非零分母排除有说明。'},
 {'locations':[49,50,54,55],'judgment':'SUPPORTED_WITH_IDENTITY_UPDATE_PENDING','basis':'当前文本明示A/B初判及T为AI，B对应原C；κ用于一致性，P/R/F1为相对所列参照的分类表现；概率180与挑战60分开。root随后按用户最新确认将人工事后复核改为真人裁定后保留，不倒记真人独立盲评。'},
 {'locations':[51],'judgment':'SUPPORTED','basis':'0.80明示项目检查目标，未称期刊统一线；真实I/F低值和参照改变导致方法排序变化均保留。'},
 {'locations':[71,72,75,76,82,89,90,105],'judgment':'SUPPORTED','basis':'14/15点方向一致且NHSA/NHC创新性例外保留；共同表达及双向分母明确，6定义/56构成分别陈述，未据共现比例声称强关联、实际公平或因果；结语同段保留文本范围与执行效果未验证。'},
 {'locations':[109],'judgment':'BLOCKED_SOURCE_BINDING','basis':'当前完整LLM03冻结profile注明stability_not_evaluated；本段75/74/73稳定性数值未有03同版实际复测绑定，不能沿用旧02。root已接受删除，此审计不改稿。','text':doc.paragraphs[109].text},
 {'locations':['explanation p4'],'judgment':'WORDING_CORRECTION_REQUIRED','basis':'“其59次重抽样”混淆每次抽取单元数量与重复次数；原方案每次抽59个单元、共5000次。root已接受澄清。','text':note.paragraphs[4].text},
 {'locations':['explanation p6'],'judgment':'SUPPORTED_SCOPE_LIMITATION','basis':'明确已有人工/模型语境接触、完整句匹配为下限及两条目标暴露，未声称未命中即未暴露；未读取267预留正文或答案。该审查只核陈述范围，不重新审计暴露计数。'},
 {'locations':['explanation p13'],'judgment':'SUPPORTED_SCOPE_LIMITATION','basis':'明确.80非统一发表线、固定参照下I双侧目标不可行及不改参照/挑样本/拼接；0.710311沿用先前已独立精确QA，未在本任务重做枚举。'},
]
for m in (RK,LK):
    for d,g in [('S','NMPA'),('E','NMPA'),('I','NMPA'),('Ec','NHSA'),('F','NHC')]:
        check('main71_105_highest_dimension_department',q(m,g,d)==max(q(m,h,d) for h in ('NMPA','NHSA','NHC')),True,f'{m}/{d}')
    for g,d in [('NMPA','S'),('NHSA','Ec'),('NHC','F')]:
        check('conclusion_department_highest_dimension',q(m,g,d)==max(q(m,g,k) for k in DS),True,f'{m}/{g}')
for p,h in EXPECTED.items():assert sha(p)==h,(p,'changed_during_audit')
for item in freeze['files']:assert sha(item['path'])==item['sha256'],item['path']
report={'status':'NUMERIC_REFERENCE_CORE_PASS_WITH_CONTENT_BLOCKER','performed_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'independence':'单独实际审读新稿和说明并独立映射统计JSON；已可见root数值/视觉结论及说明旧FAIL修正信息；未复用旧稿QA、同步作者数值函数或原个体预测。',
 'inputs':[{'path':str(p),'sha256':h} for p,h in EXPECTED.items()],'source_freeze_file_n':len(freeze['files']),'frozen_sources_unchanged':True,'main_paragraphs_actual_read_scope':'p0 0–113; all four tables; new bibliography mechanically compared in full; root handles actual full-page visual QA','old_read_scope':'only bibliography and citation token sequence comparison',
 'numeric_status':'PASS' if not errors else 'FAIL','errors':errors,'checks_n':len(checks),'failed_checks_n':sum(not c['pass'] for c in checks),'table2_80':table2,'table3_30':table3,'explanation25_kappa':note25,'csv_cell_audits':csv_reviews,'bibliography_and_citations':refaudit,'paragraph_numeric_checks':paragraph_reviews,'15_department_dimension_directions':directions,'semantic_reviews':semantic_reviews,
 'history_preserved':{'old_explanation_sha256':'1d4a787965db122570160347ec6318c77cc645d23f4411b992386654fd779cde','old_explanation_status':'FAIL_25_KAPPA_STALE_ROOT_FOUND_AND_ARCHIVED','current_explanation_sha256':EXPECTED[NOTE],'current_25_kappa_status':'PASS' if not any(x['check']=='explanation25_kappa' for x in errors) else 'FAIL'},
 'not_performed':['no model/classifier/statistical rerun','no raw prediction semantic reading','no reserved267 text/answer reading','no new literature/policy claim verification','no visual duplication','no input/document/CSV edits'],
 'delivery_gate':'待root删除p109无03来源稳定性主张、澄清说明59个/5000次并落实用户最新A/B真人裁定身份后，只补审改动段、25κ、完整引用序列和新SHA；本轮完整80/30/420/120结果可按未变字节继续绑定。'}
assert not (OUT/'independent_main_and_companion_review_INITIAL.json').exists()
(OUT/'independent_main_and_companion_review_INITIAL.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
print(json.dumps({'status':report['status'],'numeric_status':report['numeric_status'],'errors':errors,'table2_n':len(table2),'table3_n':len(table3),'note_kappa_n':len(note25),'csv_checks':[{'rows':z['rows'],'cells':z['cells_compared'],'mismatches':z['mismatch_n']} for z in csv_reviews],'references':{'n':len(newrefs),'exact':newrefs==oldrefs,'citation_tokens':len(newc),'citation_sequence_exact':newc==oldc},'report':str(OUT/'independent_main_and_companion_review_INITIAL.json'),'report_sha256':sha(OUT/'independent_main_and_companion_review_INITIAL.json')},ensure_ascii=False))
