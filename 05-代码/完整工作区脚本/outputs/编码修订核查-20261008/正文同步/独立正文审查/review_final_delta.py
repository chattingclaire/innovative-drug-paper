"""Final document delta audit; source inputs remain read-only."""
from pathlib import Path
import datetime, hashlib, json, math, re, sys
sys.dont_write_bytecode=True
from docx import Document
ROOT=Path('/Users/rainbow/Desktop/战略院/药企项目-创新药')
BASE=ROOT/'outputs/编码修订核查-20261008';OUT=BASE/'正文同步/独立正文审查';DEL=BASE/'本轮结果'
STAT=BASE/'统计复算/完整39与LLM03-同步复算01'
ARCH=BASE/'正文同步/接入裁定与清除非同源复测前记录'
MAIN=DEL/'创新药评价的注意力配置差异与协调路径.docx';NOTE=DEL/'编码复核说明与真人检验事项.docx'
OLD=ROOT/'outputs/final-three-department-20261008-human-validation-template/创新药评价的注意力配置差异与协调路径.docx'
INITIAL=OUT/'independent_main_and_companion_review_INITIAL.json'
HASHES={MAIN:'9f817ee2cec4c1696fc2e0b724eea49670e48190911943fe1cfb3b9ed7561489',NOTE:'441b742713b1f7abc6b3ae2272a413949d9779ad94122e9807c0335c44e2f874',INITIAL:'95f28187419b8ec6a45bd96c8ae73b9fa94baf8a0b5c7e57145489ce718df4fa',STAT/'output_freeze.json':'f411935421447c6fd39fdf1887078dc6ba297e201ac5bed6498befe2ba753c3a',OLD:'4d61c921a8ea9837f9d31eb4cd09ace0af1839b51e2793232ee0f402ede905e0'}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(p):return json.loads(Path(p).read_text())
def canonical(x):return hashlib.sha256(json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
for p,h in HASHES.items():assert sha(p)==h,(p,'bound_SHA_mismatch')
initial=read(INITIAL);assert initial['numeric_status']=='PASS' and initial['failed_checks_n']==0
doc=Document(MAIN);note=Document(NOTE);prev=Document(ARCH/MAIN.name);old=Document(OLD)
freeze=read(STAT/'output_freeze.json')
for b in freeze['files']:assert sha(b['path'])==b['sha256']
diff=[{'p0':i,'p1':i+1,'before':a.text,'after':b.text} for i,(a,b) in enumerate(zip(prev.paragraphs,doc.paragraphs)) if a.text!=b.text]
assert len(prev.paragraphs)==len(doc.paragraphs) and {r['p0'] for r in diff}=={49,50,55,109}
assert len(prev.tables)==len(doc.tables)==4
assert [[[c.text for c in r.cells] for r in t.rows] for t in prev.tables]==[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
classrows=read(STAT/'machine_reference_classification.json')['comparisons']+read(STAT/'A_C_before_third_agreement.json')['comparisons']
lookup={(r['component'],r['group'],r['weighting'],r['method'],r['reference'],r['dimension']):r for r in classrows}
full=read(STAT/'full4481/完整定义构成核查.json')
finalvalues=[];fail=[]
def result(k,actual,expected,source):
    ok=actual==expected
    z={'location':k,'actual':actual,'expected':expected,'pass':ok,'source':source};finalvalues.append(z)
    if not ok:fail.append(z)
for z in initial['table2_80']:
    key=z['source_key'];r=lookup[tuple(key[:6])];v=r['metrics'][key[7]]['estimate']
    expected='NA' if v is None else f'{v:.3f}';side=0 if z['display_reference']=='A' else 1
    result(f'table2.r{z["row"]}.c{z["column"]}.{z["display_reference"]}',doc.tables[1].rows[z['row']].cells[z['column']].text.split('/')[side],expected,key)
for z in initial['table3_30']:
    profile=full['primary_full4481'][z['method']]['departments'][z['department']]['profile']
    numerator=int(profile['coding_count'][z['dimension']]);den=sum(int(v) for v in profile['coding_count'].values())
    result(f'table3.r{z["row"]}.c{z["column"]}',doc.tables[2].rows[z['row']].cells[z['column']].text,f'{100*numerator/den:.2f}',z['source_key'])
for z in initial['explanation25_kappa']:
    r=lookup['probability_canonical','ALL','eligible_design_weighted',z['method'],z['reference'],z['dimension']];v=r['metrics']['kappa']['estimate']
    result(f'note.table0.r{z["row"]}.c{z["column"]}',note.tables[0].rows[z['row']].cells[z['column']].text,'NA' if v is None else f'{v:.3f}',['probability_canonical','ALL','eligible_design_weighted',z['method'],z['reference'],z['dimension'],'metrics','kappa','estimate'])
assert len(finalvalues)==135 and not fail
for z in initial['csv_cell_audits']:assert sha(z['file'])==z['sha256'],('CSV_changed',z['file'])
def ridx(d):return next(i for i,p in enumerate(d.paragraphs) if p.text.strip()=='参考文献')
def refs(d):return [p.text for p in d.paragraphs[ridx(d)+1:] if p.text.strip()]
def cites(d):return [m.group(0) for p in d.paragraphs[:ridx(d)] for m in re.finditer(r'\[[0-9,， \-—–]+\]',p.text)]
assert len(refs(doc))==66 and refs(doc)==refs(old)
assert cites(doc)==cites(old) and len(cites(doc))==95
assert '[66]' in doc.paragraphs[109].text
assert not any(x in doc.paragraphs[109].text for x in ['75条','74条','73条','复测','重复性'])
assert '经真人裁定后保留' in doc.paragraphs[49].text and '初判由AI代理' in doc.paragraphs[49].text and 'AI复核者T' in doc.paragraphs[49].text
assert '两套人工审定参照的一致性' in doc.paragraphs[50].text and '相对于所列参照的分类表现' in doc.paragraphs[50].text
assert not ('讨论分歧前的独立结果' in doc.paragraphs[50].text)
assert '有放回抽59个单元' in note.paragraphs[4].text and '重复5000次' in note.paragraphs[4].text and '卫健层固定' in note.paragraphs[4].text
assert '已经真人裁定' in note.paragraphs[1].text and '不要求重复已完成的裁定' in note.paragraphs[1].text
assert '不倒记' in note.paragraphs[6].text and '独立真人盲评信度' in note.paragraphs[6].text
assert '不是《科学决策》统一发表线' in note.paragraphs[12].text
assert '0.710311' in note.paragraphs[13].text and '不能' in note.paragraphs[14].text
semantic=[
 {'scope':'main.p49/50/55','judgment':'PASS','basis':'实际读新段；AI初判与用户确认的真人裁定后保留分开，A/B称人工审定参照，T仍明确为AI补充参照。两套已裁定结果的一致性未倒记为新独立真人盲评信度；κ与相对参照P/R/F1区分正确，低值/参照分歧保留。'},
 {'scope':'main.p109','judgment':'PASS_BLOCKER_RESOLVED','basis':'实际读新段；删除未有完整03同源绑定的75/74/73稳定性主张，保留真实模型、推理强度、最大批量和4924范围及[66]任务记录引用；本轮未补跑稳定性。'},
 {'scope':'new explanation all22 paragraphs/table25','judgment':'PASS','basis':'实际完整读当前22段及25κ；明确真人裁定已完成、不要求重做，补记录与统一最终参照绑定分开；每次抽59个单元、共5000重复和卫健固定清楚。.80项目目标非期刊线、双侧I目标限制和真实低值保留；预留材料存在接触，不与现有裁定完成状态混淆。'},
 {'scope':'unchanged main abstract/55 numeric/66/71/72/75/76/82/conclusion','judgment':'PASS_REBOUND_TO_FINAL_BYTES','basis':'初审实际逐段、逐数、独立有理数/整数核查通过，当前除55角色措辞外文本逐字不变；55数值重新读取且135表值全部重新绑定当前统计JSON。71保留14/15及医保—卫健创新性方向不同；共同表达不作因果或实际公平解释，结语同段保留文本范围。'},
 {'scope':'papers/policy assertions newly retained in explanation','judgment':'LIMIT_OF_THIS_REVIEW','basis':'本任务未新增文献或政策检索；.80非普遍线的论文核验由root另完成。本审查核其表述范围与当前统计用途一致，不冒称重新核查论文原文。'}
]
for p,h in HASHES.items():assert sha(p)==h,(p,'changed_during_final_delta')
for z in initial['csv_cell_audits']:assert sha(z['file'])==z['sha256']
report={'status':'PASS_CURRENT_CONTENT_NUMBERS_REFERENCES_WITH_STATED_SCOPE','performed_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),
 'review_identity':'单独实际复核，已可见root初步数值/视觉意见和修复决定；不称完全盲审。未复用作者同步数值函数，不读个体预测或267预留正文/答案。',
 'current_bindings':[{'path':str(p),'sha256':h} for p,h in HASHES.items()],
 'initial_actual_review':{'path':str(INITIAL),'sha256':HASHES[INITIAL],'status':'numeric/reference core PASS with explicitly preserved p109 content blocker and note59 wording correction'},
 'actual_main_delta':diff,'main_tables_all_text_unchanged':True,'actual_current_table_value_n':135,'table2_n':80,'table3_n':30,'explanation_kappa_n':25,'current_values':finalvalues,
 'CSV_binding_recheck':[{'path':z['file'],'sha256':z['sha256'],'rows':z['rows'],'columns':z['columns'],'initial_actual_cells_compared':z['cells_compared'],'mismatches':0,'unchanged_in_final_delivery':True} for z in initial['csv_cell_audits']],
 'bibliography':{'source_old_sha256':HASHES[OLD],'reference_n':66,'all_original_reference_text_order_exact':True,'reference_list_sha256':canonical(refs(doc)),'citation_token_n':95,'complete_citation_sequence_exact':True,'citation_sequence_sha256':canonical(cites(doc))},
 'semantic_reviews':semantic,'blocking_issues':[],'preserved_history':{'original_stale25_kappa_explanation_SHA':'1d4a787965db122570160347ec6318c77cc645d23f4411b992386654fd779cde','original25_failure_not_erased':True,'initial_main_with_unsupported_stability_SHA':'042cf1f7154fa8b2ad915b6cbdb6948e8583a8f164fccff0f5e168868d151705','unsupported_stability_claim_removed_without_new_run':True,'initial_note_with59_wording_SHA':'ac7f161cd4ea2238c7a259df8314eb1932f14737b9498a751171d958ac76c046'},
 'not_performed':['no model/classifier/statistical rerun','no prediction or label edits','no original individual predictions or267 text/answers read','no repeated allCSV audit after byte-identity confirmed','no visual duplicate audit; root/team cover full current rendered pages','no claim that A/B/T are new independent human blind judgments','no manuscript/CSV edits']}
p=OUT/'independent_final_manuscript_companion_review.json';assert not p.exists();p.write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
print(json.dumps({'status':report['status'],'report':str(p),'sha256':sha(p),'tables':{'table2':80,'table3':30,'note_kappa':25},'CSV_cells_unchanged_and_initially_verified':24960,'references':66,'citation_tokens':95,'blocking_issues':[]},ensure_ascii=False))
