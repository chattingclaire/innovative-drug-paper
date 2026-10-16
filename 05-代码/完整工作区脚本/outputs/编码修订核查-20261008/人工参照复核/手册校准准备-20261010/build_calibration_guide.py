from pathlib import Path
import json
import hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE

HERE = Path(__file__).resolve().parent
data = json.loads((HERE / '既有AB创新性公平性分歧材料.json').read_text())
cases = {r['id']: r for r in data['records']}
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.1)
sec.left_margin = sec.right_margin = Cm(2.2)
for name, size in [('Normal',11),('Title',20),('Heading 1',16),('Heading 2',12)]:
    sty=doc.styles[name];sty.font.name='Times New Roman';sty.font.size=Pt(size);sty.font.color.rgb=RGBColor(0,0,0)
    rf=sty.element.get_or_add_rPr().rFonts
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rf.attrib.pop(qn('w:'+key),None)
    rf.set(qn('w:eastAsia'),'Songti SC' if name=='Normal' else 'Heiti SC')
    sty.paragraph_format.line_spacing=1.25
    sty.paragraph_format.space_after=Pt(6)
    if name!='Normal':sty.paragraph_format.space_before=Pt(10)
for sty in doc.styles:
    for border in sty.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)

def p(text, style=None):
    return doc.add_paragraph(text,style)

def h(text):
    doc.add_heading(text,2)

def page(title):
    heading=doc.add_heading(title,1)
    heading.paragraph_format.page_break_before=True

def link(par, label, url):
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),par.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245782');props.append(color)
    run.append(props);txt=OxmlElement('w:t');txt.text=label;run.append(txt);hyp.append(run);par._p.append(hyp)

def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for c,txt in zip(t.add_row().cells,values):c.text=str(txt)
    for rn,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            if rn==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EEF4');pr.append(fill)
            for par in c.paragraphs:
                par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.space_after=Pt(2)
                for run in par.runs:run.font.size=Pt(10);run.bold=rn==0
    p('')
    return t

doc.add_heading('创新性与公平性编码校准说明',0)
p('本说明供研究者与编码员A、B校准共同手册。先依据政策原文统一判据，再检验两名编码员能否独立执行。以下为拟采用的操作规则和教学判读；讨论确认后才能冻结用于新的独立检验，原有标签不覆盖。')
doc.add_heading('一 先统一测量对象',1)
p('本文记录五维评价议题及直接制度安排在政策句段中的呈现。1表示实质涉及该维，不表示药品已具备优势、患者已受益或措施已经实施。每个记录的每一维最多记1次，同一记录可涉及多维。必要语境用于补全当前命题，不另计次数，也不使全部下级条款继承标签。')
p('判读顺序固定为：当前完整命题 → 药物或治疗对象 → 实质评价或制度关系 → 必要语境 → 标签及原文依据。否定、条件、限制和撤回须完整保留。')
doc.add_heading('二 现有分歧在哪里',1)
table(['材料','维度','A记1','B记1','分歧'],[
    ['概率样本180条','创新性',1,3,2],['概率样本180条','公平性',12,12,6],
    ['挑战材料60条','创新性',12,15,5],['挑战材料60条','公平性',28,34,6]], [4.5,3,2.5,2.5,2.5])
p('表内是既有标签的未加权计数，两类材料分开。概率样本原设计加权κ为创新性0.520、公平性0.641。创新性判1案例很少，两个具体分歧即可明显影响κ；仍需处理这些分歧，并增加独立检验中相关案例的信息量。原归档C在本说明中称B。')
p('校准既要阅读分歧，也要检查双方共同判1和共同判0的材料，防止把共同误解当作正确。下列规则不以提高某个分数或调整部门占比为选择依据。')

page('三 创新性的拟统一判据')
h('实质改良与一般变化分开')
p('当前直接讨论药物或治疗属性的实质改良，包括评价、否定和不确定判断，可记I；援引确切改良资格的，按下述资格条件判断。仅适应症、剂型、给药途径或工艺变化，以及一般设计建议，须明确既有治疗或使用局限，或拟用设计所针对的操作障碍，并建立具体改变与治疗或使用表现改善的关系。无需已证优效。化学药注册分类同时强调优化和临床优势，为区分“变化”与“改良”提供依据〔1〕。不同药品类别各按其分类；纳入针对使用障碍的拟用设计是本研究的操作化选择。')
h('具体创新分类作为资格条件可以纳入')
p('明确的法定创新或改良分类被用来决定差异化准入、续约或其他资格时，可按创新分类的制度使用记I；无需重新进行创新评价。理由必须写明是哪一分类、决定什么资格，不能写成已证明临床优越或研发支持。医保局对2023年续约规则的解读明确说明了注册分类与创新支持的联系〔2〕。泛称“创新药”、目录名称或仅新获批仍不充分；中药3类等不能自动等同创新药类别。')
h('研发支持须有具体可用的安排')
p('面向明确创新研发对象的审评团队、专门技术咨询、研究方案反馈、资源或提速机会，可以记I。申请人承担一般安全义务、列出沟通方式，或发布技术指南的概括性宗旨，不因对象为创新药就另记I。当前决定启动或批准开展确指计划时，只有共同材料已给出该计划的具体支持机制，才可补足判断；不要求措施已经实施。')
h('同一门槛适用于支持的设立和限制')
p('设立、适用、限制或撤回明确创新支持都可表达I。只有文件废止、计划名称或“支持创新”的泛化目的时，不自行推导出具体支持变化。一般中药制度完善不因处于创新政策中自动纳入。')
p('以上是本研究对文本关注的操作化。分类的制度使用、属性评价和研发支持可以同属创新性，但解释结果时应说明具体内容，不能混称药品技术新颖性。')

page('四 公平性的拟统一判据')
h('明确的公平评价要求可以表达关注')
p('当前以药品价值或患者治疗为对象，实质要求评价、比较公平性，或提供、核验该维信息，可记F；不额外要求每句再列出具体障碍。此类理由标明“概括性公平评价或信息要求”，不据此推断障碍类型或公平改善。仅标题、栏目名、字段填写和材料流转仍不充分。程序公正、企业竞争公平等不同对象不据词语相同纳入。')
h('栏目内的成员仍按自身实质内容判断')
p('医保“公平性”栏目包含保障短板，也包含临床管理和价格费用等内容。总领要求评价公平，不意味着栏目内每一项都记F。成员若仅涉及价格按经济性判断，只有患者能否负担、按需保障或获取差异等联系时才另记F。这一区分统一适用于五维总领与成员，不把整栏含义复制到每句。')
h('治疗缺口须与有针对性的机会安排相连')
p('原文明确某类患者缺少有效治疗选项、境内批准选项或现实供给，并据此安排针对性的申请机会、审评资源或保障，可以记F。三类缺口应分别说明，不能把无有效药写成供应短缺。单有严重疾病、罕见病、未上市、疗效更好或一般提速不充分。向上游研发延伸这一判据是研究操作化，依据按治疗需要配置机会的公平概念〔3—4〕，不是文献直接规定的二元标签。')
h('支付手续和效率改善不自动代表患者公平')
p('基金预算阈值、谈判协议期限、结算方式、电子处方流转以及基本药身份，本身不足以证明患者获取障碍或保障连续性。患者可负担性须表达费用与支付能力的关系；地区物流措施须有可定位的获取差异或障碍。若解释依赖同政策官方解读，应把对应片段统一补入所有编码材料，不能由个别编码员自行补充。')
h('纯宗旨和实质规定分开')
p('“为规范短缺药优先审评而制定程序”属于制定目的；明确允许短缺药申请优先审评，或限制此机会的适用范围，则是实质安排。前者不从目的推措施，后者不因限制或否定语气而丢掉公平议题。')

page('五 用原文检验判据')
p('以下为开发讲解，不是独立检验题。引文保留当前句的条件和否定；完整来源与其他争议记录另存于配套分歧材料。')
examples=[
 ('RECODE25-01955','创新性建议0','适应症或功能主治重大变化，不等于已经表达实质改良；当前用途是医保支付范围申报。'),
 ('RECODE25-02220','创新性建议0','当前说明起草、发布指南及一般指导功能；没有独立设置专门技术服务或具体支持机会。'),
 ('RECODE25-00118','公平性建议1','临床急需且境内未上市的治疗选项缺口，与直接纳入优先审评的具体机会相连；不声称所有替代治疗都不存在。'),
 ('RECODE25-00099','公平性建议1','当前明确限制针对短缺药的优先机会。限制性命题仍讨论按缺口配置的资源，不能摘去“不再”。'),
 ('RECODE25-04894','公平性建议0','处方系统连通和流转效率未在当前表达患者获取差异、支付能力或按需补偿关系。'),
]
for rid,title,reason in examples:
    h(rid.replace('RECODE25-','例 ')+ '  '+title)
    p('原文：'+cases[rid]['current_unit_text'])
    p('判读：'+reason)

page('六 校准后怎样检验')
p('A、B先分别试编码开发材料，再讨论分歧并说明原文依据。凡讨论改变判据，应写成可用于同类句段的通用规则，同时保留各自原判断和变更理由。Mayring（2000）把定义、例子与编码规则结合，适合用于本轮手册校准〔5〕。')
p('手册冻结后，在未用于上述讨论的材料上独立编码；材料提供同一必要语境，遮蔽对方和机器答案。先封存两份独立判断，按维度计算κ、四格、类别数量及区间，再裁定分歧形成机器评价参照。κ≥0.80为本项目预定目标，不能将协商后的共识用于证明独立信度。高一致性还须有概念和原文依据〔6〕。')
p('概率抽样用于估计目标语料中的表现；稀少类若增补，应预先规定设计并保留纳入概率。教学和有目的挑战材料用于检验边界，不能未经处理混入总体指标。仅增加大量双方都判0的普通句，不等于充分检验创新性和公平性。')
p('两种机器方法使用冻结后的同一操作定义，各自落实为规则或提示。开发材料用于修改，独立检验用于报告最终表现；Pangakis和Wolken（2025）采用共同手册，并分开提示修订与留出人工检验〔7〕。当前未形成新的A、B独立提交，不报告新的κ或“已达标”。')
doc.add_heading('主要依据',1)
refs=[
 ('〔1〕国家药监局（2020）化学药品注册分类及申报资料要求。用于区分创新分类及实质改良。','https://english.nmpa.gov.cn/2020-06/30/c_528659.htm'),
 ('〔2〕国家医保局（2023）关于谈判药品续约规则的解读。第三部分第二项及第五问。','https://www.nhsa.gov.cn/art/2023/7/21/art_105_11062.html'),
 ('〔3〕WHO Health equity。用于公平概念，不是逐句编码规范。','https://www.who.int/health-topics/health-equity'),
 ('〔4〕Oliver A, Mossialos E（2004）Equity of access to health care: outlining the foundations for action. Journal of Epidemiology and Community Health, 58:655–658.','https://doi.org/10.1136/jech.2003.017731'),
 ('〔5〕Mayring P（2000）Qualitative Content Analysis. Forum Qualitative Sozialforschung / Forum: Qualitative Social Research, 1(2), Art.20.','https://doi.org/10.17169/fqs-1.2.1089'),
 ('〔6〕Artstein R, Poesio M（2008）Inter-Coder Agreement for Computational Linguistics. Computational Linguistics, 34(4):555–596.','https://doi.org/10.1162/coli.07-034-R2'),
 ('〔7〕Pangakis N, Wolken S（2025）Keeping Humans in the Loop: Human-Centered Automated Annotation with Generative AI. ICWSM, 19(1):1471–1492.','https://doi.org/10.1609/icwsm.v19i1.35883'),
]
for text,url in refs:
    par=p(text+' ');par.paragraph_format.space_after=Pt(5)
    for run in par.runs:run.font.size=Pt(9)
    link(par,'原文',url)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='创新性与公平性编码校准说明'
doc.core_properties.subject='共用编码手册校准与独立人工检验'
doc.core_properties.author=''
out=HERE/'创新性与公平性编码校准说明.docx';doc.save(out)
print(json.dumps({'path':str(out),'sha256':hashlib.sha256(out.read_bytes()).hexdigest()},ensure_ascii=False))
