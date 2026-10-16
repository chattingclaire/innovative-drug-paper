"""Read-only saved-workbook verification and protection checks."""
import hashlib
import json
from pathlib import Path
import openpyxl

ADMIN = Path(__file__).resolve().parent
CAL = ADMIN.parent
PROXY = CAL / '代理试填-20261010'
OUT = PROXY / '复核反馈处理-20261010'
data = json.loads((OUT / '复核处理记录.json').read_text())
payload = json.loads((ADMIN / '无答案校准内容.json').read_text())
all_checks = {'workbooks': {}, 'protected_inputs_unchanged': {}, 'new_kappa_computed': False}
for filepath, expected in data['preserved_inputs'].items():
    digest = hashlib.sha256(Path(filepath).read_bytes()).hexdigest()
    assert digest == expected, filepath
    all_checks['protected_inputs_unchanged'][filepath] = digest

for coder in ['A', 'B']:
    source = openpyxl.load_workbook(PROXY / f'代理试填-{coder}（供真人复核）.xlsx')
    path = OUT / f'校准复核工作表-{coder}.xlsx'
    saved = openpyxl.load_workbook(path)
    assert saved.sheetnames == source.sheetnames + ['复核记录', '引文依据']
    for name in ['阅读材料', '编码手册']:
        assert list(saved[name].values) == list(source[name].values), name
    sh = saved['填写']
    assert [sh.cell(r, 1).value for r in range(6, 46)] == [u['material_id'] for u in payload['units']]
    assert [sh.cell(r, 2).value for r in range(6, 46)] == [source['填写'].cell(r, 2).value for r in range(6, 46)]
    blanks = []
    for i, r in enumerate(data['workbooks'][coder]['rows']):
        row, dim = divmod(i, 5)
        value = sh.cell(row + 6, dim + 3).value
        assert value == r['working_label'], (coder, r['material_id'], r['dimension'])
        assert value is None or type(value) is int and value in (0, 1)
        if value is None:
            blanks.append([r['material_id'], r['dimension']])
        assert r['reason'] in sh.cell(row + 6, dim + 8).value
        rec = saved['复核记录']
        assert rec.cell(i + 6, 1).value == r['material_id']
        assert rec.cell(i + 6, 2).value == r['dimension_name']
        assert rec.cell(i + 6, 3).value == r['original_label']
        assert rec.cell(i + 6, 4).value == r['user_direction']
        assert rec.cell(i + 6, 5).value == r['state']
        assert rec.cell(i + 6, 7).value == r['reason']
        assert rec.cell(i + 6, 9).value is None
    expected_evidence = [(r, ev, n + 1) for r in data['workbooks'][coder]['rows'] for n, ev in enumerate(r['evidence_records'])]
    evidence = saved['引文依据']
    assert evidence.max_row == len(expected_evidence) + 5
    for i, (r, ev, n) in enumerate(expected_evidence, 6):
        assert evidence.cell(i, 1).value == r['material_id']
        assert evidence.cell(i, 2).value == r['dimension_name']
        assert evidence.cell(i, 3).value == n
        assert evidence.cell(i, 4).value == ('当前句段' if ev['scope'] == 'current' else '必要语境')
        assert evidence.cell(i, 5).value == ev['quote']
        assert evidence.cell(i, 7).value == ev['connection']
        assert '页码' in evidence.cell(i, 6).value
    for rr in range(24, 30):
        assert saved['填写说明'].cell(rr, 2).value is None
    assert sh.freeze_panes == source['填写'].freeze_panes
    assert sorted(sh.tables) == sorted(source['填写'].tables)
    for n in sh.tables:
        assert sh.tables[n].ref == source['填写'].tables[n].ref
    assert len(sh.data_validations.dataValidation) == len(source['填写'].data_validations.dataValidation) == 1
    dv = sh.data_validations.dataValidation[0]
    assert dv.type == 'whole' and str(dv.sqref) == 'C6:G45' and dv.formula1 == '0' and dv.formula2 == '1'
    errors = [(s.title, c.coordinate) for s in saved for row in s for c in row if c.data_type == 'e']
    assert not errors, errors
    assert not saved._external_links
    assert all(s.sheet_state == 'visible' for s in saved)
    all_checks['workbooks'][coder] = {
        'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'all_labels_match_processing_record': True, 'pending_labels': blanks,
        'all_200_reasons_saved': True, 'complete_evidence_rows': len(expected_evidence),
        'current_and_context_source_values_preserved': True, 'manual_values_preserved': True,
        'human_metadata_blank': True, 'native_table_validation_and_freeze_preserved': True,
        'errors': errors,
    }
(OUT / '保存与来源核查.json').write_text(json.dumps(all_checks, ensure_ascii=False, indent=2) + '\n')
(OUT / '真人复核提交').mkdir(exist_ok=True)
print(json.dumps({c: {'pending_labels': v['pending_labels'], 'complete_evidence_rows': v['complete_evidence_rows'], 'errors': v['errors']} for c, v in all_checks['workbooks'].items()}, ensure_ascii=False, indent=2))
