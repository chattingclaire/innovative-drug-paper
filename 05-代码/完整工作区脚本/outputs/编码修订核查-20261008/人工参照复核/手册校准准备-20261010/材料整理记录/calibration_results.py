"""Read-only XLSX ingestion; descriptive A/B development feedback, never a final test."""
from pathlib import Path
from fractions import Fraction
from collections import Counter
from datetime import datetime, date, timezone
import argparse
import hashlib
import json
import math
import io
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
DIMS = ('安全性','有效性','经济性','创新性','公平性')
SHEETS = ['填写','阅读材料','编码手册','填写说明']
IDS = {f'校准{i:02d}' for i in range(1,41)}

def sha_bytes(data):
    return hashlib.sha256(data).hexdigest()

def canonical_sha(value):
    return sha_bytes(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False).encode())

def scalar(value):
    return value.isoformat() if isinstance(value,(datetime,date)) else value

def protected_content(wb):
    """Ignore actual response cells; allow main rows to be sorted, join by ID."""
    result = {}
    for sheet in wb:
        fixed = {}
        for row in sheet:
            for cell in row:
                if sheet.title == '填写' and 6 <= cell.row <= 45 and cell.column <= 13:
                    continue
                if sheet.title == '填写说明' and 14 <= cell.row <= 20 and cell.column == 2:
                    continue
                if cell.value is not None:
                    fixed[cell.coordinate] = scalar(cell.value)
        if sheet.title == '填写':
            fixed['material_records'] = sorted([[sheet.cell(r,1).value,sheet.cell(r,2).value] for r in range(6,46)],key=lambda row:str(row[0]))
        result[sheet.title] = fixed
    return result

def kappa_feedback(pairs):
    if any(type(a) is not int or type(b) is not int or a not in (0,1) or b not in (0,1) for a,b in pairs):
        raise ValueError('κ输入必须为完整配对的整数0/1，不能补0或猜测缺失。')
    cells = Counter(pairs)
    n = len(pairs)
    a,b,c,d = [cells[key] for key in [(1,1),(1,0),(0,1),(0,0)]]
    result = dict(zip(('n11','n10','n01','n00'),(a,b,c,d)))
    result.update({'n':n,'A_1_n':a+b,'B_1_n':a+c})
    if n == 0:
        result.update({'agreement':None,'expected_agreement':None,'kappa':None,'na_reason':'无完整配对条目'})
        return result
    po = Fraction(a+d,n)
    pe = Fraction((a+b)*(a+c)+(c+d)*(b+d),n*n)
    k = (po-pe)/(1-pe) if pe != 1 else None
    result.update({'A_1_rate':float(Fraction(a+b,n)),'B_1_rate':float(Fraction(a+c,n)),
                   'agreement':float(po),'expected_agreement':float(pe),'kappa':float(k) if k is not None else None,
                   'kappa_exact':f'{k.numerator}/{k.denominator}' if k is not None else None,
                   'na_reason':'双方全部判同一单一类别；机会校正分母为0' if k is None else None,
                   'project_0_80_feedback':'未定义' if k is None else '达到项目校准目标' if k >= Fraction(4,5) else '低于项目校准目标'})
    return result

def inspect_submission(path, coder, baseline):
    path = Path(path).resolve()
    try:
        data = path.read_bytes()
    except OSError as error:
        return {'coder':coder,'file':str(path),'sha256':None,'ready':False,'issues':[f'文件读取失败：{type(error).__name__}'],'filled_label_cells':None},None,None
    try:
        wb = load_workbook(io.BytesIO(data),data_only=False)
    except Exception as error:
        return {'coder':coder,'file':str(path),'sha256':sha_bytes(data),'ready':False,'issues':[f'工作簿解析失败：{type(error).__name__}；收到的字节仍可封存'],'filled_label_cells':None},None,data
    issues = []
    if wb.sheetnames != SHEETS:
        return {'coder':coder,'file':str(path),'sha256':sha_bytes(data),'ready':False,'issues':['页签不匹配'],'filled_label_cells':None},None,data
    if canonical_sha(protected_content(wb)) != baseline['workbooks'][coder]['protected_content_sha256']:
        issues.append('原文、材料编号、语境、手册或固定说明发生内容变化；需核查后另存补材记录。')
    if any(s.sheet_state != 'visible' for s in wb):
        issues.append('存在隐藏页签。')
    if any(cell.data_type == 'f' for s in wb for row in s for cell in row):
        issues.append('填写文件含公式；不得以公式生成或替代实际判定。')
    sheet = wb['填写']
    ids = [sheet.cell(r,1).value for r in range(6,46)]
    if len(set(ids)) != 40 or set(ids) != IDS:
        issues.append('材料编号有缺失、重复、增加或改变。')
    labels = {}
    invalid_labels = []
    missing_reasons = []
    for row in range(6,46):
        material_id = sheet.cell(row,1).value
        labels[material_id] = {}
        for offset, dim in enumerate(DIMS):
            cell = sheet.cell(row,3+offset)
            value = cell.value
            valid = type(value) in (int,float) and math.isfinite(value) and value in (0,1) and cell.data_type != 'f'
            if not valid:
                invalid_labels.append({'material_id':material_id,'dimension':dim,'cell':cell.coordinate,'value':scalar(value),'data_type':cell.data_type})
            else:
                labels[material_id][dim] = int(value)
            reason = sheet.cell(row,8+offset).value
            if not isinstance(reason,str) or not reason.strip():
                missing_reasons.append({'material_id':material_id,'dimension':dim,'cell':sheet.cell(row,8+offset).coordinate})
    if invalid_labels:
        issues.append(f'{len(invalid_labels)}个标签尚未填写或不是严格数值0/1。')
    if missing_reasons:
        issues.append(f'{len(missing_reasons)}个判定理由尚未填写。')
    metadata = {wb['填写说明'].cell(r,1).value:scalar(wb['填写说明'].cell(r,2).value) for r in range(14,21)}
    missing_metadata = [key for key,value in metadata.items() if value is None or not str(value).strip()]
    if missing_metadata:
        issues.append('实际人员和填写过程记录尚未完整。')
    if metadata.get('确认使用本表所附手册') != '是':
        issues.append('尚未确认使用本表所附共同手册。')
    report = {'coder':coder,'file':str(path),'sha256':sha_bytes(data),'ready':not issues,'issues':issues,
              'filled_label_cells':200-len(invalid_labels),'invalid_labels':invalid_labels,
              'filled_reason_cells':200-len(missing_reasons),'missing_reasons':missing_reasons,
              'metadata':metadata,'missing_metadata':missing_metadata,
              'identity_status':'程序仅保存申报记录，实际人员及独立填写过程由研究者据原始记录核实'}
    return report,labels,data

def evaluate(a_path,b_path,baseline):
    reports,labels,raw = {},{},{}
    for coder,path in [('A',a_path),('B',b_path)]:
        reports[coder],labels[coder],raw[coder] = inspect_submission(path,coder,baseline)
    pair_issues = []
    if Path(a_path).resolve() == Path(b_path).resolve() or (reports['A']['sha256'] and reports['A']['sha256'] == reports['B']['sha256']):
        pair_issues.append('两份提交使用同一文件，不能代表两份独立判断。')
    identity_key = '实际填写人姓名或编号'
    if all(r.get('metadata',{}).get(identity_key) for r in reports.values()) and reports['A']['metadata'][identity_key] == reports['B']['metadata'][identity_key]:
        pair_issues.append('两份提交填写人标识相同，须核实真实人员记录。')
    ready = all(r['ready'] for r in reports.values()) and not pair_issues
    result = {'status':'可生成开发校准反馈' if ready else '尚未满足回收要求',
              'purpose':'40条有目的开发材料；不作为正式独立信度、总体推断或机器效度',
              'submissions':reports,'pair_issues':pair_issues,'new_kappa_generated':ready,
              'inference_intervals':'固定开发材料只作描述性反馈，不生成总体置信区间',
              'independence_status':'实际独立填写、讨论及接触答案的情况依填写记录另核；程序不以κ认证人员或独立程序',
              'content_protection_scope':'核对单元格值和材料编号对应的原文、语境、手册，不以该摘要证明格式或文件关系完全未改变',
              'baseline_manual_sha256':baseline['manual_sha256'],'project_target':0.80}
    result['calculation_code_sha256'] = sha_bytes(Path(__file__).read_bytes())
    if ready:
        result['dimensions'] = {dim:kappa_feedback([(labels['A'][rid][dim],labels['B'][rid][dim]) for rid in sorted(IDS)]) for dim in DIMS}
        result['disagreements'] = [{'material_id':rid,'dimension':dim,'A':labels['A'][rid][dim],'B':labels['B'][rid][dim]} for rid in sorted(IDS) for dim in DIMS if labels['A'][rid][dim] != labels['B'][rid][dim]]
    return result,raw,labels

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--a',required=True,type=Path)
    parser.add_argument('--b',required=True,type=Path)
    parser.add_argument('--report',type=Path,help='只读核查报告；不覆盖既有文件')
    parser.add_argument('--receive-dir',type=Path,help='封存收到的原文件及结果；目录必须尚不存在，即使未完成也保留原始提交')
    args = parser.parse_args()
    baseline = json.loads((HERE/'回收基线.json').read_text())
    result,raw,labels = evaluate(args.a,args.b,baseline)
    if args.receive_dir:
        directory = args.receive_dir.resolve()
        directory.mkdir(parents=False,exist_ok=False)
        result['snapshots'] = {}
        for coder in ['A','B']:
            if raw[coder] is None:
                result['snapshots'][coder] = None
                continue
            snapshot = directory/f'编码员{coder}-收到的原文件.xlsx'
            with snapshot.open('xb') as f:
                f.write(raw[coder])
            snapshot_sha = sha_bytes(snapshot.read_bytes())
            if snapshot_sha != result['submissions'][coder]['sha256']:
                raise ValueError('回收快照与收到的原文件字节不符。')
            result['snapshots'][coder] = {'file':str(snapshot),'sha256':snapshot_sha,'exact_received_bytes':True}
        result['received_at_utc'] = datetime.now(timezone.utc).isoformat()
        (directory/'回收核查与开发反馈.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
        if result['new_kappa_generated']:
            (directory/'收到的五维标签.json').write_text(json.dumps(labels,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    if args.report:
        with args.report.open('x') as f:
            json.dump(result,f,ensure_ascii=False,indent=2,allow_nan=False)
            f.write('\n')
    print(json.dumps({'status':result['status'],'new_kappa_generated':result['new_kappa_generated'],
                      'filled_labels':{c:r['filled_label_cells'] for c,r in result['submissions'].items()}},ensure_ascii=False))
    return 0 if result['new_kappa_generated'] else 2

if __name__ == '__main__':
    raise SystemExit(main())
