from pathlib import Path
import json
import hashlib
import zipfile
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
ROOT = OUT.parents[3]
payload = json.loads((HERE / '无答案校准内容.json').read_text())
source_record = json.loads((HERE / '校准案例选择与来源记录.json').read_text())
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
workbooks = {}
reports = {}
for coder in ['A', 'B']:
    path = OUT / f'校准试填-{coder}.xlsx'
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ['填写','阅读材料','编码手册','填写说明']
    assert all(s.sheet_state == 'visible' for s in wb)
    labels = [c.value for row in wb['填写'].iter_rows(min_row=6,max_row=45,min_col=3,max_col=7) for c in row]
    reasons = [c.value for row in wb['填写'].iter_rows(min_row=6,max_row=45,min_col=8,max_col=12) for c in row]
    assert len(labels) == len(reasons) == 200
    assert all(v is None for v in labels+reasons)
    assert all(wb['填写'].cell(r,13).value is None for r in range(6,46))
    assert all(wb['填写说明'].cell(r,2).value is None for r in range(14,21))
    assert wb['填写'].freeze_panes == 'B6'
    validation = list(wb['填写'].data_validations.dataValidation)
    assert any(v.type == 'whole' and str(v.sqref) == 'C6:G45' and v.formula1 == '0' and v.formula2 == '1' for v in validation)
    assert all(c.data_type != 'f' for s in wb for row in s for c in row)
    with zipfile.ZipFile(path) as z:
        assert not any('externalLinks/' in name for name in z.namelist())
    reading = list(wb['阅读材料'].iter_rows(min_row=6,values_only=True))
    for row_index, unit in enumerate(payload['units'],6):
        assert wb['填写'].cell(row_index,1).value == unit['material_id']
        restored = ''.join(r[2] for r in reading if r[0] == unit['material_id'] and r[1].startswith('当前'))
        assert restored == unit['current_text'], unit['material_id']
        for idx, context in enumerate(unit['contexts']):
            if context == unit['current_text']:
                continue
            prefix = f"{unit['context_descriptors'][idx]['kind']}{idx+1} "
            restored_context = ''.join(r[2] for r in reading if r[0] == unit['material_id'] and r[1].startswith(prefix))
            assert restored_context == context, (unit['material_id'],idx)
            descriptor = unit['context_descriptors'][idx]
            assert all(descriptor['source_name'] in r[3] and descriptor['source_locator'] in r[3] for r in reading if r[0] == unit['material_id'] and r[1].startswith(prefix))
    line_parts = {}
    for line_no, part_no, text in wb['编码手册'].iter_rows(min_row=6,values_only=True):
        line_parts.setdefault(line_no, []).append((part_no,text or ''))
    restored_manual = '\n'.join(''.join(text for _,text in sorted(line_parts[i])) for i in sorted(line_parts))
    assert restored_manual == payload['manual']
    for s in wb:
        assert not any(isinstance(c.value,str) and ('RECODE25-' in c.value or '[object Object]' in c.value) for row in s for c in row)
    workbooks[coder] = wb
    reports[coder] = {'file':str(path),'sha256':sha(path),'unit_count':40,'blank_labels':200,'blank_reasons':200,
                      'current_text_exact':True,'contexts_exact':True,'manual_exact':True,
                      'manual_sha256':payload['manual_sha256'],'no_hidden_sheets':True,'no_formulas':True,
                      'no_external_links':True,'labels_validation':'0/1整数','freeze_panes':'B6'}

for name in ['填写','阅读材料','编码手册','填写说明']:
    a, b = workbooks['A'][name], workbooks['B'][name]
    assert a.max_row == b.max_row and a.max_column == b.max_column
    for row in a:
        for cell in row:
            if name in ['填写','填写说明'] and cell.coordinate == 'A2':
                continue
            assert cell.value == b[cell.coordinate].value, (name,cell.coordinate)
for original, expected in source_record['bindings'].items():
    assert sha(original) == expected, original
formal = ROOT / 'outputs/编码修订核查-20261008/本轮结果/创新药评价的注意力配置差异与协调路径.docx'
assert sha(formal) == '2335ad123737735ea5ae32234de323d3c7ae1ac0c23136d4d2eb7a2742b15db6'
result = {'status':'通过','workbooks':reports,'same_content_except_coder_title':True,
          'purpose':'候选手册开发校准；不是新的真人独立检验',
          'source_context_review':payload['context_review_meaning'],
          'original_bindings_unchanged':True,'formal_manuscript_unchanged':True,
          'new_human_labels':False,'new_kappa':False,
          'payload_sha256':sha(HERE/'无答案校准内容.json')}
(OUT/'校准表交付核查.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'status':result['status'],'workbooks':2,'same_content':True,'all_400_labels_and_400_reasons_blank':True},ensure_ascii=False))
