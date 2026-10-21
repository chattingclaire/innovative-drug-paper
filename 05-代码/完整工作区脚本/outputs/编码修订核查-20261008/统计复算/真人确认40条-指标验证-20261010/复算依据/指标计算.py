from pathlib import Path
import json, hashlib, shutil, csv, importlib.util, math
from fractions import Fraction
from collections import Counter
import openpyxl

ROOT=Path('/PROJECT_ROOT')
CAL=ROOT/'outputs/编码修订核查-20261008/人工参照复核/手册校准准备-20261010'
WORK=CAL/'代理试填-20261010/复核反馈处理-20261010'
OUT=ROOT/'outputs/编码修订核查-20261008/统计复算/真人确认40条-指标验证-20261010'
OUT.mkdir(parents=True,exist_ok=True)
ARCH=OUT/'确认参照存档'; ARCH.mkdir(exist_ok=True)
DIMS=['S','E','Ec','I','F']; NAMES=dict(zip(DIMS,['安全性','有效性','经济性','创新性','公平性']))
def read(p): return json.loads(Path(p).read_text())
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,d): Path(p).write_text(json.dumps(d,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def jsonl(p,rows): Path(p).write_text(''.join(json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n' for r in rows))
def ratio(n,d): return float(Fraction(n,d)) if d else None
rule=read('/tmp/validation-rule-40.json');llm=read('/tmp/validation-llm-40.json')
materials=read(CAL/'材料整理记录/无答案校准内容.json')
mu={x['material_id']:x for x in materials['units']}
rp={r['material_id']:r for r in rule['rows']};lp={r['material_id']:r for r in llm['records']}
assert len(rp)==len(lp)==len(mu)==40
for binding in list(rule['input_bindings'].values())+list(llm['source_bindings'].values()):
    assert sha(binding['path'])==binding['sha256'],binding['path']
books={}; wb_bindings={}
for arm in ['A','B']:
    path=WORK/f'校准复核工作表-{arm}.xlsx'
    books[arm]=openpyxl.load_workbook(path,data_only=True)
    dest=ARCH/path.name
    if dest.exists(): assert sha(dest)==sha(path)
    else: shutil.copy2(path,dest)
    wb_bindings[arm]={'path':str(path),'sha256':sha(path),'archive':str(dest),'reviewer':{'A':'于小迪','B':'陈海婷'}[arm]}
references=[];longrows=[];differences=[]
for rownum in range(6,46):
    wa=books['A']['填写'];wb=books['B']['填写'];mid=wa.cell(rownum,1).value
    assert mid==wb.cell(rownum,1).value and mid in mu and mid in rp and mid in lp
    al=[wa.cell(rownum,c).value for c in range(3,8)]
    bl=[wb.cell(rownum,c).value for c in range(3,8)]
    assert al==bl and all(type(v) is int and v in [0,1] for v in al)
    assert wa.cell(rownum,2).value==wb.cell(rownum,2).value
    for arm in ['A','B']:
        reading=''.join(r[2] for r in books[arm]['阅读材料'].values if r[0]==mid and str(r[1]).startswith('当前'))
        assert reading==mu[mid]['current_text']==rp[mid]['current_text'],(arm,mid)
    assert rp[mid]['global_id']==lp[mid]['global_id']
    assert lp[mid]['match_status']['target_text_exact_match'] and rp[mid]['rule_input_text_exact_match']
    reasons={d:books['A']['复核记录'].cell(6+(rownum-6)*5+i,7).value for i,d in enumerate(DIMS)}
    ref={'material_id':mid,'global_id':rp[mid]['global_id'],'department':mu[mid]['department'],
         'current_text':mu[mid]['current_text'],'labels':dict(zip(DIMS,al)),'reasons':reasons,
         'reason_source':'A表“复核记录”G列；两表标签一致，理由可有不同表述',
         'a_cell_range':f'填写!C{rownum}:G{rownum}','b_cell_range':f'填写!C{rownum}:G{rownum}',
         'confirmation_basis':'用户报告两位真人复核结果与当前复核结果一致',
         'manual_sha256':rule['calibration_manual_sha256']}
    references.append(ref)
    context_same=lp[mid]['match_status']['context_text_sequence_exact_match']
    all_found=lp[mid]['match_status']['calibration_context_fragments_all_verbatim_present_in_machine_context_text']
    context_note='语境文本及顺序一致' if context_same else ('校准片段均可逐字找到，顺序或其他语境不同' if all_found else '含新增或改写的校准语境片段')
    for d in DIMS:
        for pred in [rp[mid]['labels'][d],lp[mid]['labels'][d]]: assert type(pred) is int and pred in [0,1]
        item={'material_id':mid,'global_id':ref['global_id'],'department':ref['department'],'dimension':d,'dimension_name':NAMES[d],
              'reference':ref['labels'][d],'rule':rp[mid]['labels'][d],'llm':lp[mid]['labels'][d],
              'reference_reason':reasons[d],'current_text':ref['current_text'],'context_comparison':context_note,
              'source_name':mu[mid]['source_name'],'source_locator':mu[mid]['source_locator']}
        longrows.append(item)
        for method in ['rule','llm']:
            if item[method]!=item['reference']:
                differences.append(dict(item,method=method,prediction=item[method],difference_type='机器1／参照0' if item[method]==1 else '机器0／参照1'))
assert len({r['material_id'] for r in references})==40

corepath=ROOT/'outputs/coding-revision-20261007-r3/statistics/statistical_core_r3.py'
spec=importlib.util.spec_from_file_location('statistical_core',corepath);core=importlib.util.module_from_spec(spec);spec.loader.exec_module(core)
metric_rows=[];summaries=[]
def metrics(method,dim=None):
    rows=[r for r in longrows if dim is None or r['dimension']==dim]
    tp=sum(r[method]==r['reference']==1 for r in rows)
    fp=sum(r[method]==1 and r['reference']==0 for r in rows)
    fn=sum(r[method]==0 and r['reference']==1 for r in rows)
    tn=sum(r[method]==r['reference']==0 for r in rows)
    n=len(rows);pe=Fraction((tp+fp)*(tp+fn)+(tn+fn)*(tn+fp),n*n);po=Fraction(tp+tn,n)
    kappa=float((po-pe)/(1-pe)) if pe!=1 else None
    result={'method':method,'dimension':dim,'n':n,'reference_positive':tp+fn,'prediction_positive':tp+fp,
            'tp':tp,'fp':fp,'fn':fn,'tn':tn,'agreement':float(po),'kappa':kappa,
            'precision':ratio(tp,tp+fp),'recall':ratio(tp,tp+fn),'f1':ratio(2*tp,2*tp+fp+fn)}
    cross=core.metrics([tp,fp,fn,tn])[0]
    for index,key in enumerate(['kappa','agreement','precision','recall','f1']):
        assert (result[key] is None and math.isnan(cross[index])) or abs(result[key]-cross[index])<1e-12,(result,key,cross)
    return result
for method in ['rule','llm']:
    subset=[metrics(method,d) for d in DIMS];metric_rows.extend(subset)
    pooled=metrics(method)
    exact=sum(all((rp if method=='rule' else lp)[r['material_id']]['labels'][d]==r['labels'][d] for d in DIMS) for r in references)
    summaries.append({'method':method,'n_units':40,'n_labels':200,'label_matches':pooled['tp']+pooled['tn'],
                      'label_agreement':pooled['agreement'],'exact_unit_matches':exact,'exact_unit_agreement':exact/40,
                      'micro_precision':pooled['precision'],'micro_recall':pooled['recall'],'micro_f1':pooled['f1'],
                      'macro_f1':sum(x['f1'] for x in subset)/5,'n_differences':pooled['fp']+pooled['fn'],
                      'pooled_counts':{k:pooled[k] for k in ['tp','fp','fn','tn']}})
manifest={'title':'40条确认参照对既有机器结果的指标验证','scope':'有目的选取的手册开发校准集，40个材料、200个维度标签；结果为固定样本内跨版本诊断',
          'confirmation_reported_on':'2026-10-10','confirmation_message':'两位复核的与目前复核结果无异 可以按照这个作为复核结论去做指标验证',
          'reviewers':{'A':'于小迪','B':'陈海婷'},'actual_individual_review_dates':None,
          'human_agreement_interpretation':'根据用户报告，两位对已完成复核的结论均确认一致；不作为独立盲法首轮编码信度',
          'reference_status':'user_confirmed_consensus','workbooks':wb_bindings,'reference_n_units':40,'reference_n_labels':200,
          'reference_reason_source':'A表“复核记录”G列；保留A、B原表的各自理由，统计只使用一致的0/1标签',
          'positive_by_dimension':{d:sum(r['labels'][d] for r in references) for d in DIMS},
          'department_counts':dict(Counter(r['department'] for r in references)),
          'machine_manual_sha256':rule['machine_manual_sha256'],'reference_manual_sha256':rule['calibration_manual_sha256'],
          'same_manual':False,'same_context_sequence_count':2,'all_reference_context_fragments_verbatim_found_count':20,
          'target_text_match_count':40,'target_source_match_count':40,
          'context_interpretation':'片段未逐字命中不等于语义未覆盖；不同手册及语境下的差异不能全部归因于识别错误。',
          'new_model_calls':0,'prediction_labels_modified':False,'source_workbooks_modified':False,
          'sampling_weights':None,'confidence_intervals':None,'population_generalization':False,
          'calculation':'每维采用40条未加权二元混淆矩阵；机器预测为正向，确认结论为参照；分母为0时指标未定义。微平均合并200个标签，宏F1为五维F1平均。',
          'metric_crosscheck':{'path':str(corepath),'sha256':sha(corepath),'result':'全部与既有统计核心一致，绝对误差小于1e-12'},
          'rule_bindings':rule['input_bindings'],'llm_bindings':llm['source_bindings']}
jsonl(OUT/'确认参照40条.jsonl',references)
dump(OUT/'确认与版本记录.json',manifest)
dump(OUT/'指标结果.json',{'metrics':metric_rows,'summary':summaries})
jsonl(OUT/'逐维对照200项.jsonl',longrows)
jsonl(OUT/'差异清单.jsonl',differences)
for name,data in [('各维指标.csv',metric_rows),('差异清单.csv',differences)]:
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(data[0]));writer.writeheader();writer.writerows(data)
shutil.copy2('/tmp/validation-rule-40.json',OUT/'规则预测及来源核查.json')
shutil.copy2('/tmp/validation-llm-40.json',OUT/'大模型预测及来源核查.json')
shutil.copy2(CAL/'共用编码手册.md',ARCH/'共用编码手册.md')
shutil.copy2(CAL/'材料整理记录/无答案校准内容.json',ARCH/'无答案校准内容.json')
dump('/tmp/validation40-workbook-data.json',{'out':str(OUT),'metrics':metric_rows,'summaries':summaries,'references':references,'data':longrows,'differences':differences,'manifest':manifest})
print(json.dumps({'output':str(OUT),'positive_by_dimension':manifest['positive_by_dimension'],'departments':manifest['department_counts'],'metrics':metric_rows,'summary':summaries},ensure_ascii=False,indent=2))
