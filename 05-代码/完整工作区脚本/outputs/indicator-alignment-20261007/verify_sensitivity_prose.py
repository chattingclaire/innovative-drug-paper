"""Read-only numeric and descriptive sensitivity audit of the aligned DOCX."""
from pathlib import Path
from docx import Document
from fractions import Fraction
import json
import hashlib

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
PATH = OUT / '论文修订稿-指标统一版-待补真人验证.docx'
STAT = ROOT / 'outputs/attention-share-audit-20261007/注意力占比-原定义重算.json'
CI = OUT / 'attention_share_bounds.json'
DIMS = ('S','E','Ec','I','F')
GROUPS = {'药品监管':'NMPA','医保':'NHSA','卫生健康':'NHC','中央综合':'GOV','总体':'ALL'}

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def number(value):
    return Fraction(value['exact_fraction']) if isinstance(value,dict) and 'exact_fraction' in value else Fraction(value)

def main():
    before = sha(PATH)
    doc = Document(PATH)
    data = json.loads(STAT.read_text())
    bounds = json.loads(CI.read_text())
    checks, errors, notes, dominance = [], [], [], []
    def check(name, actual, expected, location, scope):
        row = dict(name=name, actual=actual, expected=expected, location=location, scope=scope, passed=actual==expected)
        checks.append(row)
        if not row['passed']:
            errors.append(row)
    for ti, table_name in [(2,'表3'),(4,'表5')]:
        table = doc.tables[ti]
        check('data_row_count',len(table.rows)-1,10,table_name,'shape')
        for ri, row in enumerate(table.rows[1:],1):
            cells = [c.text.strip() for c in row.cells]
            group = GROUPS[cells[0]]
            method = cells[1]+'_AI_reference' if ti==2 else 'rules_v28_2' if cells[1]=='规则v28.2' else 'llm_gpt56_sol_v13_original'
            source = data['primary'][method][group]
            denominator = number(source['weighted_five_dimension_known_hit_sum'])
            expected = f'{float(denominator):.2f}' if ti==2 else str(int(denominator))
            check('denominator_T',cells[2],expected,f'{table_name}/row{ri}',method+'/'+group)
            for col, dimension in enumerate(DIMS,3):
                expected = f"{float(number(source['dimensions'][dimension]['attention_share']))*100:.2f}"
                check('share_'+dimension,cells[col],expected,f'{table_name}/row{ri}/col{col}',method+'/'+group)
            check('T_is_sum_of_five_HT_counts',str(denominator),str(sum(number(source['dimensions'][x]['weighted_known_hit_count']) for x in DIMS)),f'{table_name}/row{ri}',method+'/'+group)
            check('exact_five_shares_sum',str(sum(number(source['dimensions'][x]['attention_share']) for x in DIMS)),'1',f'{table_name}/row{ri}',method+'/'+group)
    for method, primary in data['primary'].items():
        sens = data['sensitivities'][method]
        scenarios = {'primary':primary}
        scenarios.update({'boundary:'+k:v for k,v in sens['definition_boundaries'].items()})
        scenarios.update({k:sens[k] for k in ('policy_file_equal_weight','policy_series_equal_weight','latest_member_of_annual_series_only','exclude_all_joint_publications')})
        scenarios.update({'leave_out:'+x['discarded_file']:x['distribution'] for x in sens['leave_one_policy_file_out']})
        for scenario, result in scenarios.items():
            for group, target in [('NMPA','S'),('NHSA','Ec')]:
                shares = {x:number(result[group]['dimensions'][x]['attention_share']) for x in DIMS}
                maxima = [x for x,v in shares.items() if v==max(shares.values())]
                dominance.append(dict(method=method,scenario=scenario,group=group,claimed_dimension=target,actual_maximum_dimensions=maxima,strict_unique_maximum=maxima==[target]))
                check('dominant_dimension',','.join(maxima),target,scenario,method+'/'+group)
        for scenario, result in sens['definition_direction_changes'].items():
            check('boundary_comparison_count',len(result['comparisons']),15,'paragraph74/80/'+scenario,method)
            check('boundary_reversal_count',result['reversal_count'],0,'paragraph74/80/'+scenario,method)
            check('boundary_direction_change_count',result['changed_direction_count'],0,'paragraph74/80/'+scenario,method)
        for scenario in ('policy_file_equal_weight','policy_series_equal_weight','latest_member_of_annual_series_only'):
            result = sens['scenario_direction_changes'][scenario]
            check('organization_comparison_count',len(result['comparisons']),15,'paragraph80/'+scenario,method)
            check('organization_reversal_count',result['reversal_count'],0,'paragraph80/'+scenario,method)
            check('organization_direction_change_count',result['changed_direction_count'],0,'paragraph80/'+scenario,method)
        expected_joint, expected_loo = {'rules_v28_2':(1,0),'llm_gpt56_sol_v13_original':(2,3),'A_AI_reference':(0,1),'C_AI_reference':(0,1)}[method]
        check('joint_exclusion_reversals',sens['scenario_direction_changes']['exclude_all_joint_publications']['reversal_count'],expected_joint,'paragraph80',method)
        reversed_joint = [r for r in sens['scenario_direction_changes']['exclude_all_joint_publications']['comparisons'] if r['reversal']]
        expected_dimensions = ['I'] if method=='rules_v28_2' else ['F','I'] if method=='llm_gpt56_sol_v13_original' else []
        check('joint_exclusion_reversed_dimensions',sorted(r['dimension'] for r in reversed_joint),expected_dimensions,'paragraph80',method)
        for row in reversed_joint:
            check('joint_exclusion_reversed_pair',row['pair'],['NHSA','NHC'],'paragraph80',method)
        check('leave_one_file_reversal_cases',sens['leave_one_policy_file_summary']['reversal_case_count'],expected_loo,'paragraph80',method)
        if method.endswith('AI_reference'):
            for row in sens['leave_one_policy_file_summary']['reversed_pairs']:
                check('AC_LOO_reversal_dimension',row['dimension'],'I','paragraph80',method)
                check('AC_LOO_reversal_pair',row['pair'],['NMPA','NHSA'],'paragraph80',method)
            check('observed_file_count',data['method_file_inventory'][method]['observed_file_count'],54,'paragraph80',method)
    check('machine_same_point_direction_count',data['machine_attention_same_direction_count'],14,'paragraph76','15 professional comparisons')
    check('machine_same_direction_all_pending_completions',data['machine_attention_same_direction_all_pending_completions_count'],14,'paragraph76','15 professional comparisons')
    llm_path = ROOT/'outputs/autocode-v25-20260927/llm_runs/gpt56-sol-v13-native-1/all_predictions.jsonl'
    llm = [json.loads(x) for x in llm_path.read_text().splitlines()]
    check('old_LLM_all_five_pending_units',sum(all(v=='待定' for v in r['labels'].values()) for r in llm),49,'paragraph50/79','old sealed LLM')
    check('old_LLM_F_positive_missing_subtype',sum(r['labels']['F']==1 and not r.get('fairness_subtypes') for r in llm),3,'paragraph80','old sealed LLM')
    for key in ('known_single_coder_directions','completed_single_coder_directions','common_department_directions'):
        check('simultaneous_summary_'+key,bounds['summary'][key],0,'paragraph65/table4','consistency with parent transformed-region JSON; not new proof of the math')
    for row in doc.tables[3].rows[1:]:
        for cell in row.cells[1:]:
            check('table4_zero_count',cell.text.strip(),'0','table4','transformed comparison summary')
    phrases = {47:['分母为同一来源组五维编码句次合计','合计为100%','五维合计为零时占比无定义，不填0','维度内容覆盖比例'],
        56:['先估计每维编码为1的句次，再除以五维估计句次合计','总体先合并四组句次后求比，不平均各组占比','固定单元数比例的区间不能直接用于注意力占比'],
        59:['五维加权编码句次'],72:['分母为该组五维加权编码句次合计'],79:['五维编码句次','并非v29重跑','不将待定改判为0']}
    for index, targets in phrases.items():
        for phrase in targets:
            check('denominator_scope_phrase',phrase in doc.paragraphs[index].text,True,'paragraph'+str(index),phrase)
    prose_numbers = {65:['45.71%','46.65%','30.93%','30.04%','69.40%','64.01%','28.89%','28.24%','30.77%','34.04%'],
        73:['7.31%','7.51%','3.40%','3.49%','1.82%','1.61%'],76:['48.63%','44.41%','33.89%','30.94%','62.94%','51.90%','30.07%','35.32%','31.65%','21.75%','34.51%']}
    for index, targets in prose_numbers.items():
        for target in targets:
            check('prose_value',target in doc.paragraphs[index].text,True,'paragraph'+str(index),target)
    for method in ('A_AI_reference','C_AI_reference'):
        for group, dimensions, rounded in [('NMPA',['S','E'],77),('NHC',['S','F'],60 if method.startswith('A') else 62)]:
            actual = sum(number(data['primary'][method][group]['dimensions'][x]['attention_share']) for x in dimensions)*100
            check('abstract_rounded_dimension_sum',round(float(actual)),rounded,'paragraph1/65/104',method+'/'+group)
    if '均为各自组内最高维度' in doc.paragraphs[76].text:
        notes.append({'priority':'wording_clarity','paragraph_index_zero_based':76,
            'text':doc.paragraphs[76].text,
            'reason':'该句先列药监安全性、有效性，末尾“均”可能误读为有效性亦最高，而有效性实际第二。数值正确；建议明确最高仅指药监安全性和医保经济性。'})
    after = sha(PATH)
    result = {'status':'FAIL' if errors else 'PASS_WITH_WORDING_CLARITY_NOTE' if notes else 'PASS',
        'document':str(PATH),'document_sha256_before':before,'document_sha256_after':after,'document_unchanged_during_audit':before==after,
        'source_sha256':{str(STAT):sha(STAT),str(CI):sha(CI),str(llm_path):sha(llm_path)},
        'table3_share_checks':50,'table5_share_checks':50,'denominator_T_checks':20,
        'dominance_scenario_group_checks':len(dominance),'numeric_or_consistency_checks':len(checks),
        'failed_checks':errors,'wording_observations':notes,'all_checks':checks,'all_dominance_scenarios':dominance,
        'scope_limits':['只读正文及表格，未修改DOCX。','未进行本稿视觉布局QA。',
            '同时区间0项结论与根代理新变换JSON一致；本轮未另证明同时区域数学有效性。',
            '组内最高仅指已确定编码比率点估计及已列敏感性情景；不等于待定任意补全后必然最高或置信结论。']}
    (OUT/'敏感性文字终核.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps({k:result[k] for k in ('status','table3_share_checks','table5_share_checks','denominator_T_checks','dominance_scenario_group_checks','numeric_or_consistency_checks','failed_checks','document_unchanged_during_audit')},ensure_ascii=False))
    print('wording_notes',len(notes))

if __name__=='__main__':
    main()
