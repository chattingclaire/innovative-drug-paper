"""Collect a final DOCX audit snapshot; never edit the manuscript.

Run only after the parent has confirmed the file is ready.
The emitted candidates require semantic review; they are not automatically errors.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import argparse
import hashlib
import json
import re

HERE=Path(__file__).resolve().parent
DEFAULT=HERE/'论文修订稿-指标统一版-待补真人验证.docx'
BASELINE=HERE.parent/'manuscript-completion-20261006/论文完整修订稿-待补真人验证.docx'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm_reference(text):return re.sub(r'\s+','',re.sub(r'^\[\d+\]\s*','',text))

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--ready',action='store_true');ap.add_argument('--path',type=Path,default=DEFAULT)
    a=ap.parse_args()
    if not a.ready:raise SystemExit('No inspection performed: wait for confirmed ready status, then use --ready.')
    p=a.path.resolve();before=sha(p);doc=Document(p)
    paragraphs=[{'index':i,'text':x.text} for i,x in enumerate(doc.paragraphs)]
    by_element={id(x._p):i for i,x in enumerate(doc.paragraphs)}
    table_elements={id(x._tbl):i for i,x in enumerate(doc.tables)}
    blocks=[];pi=ti=0
    for e in doc._element.body:
        if e.tag==qn('w:p'):
            blocks.append({'kind':'paragraph','index':pi,'text':doc.paragraphs[pi].text});pi+=1
        elif e.tag==qn('w:tbl'):
            blocks.append({'kind':'table','index':ti});ti+=1
    tables=[]
    for bi,b in enumerate(blocks):
        if b['kind']!='table':continue
        before_text=[v['text'] for v in blocks[max(0,bi-4):bi] if v['kind']=='paragraph' and v['text'].strip()]
        after_text=[v['text'] for v in blocks[bi+1:bi+5] if v['kind']=='paragraph' and v['text'].strip()]
        t=doc.tables[b['index']]
        tables.append({'index':b['index'],'preceding_text':before_text,'following_text':after_text,'rows':[[c.text for c in r.cells] for r in t.rows]})
    patterns={
      'fixed_N_near_attention':r'注意力占比.{0,180}(?:分母.{0,30}(?:全部|总体|固定|N)|固定N)|(?:分母.{0,30}(?:全部分析单元|来源组全部|来源组N)).{0,180}注意力占比',
      'old_three_directions':r'(?:3项|三项).{0,45}(?:排除零|排除0|显著|方向)|(?:排除零|排除0|显著).{0,45}(?:3项|三项)',
      'fifteen_direction_claim':r'15\s*[/／]\s*15|15项.{0,80}(?:方向|同向|一致|稳定)|(?:方向|同向|一致|稳定).{0,80}15项',
      'positive_dimension_term':r'阳性维度|维度阳性|阳性点估计|阳性比例',
      'human_coding_provenance':r'真人|人工验证|人工编码|人类编码|金标准|A/C|编码员|两位AI|两名AI|agent|AI参照',
      'denominator_contract':r'分母|五维.{0,20}(?:总和|合计|总量|总句次)|命中.{0,20}(?:合计|总和)|覆盖比例|注意力占比',
      'effect_or_weight':r'政策效果|政策效应|实际评审权重|因果|偏好|主导',
    }
    candidates={k:[x for x in paragraphs if re.search(pat,x['text'])] for k,pat in patterns.items()}
    refs=[];body=[]
    for x in paragraphs:
        m=re.match(r'^\[(\d+)\]\s*(.+)',x['text'])
        if m:refs.append({'number':int(m.group(1)),'paragraph_index':x['index'],'text':x['text']})
        else:body.append(x)
    numbers=[x['number'] for x in refs];available=set(numbers)
    cites=[]
    for x in body:
        for n in re.findall(r'\[(\d+)\]',x['text']):cites.append({'number':int(n),'paragraph_index':x['index']})
    first_order=list(dict.fromkeys(x['number'] for x in cites))
    bdoc=Document(BASELINE)
    oldrefs={int(m.group(1)):norm_reference(x.text) for x in bdoc.paragraphs if (m:=re.match(r'^\[(\d+)\]\s*',x.text))}
    changed_refs=[x for x in refs if oldrefs.get(x['number'])!=norm_reference(x['text'])]
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    images=[]
    for i,x in enumerate(doc.paragraphs):
        for e in x._p.findall('.//a:blip',ns):
            rid=e.attrib.get('{'+ns['r']+'}embed');part=doc.part.related_parts[rid]
            target=HERE/('终稿-'+Path(str(part.partname)).name)
            target.write_bytes(part.blob)
            images.append({'paragraph_index':i,'relationship':rid,'partname':str(part.partname),'extracted_path':str(target)})
    report={'path':str(p),'sha256':before,'read_only_unchanged':sha(p)==before,'paragraphs':paragraphs,'tables':tables,'images':images,
            'reference_checks':{'count':len(refs),'numbers':numbers,'continuous_1_to_51':numbers==list(range(1,52)),
                'body_citation_count':len(cites),'first_citation_order':first_order,'first_citation_order_1_to_51':first_order==list(range(1,52)),
                'unresolved_citations':[x for x in cites if x['number'] not in available],'uncited_references':sorted(available-set(first_order)),
                'changed_from_audited_oct6_by_number':changed_refs,'references':refs},
            'candidate_occurrences_require_semantic_review':candidates,'automatic_check_is_not_final_approval':True}
    (HERE/'终稿只读提取与检索.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    (HERE/'终稿正文与表格-只读提取.txt').write_text('\n'.join(f"P{x['index']:03}: {x['text']}" for x in paragraphs)+'\n\n'+json.dumps(tables,ensure_ascii=False,indent=2))
    print(json.dumps({'paragraphs':len(paragraphs),'tables':len(tables),'images':len(images),'reference_checks':{k:v for k,v in report['reference_checks'].items() if k not in ['references','numbers']},'candidate_counts':{k:len(v) for k,v in candidates.items()},'read_only_unchanged':report['read_only_unchanged']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
