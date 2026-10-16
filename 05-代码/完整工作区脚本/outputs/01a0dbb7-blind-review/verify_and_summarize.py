from pathlib import Path
import json, hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
import openpyxl
from copy import copy

root=Path(__file__).resolve().parent
dims=['S','E','Ec','I','F']
names=['安全性S','有效性E','经济性Ec','创新性I','公平性F']
manifest=json.loads((root/'input_manifest.json').read_text())
results={c:json.loads((root/f'agent-{c}'/f'result-{c}.json').read_text()) for c in ['A','C']}
raw=json.loads((root/'agent-A/待评材料.json').read_text())['rows']
ids=[r['id'] for r in raw]
maps={c:{r['id']:r for r in data['records']} for c,data in results.items()}
qa={'verified_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),'inputs_identical':manifest['same_text_and_manual'],'files':{}}
for c in ['A','C']:
    result=results[c]
    assert result['protocol_version']=='CAL-5D-20260926-v1.1' and result['protocol_acknowledged']
    assert len(result['records'])==18 and set(maps[c])==set(ids)
    source=Path(manifest[c]['source'])
    assert hashlib.sha256(source.read_bytes()).hexdigest()==manifest[c]['sha256']
    sw=openpyxl.load_workbook(source)
    file=root/f'校准盲评-{c}-AI专家.xlsx'
    ow=openpyxl.load_workbook(file)
    assert sw.sheetnames==ow.sheetnames
    allowed={(1,1),(2,1)} | {(r,k) for r in range(6,24) for k in range(3,11)}
    for ss,os in zip(sw,ow):
        assert ss.max_column==os.max_column
        assert os.max_row==(68 if ss.title=='编码手册' else ss.max_row)
        for row in ss:
            for cell in row:
                other=os.cell(cell.row,cell.column)
                if ss.title!='独立试填' or (cell.row,cell.column) not in allowed:
                    assert cell.value==other.value,(c,ss.title,cell.coordinate,'unexpected value change')
                assert copy(cell.font)==copy(other.font),(c,ss.title,cell.coordinate,'font changed')
                assert copy(cell.fill)==copy(other.fill),(c,ss.title,cell.coordinate,'fill changed')
                assert copy(cell.alignment)==copy(other.alignment),(c,ss.title,cell.coordinate,'alignment changed')
                assert cell.number_format==other.number_format
        assert ss.freeze_panes==os.freeze_panes
        assert str(ss.merged_cells)==str(os.merged_cells)
        assert ss.data_validations==os.data_validations
        assert list(ss.tables)==list(os.tables)
    os=ow['独立试填']
    for row,id_ in enumerate(ids,start=6):
        r=maps[c][id_]
        assert all(r['labels'][d] in [0,1,'待定'] for d in dims)
        assert all(isinstance(r[k],str) and r[k].strip() for k in ['evidence','material_issue','rationale'])
        assert [os.cell(row,k).value for k in range(3,8)]==[r['labels'][d] for d in dims]
        assert os.cell(row,8).value==r['evidence']
        assert r['rationale'] in os.cell(row,10).value
        assert os.cell(row,9).value in ['无','上下文不足','断词或乱码','标题或碎片','其他']
    qa['files'][c]={'output':str(file),'output_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'raw_agent_result_sha256':hashlib.sha256((root/f'agent-{c}'/f'result-{c}.json').read_bytes()).hexdigest(),'records':18,'label_cells':90,'original_unchanged':True,'source_text_and_original_manual_preserved':True,'same_execution_addendum_appended':True,'validations_panes_and_existing_cell_styles_preserved':True}

aw=openpyxl.load_workbook(root/'校准盲评-A-AI专家.xlsx')
cw=openpyxl.load_workbook(root/'校准盲评-C-AI专家.xlsx')
assert list(aw['编码手册'].values)==list(cw['编码手册'].values)

def vec(r):return ' / '.join(str(r['labels'][d]) for d in dims)
def safe(t):return str(t).replace('|','／').replace('\n','<br>')
pending={c:[id_ for id_ in ids if '待定' in maps[c][id_]['labels'].values()] for c in maps}
differences=[id_ for id_ in ids if maps['A'][id_]['labels']!=maps['C'][id_]['labels']]
stats=[]
for d,name in zip(dims,names):
    pairs=[(maps['A'][id_]['labels'][d],maps['C'][id_]['labels'][d]) for id_ in ids]
    binary=[(a,b) for a,b in pairs if a in [0,1] and b in [0,1]]
    same=sum(a==b for a,b in binary)
    stats.append({'dimension':d,'name':name,'binary_n':len(binary),'binary_agree_n':same,'binary_agreement':same/len(binary) if binary else None,'pending_pair_n':18-len(binary),'positive_A':sum(a==1 for a,b in pairs),'positive_C':sum(b==1 for a,b in pairs),'pending_A':sum(a=='待定' for a,b in pairs),'pending_C':sum(b=='待定' for a,b in pairs)})
summary={'pending_rows':pending,'different_rows':differences,'same_vector_rows':18-len(differences),'dimension_stats':stats}
qa['summary']=summary
(root/'交付核验.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
lines=['# 两位AI专家agent独立盲评汇总','',
'本次按2026-09-26候选版手册和共同执行说明CAL-5D-20260926-v1.1评审同一组18条边界材料。A与C各自从未继承主对话的空白上下文开始，只收到本人空白表、同版规则及其直接提取文本。用户补充定义一致要求后，两位在正式提交前暂停，先完成15道单独构造的教学题与定义复述，75个维度判断全部一致，再读相同执行说明并确认后逐条重判18条。教学题不属于18条盲评材料，不计信度。',
'','共享的是五维定义、语境要求和问题状态规则；18条的独立判定未共享、未查看旧标签或机器判定。两人原始提交完成后才进行本汇总，未协商盲评答案、未裁定、未改标。',
'','机器口径同步：已将完全相同协议纳入新版LLM系统提示词，并将词典的规范目标绑定到同版协议。现存v4-development-16的识别覆盖与待定处理仍属实现检查事项，旧LLM提示则已查明定义较宽。未执行全量自动或LLM重跑，旧结果不能直接作为本次同口径结果；详见统一口径目录的三种方法定义核对与版本绑定。',
'','A的评审角色为创新药监管科学、药物评价与政策文本分析；C为卫生技术评估、医保政策与定性内容分析。两者均为AI agent，未人为设置不同底层模型，不能当作两位真人专家或相互独立模型的验证。材料本身也已用于规则开发，本次结果只用于手册校准。','',
f"A提交时间：{results['A']['completed_at']}。C提交时间：{results['C']['completed_at']}。",'',
f"18条中，五维标签完全相同的条目为{18-len(differences)}条，存在至少一维不同的为{len(differences)}条（包含0/1与待定的差别）。A含待定{len(pending['A'])}条，C含待定{len(pending['C'])}条。同为待定只表示两人均无法判定，不是实质判断一致。",'',
'下表只对双方均明确填0或1的同维标签计算描述性一致率；涉及待定的配对单列，不补0、不计入该分母。未计算正式信度或据此宣布验证达标。','',
'| 维度 | 可比较0/1配对 | 一致数 | 一致率 | 含待定配对 | A的1 / 待定 | C的1 / 待定 |','|---|---:|---:|---:|---:|---:|---:|']
for s in stats:
    pct=f"{s['binary_agreement']:.1%}" if s['binary_agreement'] is not None else '不适用'
    lines.append(f"| {s['name']} | {s['binary_n']} | {s['binary_agree_n']} | {pct} | {s['pending_pair_n']} | {s['positive_A']} / {s['pending_A']} | {s['positive_C']} / {s['pending_C']} |")
lines+=['','标签顺序统一为 **S / E / Ec / I / F**。','',
'| 编号 | A独立结果 | C独立结果 | 状态 |','|---|---|---|---|']
for id_ in ids:
    a,b=maps['A'][id_],maps['C'][id_]
    status='存在差别' if id_ in differences else '相同（含待定）' if id_ in pending['A'] else '相同'
    lines.append(f'| {id_} | {vec(a)} | {vec(b)} | {status} |')
lines+=['','## 差异与待定条目','', '以下原样摘录独立理由，供后续校准使用，未提供裁定标签。','']
for x in raw:
    id_=x['id']
    if id_ not in set(differences+pending['A']+pending['C']):continue
    lines += [f'### {id_}','',f"原文：{x['text']}",'']
    for c in ['A','C']:
        r=maps[c][id_]
        lines += [f"**{c}：{vec(r)}**。材料问题：{r['material_issue']}。",f"证据：{r['evidence']}",f"理由：{r['rationale']}",'']
lines+=['## 各自提出的手册疑问','']
for c in ['A','C']:
    lines += [f'### 编码员{c}','']
    for issue in results[c].get('manual_issues',[]):
        lines += [f"- **{issue['rule']}**（{'、'.join(issue['affected_ids'])}）：{issue['issue']} 建议：{issue['suggestion']}"]
    lines+=['']
lines += ['## 交付检查','',
'两份工作簿均含18条、每条五维标签、证据片段、材料问题及判定理由。原文、编号和原编码手册逐格核对保留，手册末尾附入相同的v1.1共用执行说明；原始空白文件未被覆盖。原始agent JSON与说明分别保存在agent-A、agent-C目录。为容纳判定理由，扩宽证据与理由列并调整相应行高；其他既有单元格样式、下拉校验及冻结位置保留。若问题栏原述超出模板选项，仅将其映射到既有选项并原样保留在理由栏，未改标签。','',
'建议先核实本轮标出的文本完整性和必要父条款，再讨论手册边界；若修改定义，保留本轮为原始提交，另开新轮次。正式验证需冻结手册并另取独立样本。']
(root/'盲评结果与分歧汇总.md').write_text('\n'.join(lines)+'\n')
print(json.dumps(summary,ensure_ascii=False,indent=2))
