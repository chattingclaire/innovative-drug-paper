from pathlib import Path
import json, hashlib, shutil
import openpyxl

src=Path('/Users/rainbow/Desktop/战略院/创新药论文-完整归档-20260914/12-投稿定稿工作区/outputs/paper-calibration-20260926')
out=Path(__file__).resolve().parent
manifest={}
allrows=[]
manuals=[]
for coder in ['A','C']:
    path=src/f'校准试填-{coder}.xlsx'
    wb=openpyxl.load_workbook(path,data_only=False)
    ws=wb['独立试填']
    rows=[{'id':ws.cell(r,1).value,'text':ws.cell(r,2).value} for r in range(6,24)]
    assert len(rows)==18 and all(x['id'] for x in rows)
    assert all(ws.cell(r,c).value is None for r in range(6,24) for c in range(3,11))
    manual='\n\n'.join(str(row[0]) for row in wb['编码手册'].values if row[0] is not None)
    target=out/f'agent-{coder}'
    target.mkdir(parents=True,exist_ok=True)
    shutil.copy2(path,target/path.name)
    (target/'编码手册.txt').write_text(manual)
    (target/'待评材料.json').write_text(json.dumps({'coder':coder,'instructions':[ws.cell(r,1).value for r in range(2,5)],'rows':rows},ensure_ascii=False,indent=2))
    manifest[coder]={'source':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'manual_sha256':hashlib.sha256(manual.encode()).hexdigest(),'rows':len(rows),'input_labels_empty':True}
    allrows.append(rows)
    manuals.append(manual)
assert allrows[0]==allrows[1]
assert manuals[0]==manuals[1]
manifest['same_text_and_manual']=True
(out/'input_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False,indent=2))
