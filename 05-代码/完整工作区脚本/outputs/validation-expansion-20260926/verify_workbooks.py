import json,hashlib
from pathlib import Path
from openpyxl import load_workbook
from score_blind import validate,DIMS
OUT=Path(__file__).parent
packet=json.loads((OUT/'开发60-共同无答案材料.json').read_text())
stats=json.loads((OUT/'开发60-独立盲评统计.json').read_text())
manifest=json.loads((OUT/'开发60-盲法与版本清单.json').read_text())
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(OUT/'agent-A/待评材料.json')==sha(OUT/'agent-C/待评材料.json')==manifest['packet_sha256']
assert sha(OUT/'agent-A/五维共用编码协议.md')==sha(OUT/'agent-C/五维共用编码协议.md')==manifest['protocol_sha256']
qa={'input_and_protocol_identical':True,'rows_per_coder':60,'files':{},'raw_submissions_unchanged':True}
manuals=[];contexts=[]
for coder in ('A','C'):
    resultpath=OUT/f'agent-{coder}/result-{coder}.json'
    assert sha(resultpath)==stats['input_sha256'][f'agent-{coder}/result-{coder}.json']
    _,idx=validate(resultpath,[r['id'] for r in packet['rows']])
    p=OUT/f'扩充盲评60-{coder}-AI评审.xlsx'; w=load_workbook(p,data_only=False)
    assert w.sheetnames==['独立盲评','原文语境','共同手册']
    s=w['独立盲评']; assert s.max_row==64 and s.freeze_panes=='C5'
    assert 'AI' in s['A1'].value and '60' in s['A1'].value
    for i,r in enumerate(packet['rows'],5):
        q=idx[r['id']]; values=[s.cell(i,j).value for j in range(1,10)]
        assert values[:2]==[r['id'],r['text']]
        assert values[2:7]==[q['labels'][d] for d in DIMS]
        expected='\n'.join(f'{d}：{q["dimension_reasons"][d]}' for d in DIMS)+f'\n证据：{q["evidence"]}\n综合说明：{q["rationale"]}'
        assert values[7:]==[expected,q['material_issue']]
        assert s.row_dimensions[i].height<=409
    dv=list(s.data_validations.dataValidation)
    assert len(dv)==1 and str(dv[0].sqref)=='C5:G64' and '待定' in dv[0].formula1
    for sheet in w:
        for row in sheet:
            for cell in row:
                assert cell.data_type!='e',(coder,sheet.title,cell.coordinate)
    manuals.append(list(w['共同手册'].values));contexts.append(list(w['原文语境'].values))
    qa['files'][coder]={'path':str(p),'sha256':sha(p),'exact_labels_text_reasons':True,'data_validation':True,'freeze_panes':s.freeze_panes,'excel_errors':0}
assert manuals[0]==manuals[1] and contexts[0]==contexts[1]
qa['workbook_manuals_and_contexts_identical']=True
qa['visual_review']='每份独立盲评首部、判定理由和末部已渲染；共同上下文与手册已检查，A/C同版内容逐格一致。'
(OUT/'扩充60-交付核验.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps(qa,ensure_ascii=False))
