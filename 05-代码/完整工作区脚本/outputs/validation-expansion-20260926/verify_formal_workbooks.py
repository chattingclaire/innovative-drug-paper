"""Read-only export QA after complete H600 or R200 workbook exports exist."""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile
import openpyxl

DIMS = ('S', 'E', 'Ec', 'I', 'F')
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def sha(raw):
    return hashlib.sha256(raw).hexdigest()

def compact(value):
    return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, separators=(',', ':'))

def flatten(value, pointer, target):
    if isinstance(value, dict) and value:
        for key, item in value.items():
            flatten(item, pointer + '/' + key.replace('~', '~0').replace('/', '~1'), target)
    elif isinstance(value, list) and value:
        for i, item in enumerate(value):
            flatten(item, pointer + '/' + str(i), target)
    else:
        assert pointer not in target
        target[pointer] = value

def same_cell(actual, expected):
    if expected == '' and actual is None:
        return True
    if isinstance(expected, str) and (expected.startswith('=') or re.match(r'^\d{4}-\d{2}-\d{2}(?:T|$)', expected)) and actual == "'" + expected:
        return True
    return actual == expected and not (type(actual) is bool and type(expected) is int)

def verify(root, coder, statistics, phase='H600'):
    folder = root / phase / f'agent-{coder}'
    prefix = 'formal600' if phase == 'H600' else 'challenge200'
    title = '正式盲评600' if phase == 'H600' else '补充挑战200'
    workbook_path = root / f'{title}-{coder}-AI评审.xlsx'
    checks = json.loads((root / f'{prefix}-{coder}-input-checks.json').read_text())
    n = 600 if phase == 'H600' else statistics['sample_n']
    batches = math.ceil(n / 25)
    assert n > 0 and checks['records'] == n and checks['batches'] == batches
    materials, results, expected_context = [], {}, defaultdict(dict)
    hashes = statistics['archive']['input_snapshot_hashes']
    for batch in range(1, batches + 1):
        names = {kind: f'agent-{coder}/{kind}-{batch:02d}.json' for kind in ('材料', '结果')}
        docs = {}
        for kind, name in names.items():
            raw = (root / phase / name).read_bytes()
            expected_sha = hashes['released' if kind == '材料' else 'submissions'][name]
            assert sha(raw) == expected_sha
            doc = json.loads(raw)
            docs[kind] = doc
            metadata = {k: v for k, v in doc.items() if k != 'rows'}
            prefix = '材料批次' if kind == '材料' else '独立提交批次'
            flatten({'file': name, 'sha256': sha(raw), 'metadata': metadata}, prefix, expected_context[f'批次{batch:02d}'])
        material, result = docs['材料'], docs['结果']
        materials.extend(material['rows'])
        for row in result['rows']:
            assert row['id'] not in results
            results[row['id']] = row
            flatten(row, '原始独立提交', expected_context[row['id']])
        for row in material['rows']:
            flatten(row, '材料', expected_context[row['id']])
    assert len(materials) == len(results) == n
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    assert wb.sheetnames == ['独立盲评', '原文语境', '共同手册']
    # The exporter validly omits optional worksheet dimension metadata. Derive
    # bounds from actual XML cells before checking completeness; do not treat
    # missing cached bounds as missing records.
    for sheet in wb:
        sheet.calculate_dimension(force=True)
    main = wb['独立盲评']
    assert main.max_row == n + 4 and main.max_column == 9
    errors = []
    for material, cells in zip(materials, main.iter_rows(min_row=5, max_row=n + 4)):
        row = results[material['id']]
        reasons = '\n'.join(f'{d}：{row["dimension_reasons"][d]}' for d in DIMS)
        reasons += f'\n证据：{row["evidence"]}\n综合说明：{row["rationale"]}'
        issues = f'材料问题：{compact(row["material_issue"])}\n材料核验：{compact(row["input_assessment"])}'
        issues += f'\nF子类型：{compact(row["fairness_subtypes"])}\nI子类型：{compact(row["innovation_subtypes"])}\nS质量依据：{row["safety_quality_basis"]}'
        expected = [material['id'], material['text'], *[row['labels'][d] for d in DIMS], reasons, issues]
        for cell, value in zip(cells, expected):
            if cell.data_type in ('f', 'e') or not same_cell(cell.value, value):
                errors.append(f'{coder}/独立盲评/{cell.coordinate}')
    context = wb['原文语境']
    assert context.max_row == checks['contextRows'] + 4 and context.max_column == 3
    observed = defaultdict(dict)
    chunk_groups = defaultdict(list)
    for key_cell, pointer_cell, body_cell in context.iter_rows(min_row=5):
        assert all(cell.data_type not in ('f', 'e') for cell in (key_cell, pointer_cell, body_cell))
        match = re.fullmatch(r'(.*) \[(string|number|boolean|null|array|object)\](?: \((\d+)/(\d+)\))?', pointer_cell.value)
        assert match, pointer_cell.value
        pointer, kind, index, total = match.groups()
        body = '' if body_cell.value is None else body_cell.value
        assert isinstance(body, str), body_cell.coordinate
        chunk_groups[(key_cell.value, pointer, kind)].append((int(index or 1), int(total or 1), body))
    for (key, pointer, kind), chunks in chunk_groups.items():
        assert [c[0] for c in chunks] == list(range(1, len(chunks) + 1))
        assert all(c[1] == len(chunks) for c in chunks)
        original = expected_context[key][pointer]
        body = ''
        for _, _, part in chunks:
            # Retain genuine apostrophes; remove an export escape only when the
            # corresponding exact source substring establishes that escape.
            if kind == 'string' and not original[len(body):].startswith(part):
                assert part.startswith("'") and original[len(body):].startswith(part[1:])
                part = part[1:]
            body += part
        observed[key][pointer] = body if kind == 'string' else json.loads(body)
    assert dict(observed) == dict(expected_context), 'Some source/submission fields did not survive display export'
    manual = wb['共同手册']
    protocol = (folder / '五维共用编码协议.md').read_text()
    assembled = ''
    for row in manual.iter_rows(min_row=4):
        part = row[0].value or ''
        if not protocol[len(assembled):].startswith(part):
            assert part.startswith("'") and protocol[len(assembled):].startswith(part[1:])
            part = part[1:]
        assembled += part
    assert assembled == protocol and sha(protocol.encode()) == checks['protocol_sha256']
    wb.close()
    assert not errors, errors[:30]
    native = {}
    with ZipFile(workbook_path) as z:
        for i, name in enumerate(('独立盲评', '原文语境', '共同手册'), 1):
            tree = ET.fromstring(z.read(f'xl/worksheets/sheet{i}.xml'))
            heights = [float(x.attrib['ht']) for x in tree.findall('.//s:row', NS) if 'ht' in x.attrib]
            assert heights and max(heights) <= 409
            pane = tree.find('.//s:pane', NS)
            assert pane is not None and pane.attrib.get('state') in ('frozen', 'frozenSplit')
            native[name] = {'max_row_height': max(heights), 'pane': pane.attrib}
            if i == 1:
                assert tree.find('s:dataValidations', NS) is not None
                assert tree.find('s:conditionalFormatting', NS) is not None
    return {'coder': coder, 'path': str(workbook_path), 'sha256': sha(workbook_path.read_bytes()),
            'records': n, 'labels_compared': n * 5, 'all_main_values_match': True,
            'all_context_leaf_values_reconstructed': True, 'context_leaf_count': len(chunk_groups),
            'protocol_reconstructed_exactly': True, 'formula_or_error_cells_in_data': 0, 'native_layout': native,
            'visual_review': 'separate rendered-image inspection required'}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=Path(__file__).parent)
    parser.add_argument('--phase', choices=('H600','R200'), default='H600')
    parser.add_argument('--statistics', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    assert not args.output.exists(), 'Preserve previous issued QA'
    statistics_path = args.statistics or args.root / ('H600-正式统计.json' if args.phase == 'H600' else 'R200-描述统计.json')
    statistics = json.loads(statistics_path.read_text())
    expected_status = 'completed_frozen_ID_join_not_validation_pass' if args.phase == 'H600' else 'completed_challenge_descriptive_ID_join_not_validation_pass'
    assert statistics['status'] == expected_status
    output = {'status': 'read_only_workbook_data_and_native_layout_QA',
              'phase': args.phase,
              'created_at': datetime.now(timezone.utc).isoformat(),
              'statistics_sha256': sha(statistics_path.read_bytes()),
              'workbooks': [verify(args.root, c, statistics, args.phase) for c in ('A', 'C')]}
    with args.output.open('x') as handle:
        json.dump(output, handle, ensure_ascii=False, indent=2)
        handle.write('\n')
    print(json.dumps({'status': output['status'], 'phase': args.phase, 'coders': ['A', 'C'], 'records_each': output['workbooks'][0]['records']}, ensure_ascii=False))

if __name__ == '__main__':
    main()
