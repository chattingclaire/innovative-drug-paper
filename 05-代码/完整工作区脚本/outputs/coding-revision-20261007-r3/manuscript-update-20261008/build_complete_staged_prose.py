"""Stage the complete local prose revision without editing a DOCX or labels.

All unresolved measurement fields are explicit. A future document builder must
resolve them from one adopted, SHA-bound output package, never from old prose.
"""
from pathlib import Path
import hashlib
import json
import re
from docx import Document

OUT = Path(__file__).resolve().parent
index = json.loads((OUT / "完整段落修订与证据索引.json").read_text())
methods = json.loads((OUT / "方法段替换稿与引用映射.json").read_text())
baseline = Path(index["baseline"]["path"])
assert hashlib.sha256(baseline.read_bytes()).hexdigest() == index["baseline"]["sha256"]
doc = Document(baseline)
assert len(doc.paragraphs) == 173 and len(doc.tables) == 4
anchors = {p["zero_based_body_p_index"]: p for s in index["sections"] for p in s["paragraphs"]}
assert len(anchors) == 43
for i, a in anchors.items():
    assert doc.paragraphs[i].text == a["current_text"]
    assert hashlib.sha256(doc.paragraphs[i].text.encode()).hexdigest() == a["sha256"]

text = {i: a["current_text"] for i, a in anchors.items()}
for row in methods["replacements"]:
    i = row["zero_based_direct_body_p_index"]
    assert row["old_text"] == doc.paragraphs[i].text
    text[i] = row["replacement_text"].replace("维度覆盖率", "内容覆盖比例").replace("阳性检出", "目标内容识别")

text.update({
22: "已有研究从政策工具视角比较医保与卫健政策文本〔HeBaliRamesh2022〕，并归纳我国卫生技术评估的五维价值框架〔Guo2023Online2024〕。本文将比较推进到特定药品政策语料中同一价值维度的关注比重及内容组合，以统一句段编码定位差异，结合条文和访谈解释其判断条件及使用环节。由此，把注意力配置的经验差别与评价结论需要传递的具体信息联系起来。",
1: "摘要：创新药评价的跨部门协调需要识别共同价值在不同决策环节的具体用途。基于2017—2026年三部门51个政策来源条目，按安全性、有效性、经济性、创新性、公平性五维开展规则与LLM句段编码，结合政策原文和专家访谈比较注意力配置。药监安全性占比为{{RULE_Q_NMPA_S}}%，医保经济性占比为{{RULE_Q_NHSA_Ec}}%；卫健公平性与安全性的相对排序随文件构成变化。医保经济性关注比重更高，卫健经济性内容却更常同时涉及公平性，两方法对应比例分别为{{RULE_P_F_GIVEN_Ec_NHC}}%、{{LLM_P_F_GIVEN_Ec_NHC}}%，医保为{{RULE_P_F_GIVEN_Ec_NHSA}}%、{{LLM_P_F_GIVEN_Ec_NHSA}}%。条文与访谈将差异具体化为费用保障、药品供应和临床使用的衔接问题。协调应以同一适应症和人群为基础，对照评价结论，衔接补充研究，并共享支付、配备及使用信息。",
6: "Abstract: Cross-agency coordination in innovative drug assessment requires identifying how shared values serve different decisions. Using 51 policy source entries from China's drug regulation, medical insurance and health administration systems issued between 2017 and 2026, this study applies rule-based and large language model coding across safety, efficacy, economic considerations, innovation and fairness, alongside policy interpretation and expert interviews. Rule-based attention shares are {{RULE_Q_NMPA_S}}% for safety in drug regulation and {{RULE_Q_NHSA_Ec}}% for economic considerations in medical insurance. The relative ranking of fairness and safety in health administration varies with document composition. Although medical insurance devotes a larger attention share to economic considerations, economic units in health administration more often also address fairness: {{RULE_P_F_GIVEN_Ec_NHC}}% and {{LLM_P_F_GIVEN_Ec_NHC}}% under the two coding methods, compared with {{RULE_P_F_GIVEN_Ec_NHSA}}% and {{LLM_P_F_GIVEN_Ec_NHSA}}% in medical insurance. Policy provisions and interview records connect these content differences to payment, drug supply and clinical use. Coordination should align assessment conclusions for the same indication and population, connect supplementary evidence requirements, and share information on payment coverage, institutional availability and actual use.",
37: "本文以五维注意力占比比较关注结构，进一步计算经济性单元中同时涉及公平性的比例，观察同一维度与其他评价内容的组合。该比例以经济性单元数为分母，五维注意力占比则以五维编码句次合计为分母。政策条文与访谈解释这些内容所对应的评价对象及用途，来源条目、年度系列、联合发文和操作定义边界变化检验结果的适用范围。",
44: "本文采用领域词典与语境规则进行全量编码，并以大语言模型（large language model，LLM）gpt-5.6-sol作平行测量。规则逐项处理药品对象、适用范围、父条款和否定关系；LLM按完整代码本阅读本单元及必要语境，输出五维判断与原文依据。词典分类的任务验证参照Kroon等〔Kroon2022〕，LLM标注参照Gilardi等[41]及Halterman和Keith[42]。两条方法使用同版原文、统一五维定义和纳排边界，通用修订分别保留版本及完整运行记录。本文采用版本为{{FINAL_RULES_ID}}与{{FINAL_LLM_ID}}。",
49: "表2 概率样本中规则与LLM的逐维开发诊断",
50: "Table 2 Dimension-Specific Development Diagnostics in the Probability Sample",
51: "注：主表为180条概率样本的设计加权结果，斜线前后分别以A、C为参照。κ、P、R和F1均报告原值，P、R、F1取值为0—1；零分母记NA。候选参照、分部门四格、近似95%区间和60条挑战材料另列，挑战材料不并入概率估计。",
52: "A/C在安全性、有效性、经济性、创新性、公平性上的κ依次为{{AC_KAPPA_S_E_Ec_I_F}}。两种方法对A/C及候选参照的逐维结果见表2及配套统计。{{FINAL_ERROR_PATTERN_SUMMARY}}。本轮240条已向开发环节暴露，后续修订在此取得的改善属于开发诊断；独立真人验证尚待实施。",
55: "表3 全语料三部门五维注意力占比",
56: "Table 3 Full-Corpus Five-Dimensional Attention Shares by Department",
57: "注：规则与LLM分别覆盖三部门4481个分析单元；占比单位为%，分母为本部门五维编码句次合计。同一单元可跨维计数，同维至多一次，精确五项合计100%；显示值独立舍入。180条概率样本的同样本比较及其设计加权估计另列，推断范围为剩余规范代表框。",
59: "结合3场研讨会材料及2份专家访谈纪要，比较评价内容与决策情境。两份访谈分别来自高校药事管理领域和医保评审领域专家；药监环节主要依据公开技术文件，研讨会材料补充跨环节问题。参照Carter等[50]的多来源对照方法，将纪要陈述与政策条文相互参照。",
60: "围绕判断对象、证据要求、结果用途及跨环节实施问题整理资料，记录陈述所指的评价环节，并回查对应政策。纪要中关于支付准入、医院配备和临床使用的陈述用于解释衔接问题；涉及不同意见时保留其评价条件，具体引文、日期及来源定位列入配套证据索引。",
63: "全量编码呈现不同的五维关注结构（表3、图2）。规则结果中，药监安全性、有效性、经济性、创新性、公平性分别为{{RULE_Q_NMPA_S}}%、{{RULE_Q_NMPA_E}}%、{{RULE_Q_NMPA_Ec}}%、{{RULE_Q_NMPA_I}}%、{{RULE_Q_NMPA_F}}%；医保分别为{{RULE_Q_NHSA_S}}%、{{RULE_Q_NHSA_E}}%、{{RULE_Q_NHSA_Ec}}%、{{RULE_Q_NHSA_I}}%、{{RULE_Q_NHSA_F}}%；卫健分别为{{RULE_Q_NHC_S}}%、{{RULE_Q_NHC_E}}%、{{RULE_Q_NHC_Ec}}%、{{RULE_Q_NHC_I}}%、{{RULE_Q_NHC_F}}%。药监关注风险及疗效证据，医保突出经济价值与支付安排，卫健的公平性和安全性在本口径中占比较高。",
68: "两种机器编码在三部门15项同维度比较中，有{{DEPARTMENT_DIRECTION_MATCH_N}}项点估计方向一致。{{FINAL_DEPARTMENT_COMPARISON_SUMMARY}}。创新性仍按创新属性、注册分类和研发支持的具体命题判定，不能由相关名称在文件中的出现直接赋值。",
69: "{{PROBABILITY180_ATTENTION_SUMMARY}}。{{FINAL_FULL_CORPUS_COMPOSITION_SUMMARY}}。全量分布描述所收集政策文本的关注结构，概率样本的估计则对应其剩余规范代表框，两种范围分别列示。",
71: "维度占比之外，同一评价内容与哪些价值要求共同出现，进一步呈现部门差异。医保经济性注意力占比高于卫健，而经济性单元中同时涉及公平性的比例，规则下卫健为{{RULE_JOINT_NHC}}/{{RULE_Ec_NHC}}（{{RULE_P_F_GIVEN_Ec_NHC}}%）、医保为{{RULE_JOINT_NHSA}}/{{RULE_Ec_NHSA}}（{{RULE_P_F_GIVEN_Ec_NHSA}}%）；LLM下分别为{{LLM_JOINT_NHC}}/{{LLM_Ec_NHC}}（{{LLM_P_F_GIVEN_Ec_NHC}}%）和{{LLM_JOINT_NHSA}}/{{LLM_Ec_NHSA}}（{{LLM_P_F_GIVEN_Ec_NHSA}}%）。{{FINAL_CORE_SENSITIVITY_SUMMARY}}。这一补充比例观察经济性内容的共同表达，具体用途结合条文判断（表4）。",
76: "医保经济性内容展开为药价、治疗费用、销量预测、支付范围及基金影响。《申报指南》[37]要求报告未来两年销量与基金影响，《谈判药品续约规则》[55]以实际药品费用与预算的比值安排续约和支付标准调整。2025年官方问答将基金支付监测与可及性联系起来，并说明支付标准包含基金及参保人员共同支付的费用〔NHSAFAQ2025〕，呈现经济内容与保障安排的共同表达。卫健《管理指南》[35]将价格水平、可获得性和可负担性共同纳入可及性评价，并结合机构配备、短缺和家庭支付能力判断；药品监测政策[57]进一步要求共享信息、保障供应和应用评价结果。经济性与公平性内容由此对应费用保障、供应和使用等具体问题。",
83: "从制度逻辑视角看，药监的科学—安全逻辑对应研发与上市时对获益风险及证据充分性的判断，医保的经济—效率逻辑对应公共支付中的比较价值与资源约束，卫健的公共服务逻辑对应临床安全使用、供应及需求回应。五维占比呈现这些任务的关注重心，同句内容组合进一步显示相同维度怎样进入不同的判断条件。经济性与公平性的共同表达，以及相关支付和供应条文，为解释决策任务如何组织评价内容提供了具体材料。",
84: "从注意力基础观看，三部门共同关注安全性、有效性及临床价值，关注内容的比例与组合又随决策任务有所区别。医保申报引用药审技术评审报告[37]，体现上市证据的跨环节使用；卫健将价格与实际获取共同评价[35]，使支付条件与机构配备、临床使用建立联系。史录文教授访谈纪要所记录的“医保进了，医院不一定用”，进一步呈现支付准入与院端使用之间需要衔接的实践问题。",
85: "上述差异将协调具体落实到结论及其使用条件：以共同人群核对比较方案，标明评价证据回答了哪些问题，并连通支付范围、患者负担、机构配备和实际使用的信息。这样，各部门的专门判断可以形成后续环节可识别、可复用的证据，支付和临床反馈又能支持下一轮评价。",
87: "针对五维关注结构及同一维度的内容组合差异，协调应连接上市评价、支付判断与临床执行。以下围绕结论对照、补充证据和信息共享提出路径，并明确各部门需要提供的内容及更新节点。",
97: "现行目录调整公示申报信息，并向企业反馈评审结果、参照药和支付范围[38]；2025年措施[2]提出医疗、医保、医药信息互通。在此基础上，可按同一药品、适应症和人群维护带来源、数据字典和更新日期的证据目录，将支付范围、基金与个人费用责任、机构配备、短缺及实际使用信息相衔接，分层管理保密价格资料。医保提供费用保障条件，卫健反馈机构配备及使用情况；出现重大风险、疗效或供应变化时同步更新。该安排有助于定位支付准入后需要继续处理的供应和使用问题。",
99: "本文在统一五维框架下比较三部门创新药评价的注意力配置。{{FINAL_CONCLUSION_ATTENTION_SUMMARY}}。进一步比较经济性单元的内容组合，卫健同时涉及公平性的比例为{{RULE_P_F_GIVEN_Ec_NHC}}%、{{LLM_P_F_GIVEN_Ec_NHC}}%，医保为{{RULE_P_F_GIVEN_Ec_NHSA}}%、{{LLM_P_F_GIVEN_Ec_NHSA}}%。政策原文和访谈将该差别落实到费用保障、药品供应与临床使用环节，说明共同评价维度在不同任务中具有具体的判断条件和用途。",
100: "本文的经验贡献是，将部门五维关注比重的比较推进到同名评价维度的内容组合及使用环节，明确评价结论跨部门衔接需要传递哪些信息。相应路径包括建立同一适应症和人群下的结论对照，衔接上市后研究与支付、临床再评价，并共享支付范围、患者负担、机构配备及实际使用反馈，使专业分工形成连续使用的评价证据。",
101: "后续研究可结合药品层面的评审与使用资料，考察上述信息衔接与评价周期、证据复用及患者获得之间的联系。",
105: "LLM使用gpt-5.6-sol，推理强度medium，按完整手册输出五维判断。每批最多25条，在独立批次上下文中处理同版原文及必要语境，不提供参考标签或机器预测。实际采用的全量运行{{FINAL_LLM_EXECUTION_SUMMARY}}；首次技术有效输出、失败重试、版本和提示SHA见封存清单。依Törnberg[62]关于记录任务及提示的建议，重复性单列为{{FINAL_LLM_STABILITY_SUMMARY}}。服务未暴露的参数记为不可得。",
109: "官方问答与正文分别标记，重复附件单次计数，联合专业文件保留全部署名机关并按既定系统单次归属。文件等权时，先计算各文件内部比例，再对对应分母非零的文件求均值；零分母名单单列。系列等权沿用冻结系列分类。全量机器分布、180条概率样本、60条挑战材料分别存储，样本设计、身份与适用范围见附录A。",
})

# Minimum corrections found by the independent prose reviewer.
text[76] = text[76].replace("药品监测政策[57]", "药品监测政策〔NHCMonitoring2019〕")
text[67] = "注：据同版原文的全量规则与LLM编码分别绘制。分母为同部门五维编码句次合计，每组精确五项合计100%；完整数值见表3，同180条概率样本的参考与方法比较另列。图中保留1位小数，表3保留2位小数。"
text[59] = "结合3场研讨会材料、史录文教授访谈纪要及医保调研纪要，比较评价内容与决策情境。药监环节主要依据公开技术文件，研讨会材料补充跨环节问题。参照Carter等[50]的多来源对照方法，将纪要陈述与政策条文相互参照。"
text[79] = text[79].replace("受访医保评审专家指出", "医保调研纪要记录")
text[96] = text[96].replace("受访医保评审专家提到", "医保调研纪要提到").replace("高校药事管理专家提到", "史录文教授访谈纪要提到")

# Apply the closed 50-paragraph review to the same baseline. Two retained
# paragraphs also need scope/attribution corrections, bringing the map to52.
text[45] = doc.paragraphs[45].text.replace(
    "规则全量结果用于主描述，LLM与A/C结果用于平行比较。",
    "规则全量结果用于主描述，LLM全量编码作平行测量；A/C及候选参照用于180条概率样本的同样本比较，60条挑战材料单列诊断。")
text[81] = doc.paragraphs[81].text.replace(
    "医保评审领域受访专家则讨论", "医保调研纪要则讨论")
text[22] = text[22].replace("我国卫生技术评估的五维价值框架", "医保目录专家评审的五维价值框架")
text[54] = text[54].replace(
    "注意力占比、内容覆盖比例和同句条件比例分别使用对应分母；聚合分布稳定不能替代句级效度。",
    "内容覆盖比例为某维被编码为1的单元数占本部门全部分析单元数的比例；注意力占比及经济性单元中的公平性比例分别按前述分母计算。")
text[68] = "在3组部门配对与5个维度组成的15项注意力占比差比较中，两种机器编码有{{DEPARTMENT_DIRECTION_MATCH_N}}项点估计方向一致。{{FINAL_DEPARTMENT_COMPARISON_SUMMARY}}。"
text[100] = "本文的经验贡献是，将部门五维关注比重的比较推进到经济性内容与公平性的共同表达，并结合政策原文解释评价结论的使用环节。据此提出结论对照、补充证据与信息共享的衔接安排：建立同一适应症和人群下的结论对照，衔接上市后研究与支付、临床再评价，共享支付范围、患者负担、机构配备及实际使用反馈，以支持评价证据跨环节连续使用。"
text[96] = text[96].replace(
    "史录文教授访谈纪要提到经济测算与医疗数据获取的困难",
    "史录文教授访谈纪要指出商保获取公立医院数据困难，影响测算")
text[84] = text[84].replace(
    "使支付条件与机构配备、临床使用建立联系", "在评价内容上连接支付条件与机构配备、临床使用").replace(
    "史录文教授访谈纪要所记录的", "2025年6月9日史录文教授访谈纪要所记录的")
text[48] = text[48].replace(
    "5000次重新计算指标后", "卫健层为全查，重复权重恒为1；5000次重新计算指标后")
text[48] = text[48].replace(
    "分报T裁定前A/C的κ与观察一致率，以及规则、LLM分别对A、C和候选参照的κ、精确率、召回率及F1。",
    "采用Cohen的κ[44]报告T裁定前A/C的一致性及两条方法分别对各参照的一致性，另报观察一致率。参照Kristensen-McLachlan等[45]，以精确率、召回率及F1报告分类表现；三者分别反映所识别内容与参照相符的比例、参照内容被识别的比例及二者的调和均值。")
text[106] = text[106].replace("两者均不替代真人判断或语料外验证。", "")
text[53] = text[53].replace("，不能替代单元标签验证", "")
text[1] = text[1].replace("卫健公平性与安全性的相对排序随文件构成变化。", "")
text[6] = text[6].replace("The relative ranking of fairness and safety in health administration varies with document composition. ", "")

records = []
for i in sorted(text):
    old = doc.paragraphs[i].text
    records.append({
        "zero_based_body_p_index": i,
        "old_text_sha256": hashlib.sha256(old.encode()).hexdigest(),
        "old_text": old,
        "replacement_text": text[i],
        "unchanged_after_review": old == text[i],
        "measurement_fields": sorted(set(re.findall(r"\{\{([^{}]+)\}\}", text[i]))),
        "unresolved_source_keys": sorted({key for token in re.findall(r"〔([^〕]+)〕", text[i]) for key in re.split(r"[；;]", token)}),
    })

report = {
    "status": "complete_staged_local_prose_revision_not_a_finished_manuscript",
    "date": "2026-10-08",
    "baseline": index["baseline"],
    "DOCX_modified": False,
    "labels_or_statistics_modified": False,
    "all_43_mapped_paragraphs_covered": set(anchors).issubset(text),
    "replacements": records,
    "measurement_fields": sorted({k for x in records for k in x["measurement_fields"]}),
    "source_keys": sorted({k for x in records for k in x["unresolved_source_keys"]}),
    "source_definitions": {
        "HeBaliRamesh2022": "He, Bali, Ramesh (2022), Active stewardship in healthcare: Lessons from China's health policy reforms. Social Policy & Administration 56(6):925–940, DOI10.1111/spol.12832. Closest-comparison scope verified in methods-20261008/先行研究与具体增量对照.md.",
        "Guo2023Online2024": "Guo et al., Health Technology Assessment in China's Health Care Sector: Development and Applications. Health Systems & Reform9(3):2327099, volume year2023, online2024-05-08, DOI10.1080/23288604.2024.2327099. Table1five-value framework verified; publication year and online date kept distinct.",
        "Kroon2022": "Kroon, van der Meer, Vliegenthart (2022), Computational Communication Research 4(2):528–570, DOI10.5117/CCR2022.2.006.KROO; methods-20261008 source audit holds full-text scope.",
        "NHSAFAQ2025": "Actual 2025 directory-adjustment official FAQ attachment, units00647/00752. Resolve exact primary entry and component URL from frozen source inventory before final bibliography; do not substitute a 2026 FAQ merely because the provisions look similar.",
        "NHCMonitoring2019": "国家卫生健康委关于开展药品使用监测和临床综合评价工作的通知，国卫药政函〔2019〕80号；官网发布日期2019-04-09，正式存档正文与单元00337/00338/00347核对。不能误指基稿[57]的2018处方审核规范。",
    },
    "table_contracts": {
        "table_1": "Preserve theory/task mapping; all factual statements source-checked, no attention numbers here.",
        "table_2": "180 probability, eligible_design_weighted ALL; each method/dimension vs A and C, κ/P/R/F1 all original decimal units; same adopted outputs and intervals. Candidate and department comparisons remain in full supplement.",
        "table_3": "6 rows: 3departments×2full-corpus methods, denominator five-dimension count sum; all five dimensions separately. No sample references mixed into full-corpus rows.",
        "table_4": "Preserve 3department×5dimension substantive content mapping; crosscheck any revised example against source units.",
        "sample_attention_supplement": "All A/C/candidate/machine methods on same180 with weights and proper frame scope; do not present as4481 census.",
    },
    "figure_contracts": {
        "figure_1": "Preserve theoretical framework after text/label consistency check.",
        "figure_2": "Redraw five separate attention shares for adopted full-corpus rule and LLM outputs; exact five shares sum100. Include method panels or labels, no S+E aggregate.",
    },
    "integration_gates": [
        "Root decides explicit adopted methods after complete candidate03 comparison and semantic audit; no output splicing or score selection.",
        "Resolve all numeric/summary fields from the same SHA-bound statistical bundle and verify direction wording against actual results; change prose if evidence contradicts it.",
        "Resolve all new source keys and renumber citations in first-appearance order; existing numbers remain baseline references until that step.",
        "Preserve AI identity and pending true-human validation. No claim all metrics≥0.8, publication guaranteed, global first, causality, or actual patient fairness inferred from text.",
        "Generate final DOCX only after measurement decision; fresh page render, word count, all-table/figure/reference checks.",
    ],
}
(OUT / "完整论文局部替换稿.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
md = ["# 完整论文局部替换稿", "", "2026年10月8日。覆盖既定43处段落及相关表题表注；已交付DOCX未修改。双花括号必须由最终同版统计填充，方括号引文暂沿用基稿编号，新来源key待统一排号。此稿不是已完成终稿。", ""]
for x in records:
    md.extend([f"## 原稿直接body段落{x['zero_based_body_p_index']}" + ("（核查后保留）" if x['unchanged_after_review'] else ""), "", x["replacement_text"], ""])
md.extend(["## 接入要求", "", *["- " + t for t in report["integration_gates"]]])
(OUT / "完整论文局部替换稿.md").write_text("\n".join(md) + "\n")
print(json.dumps({"paragraph_n": len(records), "mapped_43_covered": report["all_43_mapped_paragraphs_covered"], "changed_n": sum(not r['unchanged_after_review'] for r in records), "numeric_fields_n": len(report["measurement_fields"]), "source_keys": report["source_keys"], "DOCX_modified": False}, ensure_ascii=False))
