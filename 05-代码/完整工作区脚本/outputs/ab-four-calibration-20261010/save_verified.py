from pathlib import Path
import hashlib, json, os, re, zipfile
import openpyxl

root=Path('/Users/rainbow/Desktop/战略院/药企项目-创新药')
out=root/'outputs/ab-four-calibration-20261010'
target=root/'outputs/goal-calibration-20261010/09-delivery/agent-review-20261010/剩余4项单独校准表.xlsx'
before=out/'before.xlsx'
generated=out/'artifact-edited.xlsx'
assert target.read_bytes()==before.read_bytes(), 'Source changed during review; do not replace.'
addresses={f'{c}{r}' for r in range(7,11) for c in ['D','E']}
def mask_authorized(xml):
    for address in sorted(addresses):
        matches=list(re.finditer(r'<x:c\b[^>]*\br="'+address+r'"[^>]*>', xml))
        assert len(matches)==1, (address,len(matches))
        match=matches[0]
        if match.group().rstrip().endswith('/>'):
            end=match.end()
        else:
            end=xml.index('</x:c>',match.end())+len('</x:c>')
        xml=xml[:match.start()]+'<authorized-cell/>'+xml[end:]
    return xml
sheet='xl/worksheets/sheet1.xml'
with zipfile.ZipFile(before) as src, zipfile.ZipFile(generated) as authored:
    old=src.read(sheet).decode()
    new=authored.read(sheet).decode()
    assert mask_authorized(old)==mask_authorized(new), 'Changes outside D7:E10'
    assert src.read('xl/styles.xml')==authored.read('xl/styles.xml')
    # Carry only the Artifact Tool-authored worksheet into the original package,
    # keeping original drawings, relationship IDs and every unrelated part exact.
    temp=out/'verified.xlsx'
    with zipfile.ZipFile(temp,'w') as dst:
        for info in src.infolist():
            dst.writestr(info, authored.read(sheet) if info.filename==sheet else src.read(info.filename))
with zipfile.ZipFile(before) as src, zipfile.ZipFile(temp) as dst:
    changed=[n for n in src.namelist() if src.read(n)!=dst.read(n)]
    assert changed==[sheet]
oldbook=openpyxl.load_workbook(before)
newbook=openpyxl.load_workbook(temp)
assert oldbook.sheetnames==newbook.sheetnames
diff=[]
for s in oldbook:
    t=newbook[s.title]
    assert s.max_row==t.max_row and s.max_column==t.max_column
    assert len(s._images)==len(t._images)
    for row in s:
        for c in row:
            n=t[c.coordinate]
            assert c._style==n._style
            if c.value!=n.value:
                diff.append((s.title,c.coordinate))
assert set(diff)=={('校准表',a) for a in addresses}
expected=json.loads((out/'judgments.json').read_text())
assert [[newbook['校准表'].cell(r,c).value for c in [4,5]] for r in range(7,11)]==expected
assert target.read_bytes()==before.read_bytes(), 'Source changed before save.'
os.replace(temp,target)
assert zipfile.is_zipfile(target)
(out/'verification.json').write_text(json.dumps({
    'output':str(target),'changed_cells':sorted(addresses),
    'codes':[1,1,1,1],'other_package_parts_byte_identical':True,
    'styles_preserved':True,'source_images_preserved':True,
    'sha256_before':hashlib.sha256(before.read_bytes()).hexdigest(),
    'sha256_after':hashlib.sha256(target.read_bytes()).hexdigest(),
    'review_identity':'AI semantic review'
},ensure_ascii=False,indent=2))
print(json.dumps({'saved':str(target),'changed_cells':sorted(addresses),'codes':[1,1,1,1],'preservation_verified':True},ensure_ascii=False))
