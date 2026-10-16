"""Read-only workbook delivery comparison. Never saves a workbook."""
import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validation_signature(sheet):
    return sorted(tostring(item.to_tree(), encoding="unicode")
                  for item in sheet.data_validations.dataValidation)


def columns_signature(sheet):
    return {name: {
        "width": column.width, "hidden": column.hidden,
        "bestFit": column.bestFit, "min": column.min, "max": column.max,
        "outlineLevel": column.outlineLevel, "collapsed": column.collapsed,
        "style": column.style,
    } for name, column in sheet.column_dimensions.items()}


def hyperlink_signature(cell):
    link = cell.hyperlink
    if link is None:
        return None
    return {"target": link.target, "location": link.location,
            "tooltip": link.tooltip, "display": link.display}


def local_target(workbook_path, target):
    if not target:
        return None
    parsed = urlsplit(target)
    if parsed.scheme not in ("", "file"):
        return None
    target_path = Path(unquote(parsed.path))
    if not target_path.is_absolute():
        target_path = workbook_path.parent / target_path
    return target_path.resolve()


def check(original, final, labels_path, report_path):
    inputs = {"original": original, "final": final, "labels": labels_path}
    for name, path in inputs.items():
        if not path.is_file():
            raise FileNotFoundError(f"Missing {name}: {path}")
    before = {name: sha256(path) for name, path in inputs.items()}
    records = {}
    duplicate_labels = []
    # Only the exported values are extracted; no semantic reevaluation occurs.
    for line in labels_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        item = json.loads(line)
        key = item["blind_id"]
        if key in records:
            duplicate_labels.append(key)
        records[key] = {dim: judgment["value"]
                        for dim, judgment in item["judgments"].items()}

    # Normal loading exposes hyperlinks, validation, widths and panes. No save
    # operation is used; byte hashes are checked again after all comparisons.
    original_wb = load_workbook(original, read_only=False, data_only=False)
    final_wb = load_workbook(final, read_only=False, data_only=False)
    failures = []
    totals = Counter()
    sheet_reports = []

    def require(ok, message):
        if not ok:
            failures.append(message)

    require(original_wb.sheetnames == final_wb.sheetnames,
            "Worksheet names or order differ from frozen original")
    require(not duplicate_labels, "Duplicate blind IDs in final JSONL")
    require(set(records) == {f"A-{n:06}" for n in range(1, 801)},
            "Final JSONL does not contain exactly A-000001 through A-000800")

    observed_keys = []
    for original_sheet in original_wb:
        name = original_sheet.title
        if name not in final_wb:
            continue
        final_sheet = final_wb[name]
        headers_equal = ([c.value for c in original_sheet[1]] ==
                         [c.value for c in final_sheet[1]])
        shape_equal = ((original_sheet.max_row, original_sheet.max_column) ==
                       (final_sheet.max_row, final_sheet.max_column))
        validations_equal = (validation_signature(original_sheet) ==
                             validation_signature(final_sheet))
        panes_equal = original_sheet.freeze_panes == final_sheet.freeze_panes
        widths_equal = columns_signature(original_sheet) == columns_signature(final_sheet)
        for ok, label in [(headers_equal, "headers"), (shape_equal, "shape"),
                          (validations_equal, "data validation"),
                          (panes_equal, "freeze panes"),
                          (widths_equal, "column dimensions")]:
            require(ok, f"{name}: changed {label}")

        abc_mismatches = []
        hyperlink_mismatches = []
        missing_targets = []
        labels_mismatches = []
        incomplete = []
        for row in range(2, original_sheet.max_row + 1):
            for col in (1, 2, 3):
                left = original_sheet.cell(row, col)
                right = final_sheet.cell(row, col)
                if left.value != right.value or left.data_type != right.data_type:
                    abc_mismatches.append(right.coordinate)
            blind_id = original_sheet.cell(row, 1).value
            dim = original_sheet.cell(row, 3).value
            observed_keys.append((blind_id, dim))
            totals["judgment_rows"] += 1
            expected = records.get(blind_id, {}).get(dim, "MISSING")
            value_cell = final_sheet.cell(row, 4)
            reason_cell = final_sheet.cell(row, 5)
            value, reason = value_cell.value, reason_cell.value
            if expected is None:
                lawful_null = (value is None and isinstance(reason, str) and
                               (reason.startswith("待定：") or reason.startswith("待定:"))
                               and bool(reason[3:].strip()))
                if lawful_null:
                    totals["lawful_null"] += 1
                else:
                    incomplete.append(row)
            else:
                typed_binary = (type(value) is int and value in (0, 1)
                                and value_cell.data_type != "f")
                reason_complete = isinstance(reason, str) and bool(reason.strip())
                if typed_binary and reason_complete:
                    totals["complete_binary"] += 1
                else:
                    incomplete.append(row)
                if not typed_binary or value != expected:
                    labels_mismatches.append(row)
            left_link = hyperlink_signature(original_sheet.cell(row, 1))
            right_link = hyperlink_signature(final_sheet.cell(row, 1))
            if left_link:
                totals["original_a_hyperlinks"] += 1
            if right_link:
                totals["final_a_hyperlinks"] += 1
            if not left_link or left_link != right_link:
                hyperlink_mismatches.append(row)
            target = local_target(final, right_link["target"] if right_link else None)
            if target is not None and target.is_file():
                totals["existing_local_a_hyperlink_targets"] += 1
            else:
                missing_targets.append({"row": row, "target": str(target)})

        for values, label in [(abc_mismatches, "A/B/C values or order"),
                              (hyperlink_mismatches, "A hyperlinks"),
                              (missing_targets, "local hyperlink targets"),
                              (labels_mismatches, "JSONL label matches"),
                              (incomplete, "D/E completeness")]:
            require(not values, f"{name}: failed {label}; count={len(values)}")
        sheet_reports.append({
            "sheet": name,
            "rows_including_header": final_sheet.max_row,
            "columns": final_sheet.max_column,
            "headers_equal": headers_equal,
            "shape_equal": shape_equal,
            "abc_value_order_mismatches": abc_mismatches,
            "label_mismatch_rows": labels_mismatches,
            "incomplete_rows": incomplete,
            "hyperlink_mismatch_rows": hyperlink_mismatches,
            "missing_local_targets": missing_targets,
            "data_validation_equal": validations_equal,
            "data_validation_count": len(final_sheet.data_validations.dataValidation),
            "freeze_panes_equal": panes_equal,
            "freeze_panes": final_sheet.freeze_panes,
            "column_dimensions_equal": widths_equal,
            "column_dimensions": columns_signature(final_sheet),
        })

    expected_keys = {(f"A-{n:06}", dim) for n in range(1, 801)
                     for dim in ("S", "E", "Ec", "I", "F")}
    require(len(observed_keys) == 4000 and set(observed_keys) == expected_keys,
            "Workbook key coverage is not exactly 800 records × five dimensions")
    for field in ("judgment_rows", "original_a_hyperlinks", "final_a_hyperlinks",
                  "existing_local_a_hyperlink_targets"):
        require(totals[field] == 4000, f"{field}: expected 4000; got {totals[field]}")
    require(totals["complete_binary"] + totals["lawful_null"] == 4000,
            "Not all 4000 D/E rows are complete")
    original_wb.close()
    final_wb.close()
    after = {name: sha256(path) for name, path in inputs.items()}
    require(before == after, "Input file bytes changed during read-only check")
    result = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "reviewer": "AI-Codex independent delivery check",
        "scope": "Technical export verification only; no semantic relabeling",
        "inputs": {name: str(path) for name, path in inputs.items()},
        "input_sha256_before": before, "input_sha256_after": after,
        "read_only_inputs_unchanged": before == after,
        "final_jsonl_records": len(records), "totals": dict(totals),
        "worksheets": sheet_reports, "failures": failures,
        "pass": not failures,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"pass": result["pass"], "totals": dict(totals),
                      "failures": failures, "report": str(report_path)}, ensure_ascii=False))
    return 0 if result["pass"] else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--original", required=True, type=Path)
    parser.add_argument("--final", required=True, type=Path)
    parser.add_argument("--labels", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    arguments = parser.parse_args()
    raise SystemExit(check(arguments.original.resolve(), arguments.final.resolve(),
                           arguments.labels.resolve(), arguments.report.resolve()))
