from pathlib import Path
from collections import Counter
from html.parser import HTMLParser
from urllib.parse import urlsplit, unquote
import csv
import hashlib
import io
import json
import posixpath
import zipfile

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent.parent
PACKS = ROOT / 'outputs/goal-calibration-20261010/09-delivery/formal800-human-distribution-packages-20261010-sol'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def save_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

summary = json.loads((BASE / '抽检汇总.json').read_text())
comparison = json.loads((BASE / '与规则法对照汇总.json').read_text())
technical = json.loads((BASE / 'checks/independent-delivery-check.json').read_text())
assert technical['pass'] and not technical['failures']
assert sha(BASE / 'AI原始判读-800条.jsonl') == summary['initial_jsonl_sha256']
reviews = {r['blind_id']: r for r in map(json.loads, (BASE / 'AI原始判读-800条.jsonl').read_text().splitlines())}
metadata = json.loads((BASE / 'checks/sample-metadata.json').read_text())['records']
rules = {r['id']: r for r in map(json.loads, Path(comparison['rule_predictions_path']).read_text().splitlines())}

with (BASE / '规则法未给出二元值清单.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['A盲号', '原记录号', '原样本分组', '维度', '规则法值', '本轮AI值', '当前完整命题', 'AI依据', 'AI定位', 'AI理由', '规则法理由'])
    for blind, sid, domain, dim, rule_value, ai_value in comparison['not_joint_binary']:
        item = reviews[blind]['judgments'][dim]
        source = json.loads((BASE / f'source-A/coder-A/materials/{blind}.json').read_text())['source_record']
        writer.writerow([blind, sid, domain, dim, '未给出二元值', ai_value, source['text'], item['evidence'], item['locator'], item['reason'], rules[sid]['decisions'][dim].get('reason', '')])

originals = {}
for code in 'AB':
    path = PACKS / f'formal800-{code}-independent.zip'
    checked = 0
    with zipfile.ZipFile(path) as original:
        assert original.testzip() is None
        for item in original.infolist():
            if item.is_dir():
                continue
            target = BASE / f'source-{code}' / item.filename
            assert target.read_bytes() == original.read(item.filename), item.filename
            checked += 1
    originals[code] = {'path': str(path), 'sha256': sha(path), 'unchanged_extracted_files': checked}
assert originals['A']['sha256'] == summary['source_package_sha256']

source_sets = {}
for code in 'AB':
    source_sets[code] = {}
    for path in (BASE / f'source-{code}/coder-{code}/materials').glob('*.json'):
        source = json.loads(path.read_text())['source_record']
        source_sets[code][source['id']] = source['text']
assert len(source_sets['A']) == 800 and source_sets['A'] == source_sets['B']

report = '''800条AI辅助抽检结果说明
完成日期：2026年10月10日

本轮已对选定的800条材料完成S、E、Ec、I、F五个维度，共4000项语义判断；每项均填写了0/1值、判断理由、依据和位置。AI待定项为0。800条是本项目此前选定的抽检规模，不是方法规定的固定数量。

样本原有用途划分保持不变：600条用于独立检验，200条用于已接触材料的补充检查。本次执行者为AI-Codex，结果属于一轮AI辅助抽检。三个子任务和主任务分别阅读不重叠的记录范围，合并为这一轮结果；不计作A/B两位真人结果，也不替代规定参数和批次流程的正式LLM生产主运行或人类信度检验。

A、B原包经核对包含相同的800条记录和文本，排列不同。本轮按A盲号填写一份结果表。两份原始真人分发ZIP及其解压内容均保留原样。

各维度判为1的数量（分母均为各组条数）：
组别                         S     E     Ec    I     F
独立检验用途600条           111    61    118   21     8
已接触补充检查200条          34    30     62   29    10
合计800条                  145    91    180   50    18
五个维度可同时成立，不能将维度数量视为互斥类别数量。

与现有规则法的对照：
先冻结本轮AI初判，再读取规则法结果进行对照。对照对象是rule-full4481-v2-root/run4481-first的规则法结果，其手册、提示和共同协议版本与本轮材料包一致；这不是与正式LLM结果的对照。
4000项中，3975项双方均有二元值，其中3697项一致，278项不一致，分歧涉及226条材料。另25项（5条材料的五个维度）规则法未给出二元值，单列清单，不计入3975项的一致性分母。

组别                    双方有二元值    一致    分歧    规则法未给二元值
独立检验用途600条              2980    2787     193                  20
已接触补充检查200条             995     910      85                   5
合计                          3975    3697     278                  25

维度分歧数量：S 89项、E 49项、Ec 87项、I 44项、F 9项。
以上是所选样本的未加权描述性对照。分歧尚未经过真人裁定，不能将分歧直接认定为规则法错误，也不能将双方一致比例称为准确率或总体正确率。正式人类一致性仍需真人独立填写和后续裁决。

如何使用完整包：
1. 完整解压后，打开coder-A/AI抽检800条-五维结果.xlsx。
2. 保留coder-A/materials与shared-source的相对目录结构，点击表格A列盲号即可查看相应材料及链接。
3. 表格D列是本轮AI值；E列是理由、依据和定位。原表的编号、命题、维度及行顺序均保持不变。
4. “与规则法分歧清单.csv”列出278项分歧，供集中复核；“规则法未给出二元值清单.csv”列出其余25项。
5. “AI抽检800条-逐维明细.csv”是4000项完整明细；“AI原始判读-800条.jsonl”是冻结的初判记录。“抽检汇总.json”“与规则法对照汇总.json”和checks内记录用于追溯。
6. coder-A内的manual.txt、execution-prompt.txt、common-execution-protocol.txt为原材料包的定义及流程参考，原文保留；本轮实际执行身份与范围以本说明及audit-plan.txt为准。

交付检查：
已核验800条/4000项覆盖、编号与来源对应、原文引文匹配、标签与输出一致。结果表4000个材料链接及材料页3405处本地资源链接目标均存在，数据验证、冻结窗格、列宽保留，并检查了首段、中段和末段的表格渲染。这些检查确认交付完整性，不替代对语义判断正确性的复核。

原始初判已冻结。后续若复核后调整标签，应另存修订记录并保留本轮初判。
'''
(BASE / '抽检结果说明.txt').write_text(report, encoding='utf-8')

included = {}
source_root = BASE / 'source-A'
for folder in ['coder-A/materials', 'shared-source']:
    for path in (source_root / folder).rglob('*'):
        if path.is_file():
            included[path.relative_to(source_root).as_posix()] = path
for name in ['AI抽检800条-五维结果.xlsx', 'manual.txt', 'execution-prompt.txt', 'common-execution-protocol.txt', 'source-assets-index.json']:
    included['coder-A/' + name] = source_root / 'coder-A' / name
for name in ['抽检结果说明.txt', 'AI抽检800条-逐维明细.csv', '与规则法分歧清单.csv', '规则法未给出二元值清单.csv', 'AI原始判读-800条.jsonl', '抽检汇总.json', '与规则法对照汇总.json', '待定与补材清单.json', 'audit-plan.txt']:
    included[name] = BASE / name
for name in ['independent-delivery-check.json', 'source-link-check.json', 'progress.json']:
    included['checks/' + name] = BASE / 'checks' / name

manifest = {
    'identity': 'AI-Codex辅助抽检；一轮800条4000项，非真人结果',
    'records': 800,
    'judgments': 4000,
    'source_packages': originals,
    'same_800_source_ids_and_texts_in_A_B': True,
    'frozen_initial_sha256': summary['initial_jsonl_sha256'],
    'delivery_file_count_excluding_manifest': len(included),
    'files': {name: {'bytes': path.stat().st_size, 'sha256': sha(path)} for name, path in sorted(included.items())},
}
save_json(BASE / '交付文件清单.json', manifest)
included['交付文件清单.json'] = BASE / '交付文件清单.json'

archive = BASE / 'AI抽检800条-完整结果包.zip'
with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
    for name, path in sorted(included.items()):
        z.write(path, name)

class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
    def handle_starttag(self, tag, attrs):
        self.links.extend(value for key, value in attrs if key in ['href', 'src'] and value)

with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    names = set(z.namelist())
    assert names == set(included)
    for name, path in included.items():
        assert z.read(name) == path.read_bytes(), name
    local_links = 0
    for name in sorted(n for n in names if n.startswith('coder-A/materials/') and n.endswith('.html')):
        parser = Links()
        parser.feed(z.read(name).decode('utf-8'))
        for link in parser.links:
            parts = urlsplit(link)
            if parts.scheme or parts.netloc or not parts.path:
                continue
            target = posixpath.normpath(posixpath.join(posixpath.dirname(name), unquote(parts.path)))
            assert target in names, (name, link, target)
            local_links += 1
    workbook_bytes = z.read('coder-A/AI抽检800条-五维结果.xlsx')
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(workbook_bytes))
    workbook_links = 0
    for row in wb.worksheets[0].iter_rows(min_row=2):
        link = row[0].hyperlink.target
        target = posixpath.normpath(posixpath.join('coder-A', unquote(urlsplit(link).path)))
        assert target in names
        workbook_links += 1
    assert workbook_links == 4000 and local_links == 3405

seal = {
    'pass': True,
    'archive': str(archive),
    'archive_bytes': archive.stat().st_size,
    'archive_sha256': sha(archive),
    'files_checked': len(included),
    'workbook_links_checked_inside_archive': workbook_links,
    'material_local_links_checked_inside_archive': local_links,
    'frozen_initial_sha256': sha(BASE / 'AI原始判读-800条.jsonl'),
    'workbook_sha256': sha(source_root / 'coder-A/AI抽检800条-五维结果.xlsx'),
}
save_json(BASE / 'checks/交付包核验.json', seal)
print(json.dumps(seal, ensure_ascii=False, indent=2))
