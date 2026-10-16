from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import json, re, hashlib, sys
from lxml import etree as E
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
BASE=Path('/Users/rainbow/Desktop/战略院/创新药论文-完整归档-20260914/12-投稿定稿工作区/重跑修订稿-20260927/论文修订稿-匿名审稿版.docx')
MET=ROOT.parent/'autocode-v28-review-20261004-111306'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def W(s): return '{'+NS['w']+'}'+s
def plain(e): return ''.join(e.itertext()) if False else ''.join(e.xpath('.//w:t/text()',namespaces=NS))
records=json.loads((MET/'extended_metrics.json').read_text())['records']
index={(r['scope'],r['dimension'],r['method'],r['reference']):r for r in records}
def val(scope,dim,method,ref,key):return index[scope,dim,method,ref]['metrics'][key]['estimate']
def pair(scope,dim,method,key,percent=False):
    v=[val(scope,dim,method,r,key) for r in ['A','C']]
    return '/'.join(f'{x*100:.1f}' if percent else f'{x:.3f}' for x in v)

with ZipFile(BASE) as z:
    doc=E.fromstring(z.read('word/document.xml'))
body=doc.find('w:body',NS)
paras=body.findall('w:p',NS)
tables=body.findall('w:tbl',NS)

def replace(p,text,refs=True):
    pr=p.find('w:pPr',NS)
    rp=None
    for r in p.findall('w:r',NS):
        q=r.find('w:rPr',NS)
        if q is not None and q.find('w:vertAlign',NS) is None:
            rp=deepcopy(q);break
    for child in list(p):
        if child is not pr:p.remove(child)
    for part in re.split(r'(\[\d+(?:[,–—-]\d+)*\])',text) if refs else [text]:
        if not part:continue
        r=E.SubElement(p,W('r'))
        if rp is not None:r.append(deepcopy(rp))
        if refs and re.fullmatch(r'\[\d+(?:[,–—-]\d+)*\]',part):
            q=r.find('w:rPr',NS)
            if q is None:q=E.SubElement(r,W('rPr'))
            v=E.SubElement(q,W('vertAlign'));v.set(W('val'),'superscript')
        for line_no,line in enumerate(part.split('\n')):
            if line_no:E.SubElement(r,W('br'))
            t=E.SubElement(r,W('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=line

def cloned(p,text,refs=True):
    x=deepcopy(p);replace(x,text,refs);return x
def cellset(cell,text):
    ps=cell.findall('w:p',NS)
    if not ps:ps=[E.SubElement(cell,W('p'))]
    replace(ps[0],text,False)
    for p in ps[1:]:cell.remove(p)

abstract='摘要：基于2017—2026年55个政策来源条目，结合意义单元编码、抽样复核与专家访谈，考察创新药评价的注意力配置及跨环节衔接。研究发现，药品注册、医保准入与临床使用共同关注安全性、有效性和临床价值，但注册创新身份、上市获益风险与支付比较价值服务于不同决策，评价结论不能直接换算。药监来源安全性、有效性及医保来源经济性呈现较高的阳性点估计；部门关注分布的解释还需结合分类口径、材料构成与区间结果。同名维度在不同环节的评价对象、比较方案和证据要求存在差别，是理解评价差异与结论适用范围的关键。协调应以明确评价结论的适用条件为基础，推动已有证据复用、补充研究衔接和评审理由共享，为创新药跨环节评价提供可对照的信息。'
replacements={
1:abstract,
43:'本文采用领域词典与语境规则，并以大语言模型（LLM）作平行测量。Sevenans等[36]的政策议程编码、Loughran和McDonald[37]的领域词典研究及Grimmer和Stewart[29]的任务验证论述提供方法参考。两种方法遵循CAL-5D-20260926-v1.3五维规范，对同一4924个完整单元分别编码。当前规则为v28.2；LLM保留gpt-5.6-sol此前封存运行，不能将规则修订后的改善写成LLM改善。规则对药品对象、语义条件、条件与否定、名称提及及排除项作通用修订，并对104条输入补充经官方原文核验的语境。两者的上下文版本因此并非完全相同，另以规则在旧输入上的结果分离代码与补料影响。具体执行记录见附录A。',
46:'本研究将五维编码分别视为二元分类任务，报告A/C协商前一致性，以及规则、LLM相对各参照的分类表现。Cohen’s κ用于扣除边际分布决定的机会一致[46]，同时报告观察一致率及混淆矩阵；参照Kristensen-McLachlan等[47]的报告方式，以精确率、召回率和F1检验阳性内容识别，避免多数阴性抬高整体一致率。κ≥0.80是本项目后设的开发工作目标，Landis和Koch[39]的解释区间并非统一发表标准。现有A/C为agent首次独立判读，不能作为已完成的真人金标准。',
47:'1268条与H427均已用于当前规则的差错诊断，因此本版将二者列为开发复核材料。H427在旧版本预测封存时属于未暴露的有限剩余框，其历史验证地位不外推至v28.2。各维度按固定共同可判掩码比较，不将历史待定填0。当前规则全量待定为0，旧LLM仍保留49条五维待定。表2报告H427的κ、P、R和F1，1268条逐维κ另见表A1；全部观察一致率、TP/FP/FN/TN、阳性比例及排除记录见配套统计。依照定稿后独立验证原则[48]，正式真人验证拟由两名领域编码员先依据同版手册校准，在未参与方法调整的样本中独立判读完整句段及必要上下文，不查看机器或对方答案；协商前信度单列，分歧裁定后形成共同参考。实施记录和真人结果尚待补齐，不以现有A/C数值代填。',
48:'表2 当前规则与旧LLM在H427开发材料中的表现',
49:'Table 2 Performance of Current Rules and the Previous LLM Run in H427 Development Material',
50:'注：H427共427条，各维以A/C为参照的固定共同可判n为394/391；斜线前后分别为A/C。P、R单位为%，κ和F1为0—1。A/C自身一致性另按双方可判391条计算，不以机器结果筛选。95%区间采用来源文件聚类bootstrap，5000次重复，种子及无定义重复数见配套统计[49]；该区间只描述固定开发材料，不覆盖规则选择、系统性标注误差或语料外推广。历史待定及可判覆盖另报。',
51:'H427的A/C间κ按安全性、有效性、经济性、创新性、公平性依次为0.974、0.960、1.000、0.841、0.935。当前规则在两组材料的20项κ为0.874—0.974，精确率0.846—1.000、召回率0.816—0.986、F1为0.879—0.976；20项κ点估计均达到开发目标，但其中5项95%区间下限低于0.80。旧LLM的20项κ中11项达到该目标，创新性与公平性的8项均未达到，H427有效性—C亦未达到。1268条中，旧LLM创新性精确率为60.9%/64.8%，公平性为49.3%/55.5%，提示其阳性判定边界需进一步核查。上述结果支持当前规则相对固定参照的开发期一致性改善，不构成真人效度或新独立验证通过的证据。',
53:'总体估计沿用既定1268个唯一抽样单元及29层权重：从1531条框按25层抽600条，从3393条补集按四来源组抽668条，各层以总体数/样本数加权，状态总量除以来源组固定总体数（表3）。A/C分别估计，不将两份编码当作样本量翻倍。各层区间采用有限总体超几何检验反演[50]；边际区间在来源组或总体各层间分配0.05后合并。部门比较保持既定两位参照、五维、29层及阳性和阳性加待定两端点，共580项家族，以Bonferroni原理[51]构造至少95%的同时覆盖界。允许待定任意补全后，只有A/C均在同方向严格排除零，才记共同支持的方向；全部30项部门对比较均保留（表4、图2）。该合并方案是本研究的保守构造，并非文献规定的统一算法。内部衔接计划已确定保留此主分析，H427不并入总体估计。',
76:'注：单元格为阳性比例（待定比例），单位为%；每维分母为来源组全部N，五维可重叠。当前规则待定为0；旧LLM保留49条五维待定。数据描述固定语料的全量标签，版本与上下文差异见4.1节；不使用抽样主分析的区间，也不替换其权重或参考标注。',
101:'本文结合政策文本编码与专家访谈，比较创新药评价的关注内容与决策用途。文本分布呈现药监来源安全性、有效性及医保来源经济性阳性点估计较高的描述特征，但既定同时比较与待定处理未形成两份参照共同支持的稳健部门排序。政策条文则表明，三部门对安全、有效和临床价值具有共同关注，注册创新身份、上市获益风险及支付比较价值分别回答不同决策问题。创新药评价的协调应以这些具体判断及其适用条件为基础，将目标人群、比较方案、证据要求与不确定性纳入结论对照，推动证据复用、补充研究衔接和评审理由共享。',
103:'本研究的测量限于纳入公开文本的注意力分布，不代表个人认知或实际评审权重；文件体裁、年度材料及来源归属可能影响比例。医保五维框架对其他评价环节存在覆盖限制，低比例不等于不重视。两份访谈提供解释线索，不能代表部门统一立场或验证句段标签，药监环节也缺少一对一访谈。当前A/C为agent参照，规则改善及其区间属于已接触材料的开发复核；独立真人验证尚待实施。旧LLM尚未完成修订重跑，底层revision和temperature不可得，片内上下文与模板理由也限制复现。相关不确定性不由较高κ消除。',
105:'A1 编码规范与实现。五维操作定义固定为CAL-5D-20260926-v1.3，全文、父条款和表头解释药品对象、条件及引用，不复制其他单元标签。每单元每维至多计一次，可多标签。S限药物损害、风险控制及明确药品质量安全；E限防治效果、临床获益及对应证据要求；Ec限药品价格、费用与支付等；I区分实质创新属性、研发支持及创新资格比较优势；F按本研究保障获取、可负担与临床管理难度的操作口径判断。一般名称、事务性程序或其他对象的“安全”“有效”“创新”不自动赋标签。完整纳入、排除及子类型规则见配套手册。',
111:'语料与语境。4924个分析单元及621个原有语境单元冻结；本轮对104条输入增加官方原文核验的必要上下文，正文与ID均不改。固定旧输入仅更换代码改变257条，固定新代码补充上下文再改变13条；相对v27合计270条至少一个标签改变。自动规则通过168项回归检查，全量待定为0，原A/C与旧LLM保留首次结果。最终配对掩码、上下文版本与数据哈希分别封存。',
112:'操作定义边界。五项敏感性分别排除仅质量安全、仅创新资格比较优势、仅临床管理、仅未细分公平资料要求，或仅保留保障获取与可负担依据；混合独立依据保留。原A/C主分析的五项情景已计算。当前规则及旧LLM情景须依各自版本的子类型重新计算，不能把旧规则边界数值沿用至v28.2。',
113:'LLM运行。现稿所用为此前gpt-5.6-sol封存结果，共493批、12片，前11片各420条、末片304条；每片新上下文，片内延续，每批最多10条。推理强度high，底层revision未知，temperature未暴露；没有向模型提供A/C标签。保留首次输出与后续字段更正记录。v29仅完成候选提示词和结构检查，尚无新全量模型结果。',
114:'A2 参照与验证。A/C原始判读、抽样权重、待定和各维实际配对数分别保存。H427为历史剩余框全数判读，不是全语料代表性随机样本；1268条用于加权主分析及开发对照，两者不合并估计总体。真人验证将另锁定未参与规则和提示词调整的样本，记录编码员背景、首次独立提交、分歧理由及裁定过程，并以裁定参考检验两条机器方法；既有机器—A/C比较作为参考敏感性保留。真人流程已有待确认模板，其实施结果不由agent判断替代。',
169:'Abstract: Drawing on 55 policy source entries from 2017 to 2026, this study combines meaning-unit coding, sample-based review and expert interviews to examine attention allocation and coordination in innovative drug evaluation. Drug registration, reimbursement and clinical use share concerns about safety, efficacy and clinical value. However, innovation status in registration, benefit–risk assessment for marketing authorization and comparative value for reimbursement address different decision questions and cannot be directly converted into one another. Positive point estimates are relatively high for safety and efficacy in regulatory sources and for economic considerations in reimbursement sources; their interpretation depends on classification boundaries, corpus composition and interval results. Differences in evaluation objects, comparators and evidence requirements help explain the meanings of shared dimensions and the scope of assessment conclusions. Coordination should clarify the conditions under which conclusions can be used across stages, link reusable evidence with supplementary studies and share assessment reasons. The study provides a traceable basis for comparing evaluation content and connecting decisions across stages.',
}

# Complete-version numbers are injected only from audited files; no imputed human results.
extra_path=ROOT/'paragraph_updates.json'
if extra_path.exists():replacements.update({int(k):v for k,v in json.loads(extra_path.read_text()).items()})
for i,text in replacements.items():replace(paras[i],text)
replace(paras[168],'Differences in Attention Allocation in Innovative Drug Evaluation\nand Coordination Pathways',False)
# Keep the established LLM-method source cited, rather than leaving an uncited reference.
replace(paras[43],replacements[43].replace('并以大语言模型（LLM）作平行测量。','并以大语言模型（LLM）作平行测量，参考Gilardi等[38]的文本标注研究。'))
# The detailed codebook stays in the already-delivered supplement rather than repeating five definitions.
for i in range(106,111):body.remove(paras[i])

rows=tables[1].findall('w:tr',NS)
for c,text in zip(rows[0].findall('w:tc',NS),['维度','方法','κ A/C','P A/C','R A/C','F1 A/C']):cellset(c,text)
for j,dim in enumerate(['S','E','Ec','I','F']):
    for off,method,label in [(1,'rules_v28','规则v28.2'),(2,'llm_unchanged','旧LLM')]:
        cells=rows[2*j+off].findall('w:tc',NS)
        texts=[['安全性','有效性','经济性','创新性','公平性'][j],label,pair('H427',dim,method,'kappa'),pair('H427',dim,method,'precision',True),pair('H427',dim,method,'recall',True),pair('H427',dim,method,'f1')]
        for c,t in zip(cells,texts):cellset(c,t)

dist_path=ROOT/'table5_values.json'
if dist_path.exists():
    vals=json.loads(dist_path.read_text())
    for row,texts in zip(tables[4].findall('w:tr',NS),vals):
        for c,t in zip(row.findall('w:tc',NS),texts):cellset(c,t)

# Original identifiers here are remapped to first-citation order below.
newrefs={
46:'[46] COHEN J. A coefficient of agreement for nominal scales[J]. Educational and Psychological Measurement, 1960, 20(1): 37-46. DOI:10.1177/001316446002000104.',
47:'[47] KRISTENSEN-MCLACHLAN R D, CANAVAN M, KÁRDOS M, et al. Are chatbots reliable text annotators? Sometimes[J]. PNAS Nexus, 2025, 4(4): pgaf069. DOI:10.1093/pnasnexus/pgaf069.',
48:'[48] TÖRNBERG P. Best practices for text annotation with large language models[J]. Sociologica, 2024, 18(2): 67-85. DOI:10.6092/issn.1971-8853/19461.',
49:'[49] DAVISON A C, HINKLEY D V. Bootstrap methods and their application[M]. Cambridge: Cambridge University Press, 1997. DOI:10.1017/CBO9780511802843.',
50:'[50] BUONACCORSI J P. A note on confidence intervals for proportions in finite populations[J]. The American Statistician, 1987, 41(3): 215-218. DOI:10.1080/00031305.1987.10475484.',
51:'[51] DUNN O J. Multiple comparisons among means[J]. Journal of the American Statistical Association, 1961, 56(293): 52-64. DOI:10.1080/01621459.1961.10482090.',
}
refupdates=ROOT/'reference_updates.json'
if refupdates.exists():
    for k,v in json.loads(refupdates.read_text()).items():
        n=int(k)
        if n==43:v=v.replace('2026-09-25','2026-10-06')
        if n<=45:replace(paras[122+n],v,False)
        else:newrefs[n]=v
for n,text in newrefs.items():paras[168].addprevious(cloned(paras[167],text,False))

# Compact appendix gives the comparison specifically requested by the author.
anchor=paras[119]
anchor.addprevious(cloned(paras[115],'A4 开发复核的逐维一致性',False))
anchor.addprevious(cloned(paras[48],'表A1 1268条材料中各方法与固定A/C参照的κ',False))
at=deepcopy(tables[1])
for r in list(at.findall('w:tr',NS))[6:]:at.remove(r)
ars=at.findall('w:tr',NS)
for c,t in zip(ars[0].findall('w:tc',NS),['维度','A—C','规则—A','规则—C','旧LLM—A','旧LLM—C']):cellset(c,t)
for j,dim in enumerate(['S','E','Ec','I','F']):
    ac=[r for r in records if r['scope']=='current1268' and r['dimension']==dim and r['method']=='A_C'][0]['metrics']['kappa']['estimate']
    texts=[['安全性','有效性','经济性','创新性','公平性'][j],f'{ac:.3f}']+[f'{val("current1268",dim,m,ref,"kappa"):.3f}' for m,ref in [('rules_v28','A'),('rules_v28','C'),('llm_unchanged','A'),('llm_unchanged','C')]]
    for c,t in zip(ars[j+1].findall('w:tc',NS),texts):cellset(c,t)
anchor.addprevious(at)
anchor.addprevious(cloned(paras[50],'注：A/C为agent首次独立参照；该表属于开发复核，不是真人验证结果。机器—A有效n按S/E/Ec/I/F为1214/1214/1214/1209/1214，机器—C为1216/1216/1216/1213/1215；A—C为1216/1216/1216/1211/1215。原始待定及共同可判掩码保持封存口径，五项指标、混淆矩阵与95%区间见配套统计。',False))

# Keep the compact six-row comparison table intact rather than leaving its final row alone.
for row_no,row in enumerate(tables[3].findall('w:tr',NS)):
    for p in row.xpath('.//w:p',namespaces=NS):
        pp=p.find('w:pPr',NS)
        if pp is None:pp=E.SubElement(p,W('pPr'))
        kn=pp.find('w:keepNext',NS)
        if kn is None:kn=E.SubElement(pp,W('keepNext'))
        kn.set(W('val'),'1' if row_no<5 else '0')

# References in order of first citation, with a crosswalk retained for the authority audit.
refparas={}
for p in body.findall('w:p',NS):
    # Retained Word styles use numeric style IDs; match the bibliography's verified prefix.
    match=re.match(r'^\[(\d+)\]\s',plain(p))
    if match:refparas[int(match.group(1))]=p
order=[]
for p in body.findall('w:p',NS):
    if p in refparas.values():continue
    for n in re.findall(r'\[(\d+)\]',plain(p)):
        n=int(n)
        if n in refparas and n not in order:order.append(n)
order.extend(n for n in sorted(refparas) if n not in order)
crosswalk={n:i+1 for i,n in enumerate(order)}
for p in body.findall('w:p',NS):
    if p in refparas.values():continue
    for t in p.xpath('.//w:t',namespaces=NS):
        if t.text:t.text=re.sub(r'\[(\d+)\]',lambda m:f'[{crosswalk[int(m.group(1))]}]' if int(m.group(1)) in crosswalk else m.group(0),t.text)
for n in order:
    p=refparas[n];txt=re.sub(r'^\[\d+\]',f'[{crosswalk[n]}]',plain(p));replace(p,txt,False)
    body.remove(p);paras[168].addprevious(p)
(ROOT/'reference_number_crosswalk.json').write_text(json.dumps(crosswalk,ensure_ascii=False,indent=2))

# Explicitly permitted title/abstract adjustments to the current journal format.
for p in [paras[0]]:
    for r in p.findall('w:r',NS):
        pr=r.find('w:rPr',NS)
        if pr is None:pr=E.SubElement(r,W('rPr'))
        fonts=pr.find('w:rFonts',NS)
        if fonts is None:fonts=E.SubElement(pr,W('rFonts'))
        fonts.set(W('eastAsia'),'黑体')
        for key in ['sz','szCs']:
            sz=pr.find('w:'+key,NS)
            if sz is None:sz=E.SubElement(pr,W(key))
            sz.set(W('val'),'28')
    pp=p.find('w:pPr',NS);jc=pp.find('w:jc',NS)
    if jc is None:jc=E.SubElement(pp,W('jc'))
    jc.set(W('val'),'center')
for r in paras[1].findall('w:r',NS):
    pr=r.find('w:rPr',NS)
    if pr is None:pr=E.SubElement(r,W('rPr'))
    fonts=pr.find('w:rFonts',NS)
    if fonts is None:fonts=E.SubElement(pr,W('rFonts'))
    fonts.set(W('eastAsia'),'宋体')
    for b in pr.findall('w:b',NS)+pr.findall('w:bCs',NS):pr.remove(b)

# September 2026 journal template places the English front matter before the introduction.
for i in [168,169,170]:
    body.remove(paras[i]);paras[5].addprevious(paras[i])
for r in paras[168].findall('w:r',NS):
    pr=r.find('w:rPr',NS)
    if pr is None:pr=E.SubElement(r,W('rPr'))
    fonts=pr.find('w:rFonts',NS)
    if fonts is None:fonts=E.SubElement(pr,W('rFonts'))
    fonts.set(W('ascii'),'Times New Roman');fonts.set(W('hAnsi'),'Times New Roman')
    sz=pr.find('w:sz',NS)
    if sz is None:sz=E.SubElement(pr,W('sz'))
    sz.set(W('val'),'28')
    b=pr.find('w:b',NS)
    if b is None:b=E.SubElement(pr,W('b'))
    b.set(W('val'),'1')
pp=paras[168].find('w:pPr',NS);jc=pp.find('w:jc',NS)
if jc is None:jc=E.SubElement(pp,W('jc'))
jc.set(W('val'),'center')
for i in [169,170]:
    pp=paras[i].find('w:pPr',NS)
    for key in ['keepNext','keepLines']:
        for q in pp.findall('w:'+key,NS):pp.remove(q)
    ind=pp.find('w:ind',NS)
    if ind is None:ind=E.SubElement(pp,W('ind'))
    ind.set(W('firstLine'),'0')
    ind.set(W('firstLineChars'),'0')
    jc=pp.find('w:jc',NS)
    if jc is None:jc=E.SubElement(pp,W('jc'))
    jc.set(W('val'),'left')
    for r in paras[i].findall('w:r',NS):
        pr=r.find('w:rPr',NS)
        if pr is None:pr=E.SubElement(r,W('rPr'))
        for b in pr.findall('w:b',NS)+pr.findall('w:bCs',NS):pr.remove(b)
        fonts=pr.find('w:rFonts',NS)
        if fonts is None:fonts=E.SubElement(pr,W('rFonts'))
        fonts.set(W('ascii'),'Times New Roman');fonts.set(W('hAnsi'),'Times New Roman')
        sz=pr.find('w:sz',NS)
        if sz is None:sz=E.SubElement(pr,W('sz'))
        sz.set(W('val'),'21')

main=ROOT/'论文完整修订稿-待补真人验证.docx'
xml=E.tostring(doc,encoding='UTF-8',xml_declaration=True,standalone=True)
with ZipFile(BASE) as zin, ZipFile(main,'w') as zout:
    for info in zin.infolist():zout.writestr(info,xml if info.filename=='word/document.xml' else zin.read(info.filename))
preserve={}
with ZipFile(BASE) as a, ZipFile(main) as b:
    preserve={'source_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),'changed_parts':[n for n in a.namelist() if a.read(n)!=b.read(n)],'tables_after':len(body.findall('w:tbl',NS)),'abstract_characters':len(abstract)-3,'reference_crosswalk':crosswalk}
(ROOT/'manuscript_edit_audit.json').write_text(json.dumps(preserve,ensure_ascii=False,indent=2))
if '--main-only' in sys.argv:
    print(json.dumps({'main':str(main),'main_only':True,'changed_parts':preserve['changed_parts']},ensure_ascii=False));sys.exit(0)

# Separate confirmation document keeps actual status and a complete human-validation reporting form.
m=Document();s=m.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    sty=m.styles[name];sty.font.name='Times New Roman';sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
m.styles['Normal'].font.size=Pt(11);m.styles['Normal'].paragraph_format.line_spacing=1.3
m.styles['Title'].font.size=Pt(20);m.styles['Heading 1'].font.size=Pt(14)
# Word's default Title sometimes carries a decorative blue border.
for sty in m.styles:
    for e in sty._element.xpath('.//w:pBdr'):e.getparent().remove(e)
def p(t,style=None):return m.add_paragraph(t,style)
def h(t):return p(t,'Heading 1')
def tab(headers,rows,widths=None):
    t=m.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    for vals in rows:
        for c,v in zip(t.add_row().cells,vals):c.text=str(v)
    for i,r in enumerate(t.rows):
        pr=r._tr.get_or_add_trPr()
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        pr.append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(r.cells):
            if widths:c.width=Cm(widths[j])
            tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                q=OxmlElement('w:'+edge);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');b.append(q)
            tc.append(b)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECF1');tc.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
                for rr in pp.runs:rr.font.size=Pt(10)
    return t
p('创新药评价论文待确认事项','Title')
p('截至2026年10月6日，论文结构、政策分析与当前指标已接入完整稿。当前规则达到开发期κ点估计目标；旧LLM仍需修订重跑，真人独立验证尚未实施。因此这份完整稿可用于继续修改与补齐材料，尚不能声称完成真人效度验证或直接承诺期刊录用。')
h('现有指标与可以支持的结论')
tab(['维度','A—C','规则—A','规则—C','旧LLM—A','旧LLM—C'],[[['安全性','有效性','经济性','创新性','公平性'][j]]+[f"{([r for r in records if r['scope']=='current1268' and r['dimension']==dim and r['method']=='A_C'][0]['metrics']['kappa']['estimate']):.3f}"]+[f'{val("current1268",dim,method,ref,"kappa"):.3f}' for method,ref in [('rules_v28','A'),('rules_v28','C'),('llm_unchanged','A'),('llm_unchanged','C')]] for j,dim in enumerate(['S','E','Ec','I','F'])],[2.2,2.1,2.4,2.4,2.4,2.4])
p('上表为1268条开发材料的固定配对口径。A/C为agent原始判读，不能改称真人数据。两组材料共20项规则κ为0.874—0.974；5项区间下限仍低于0.80。旧LLM20项中11项达到开发目标，创新性与公平性的8项均低于目标。κ说明一致性；政策结论另依据冻结主分析、条文比较及敏感性检验。')
p('论文目前支持：三个评价环节具有安全、有效和临床价值的共同基础，注册创新身份、上市获益风险及支付比较价值服务于不同决策，衔接应明确人群、比较方案、证据和不确定性。主分析尚未形成A/C共同支持的稳健部门排序，不能写成各部门实际赋权已被量化证明。')
h('真人独立验证的方法报告模板')
p('以下段落只有在对应步骤真实完成后，才改为过去时并填入真实人数、样本和结果。两名真人编码员不等于两名agent；现有κ不能迁入真人结果表。')
p('正式真人验证拟邀请两名熟悉药品评价与政策文本分析的编码员A和C，依据CAL-5D-20260926-v1.3手册统一五维操作定义、句段边界与多标签规则。初始18条仅用于校准，不计入最终验证。规则和提示词冻结后，另锁定未参与开发的独立验证样本，保留来源组、政策文件及年度系列信息；两名编码员分别阅读完整句段及必要上下文，不查看机器输出或对方标签，给出0/1与判定依据。材料不足先回查原文，不把缺料直接编码为0。独立提交封存后，逐维报告Cohen’s κ及观察一致率；分歧经原文回查与透明裁定形成共同参考，保留首次判断和裁定理由。自动规则与LLM均与该参考比较，逐维报告TP、FP、FN、TN、精确率、召回率、F1和适当的不确定性，并以对A、C的分别比较作为参考敏感性。')
tab(['需填字段','实际记录'],[['编码员A背景与独立提交日期','待真人执行后确认'],['编码员C背景与独立提交日期','待真人执行后确认'],['独立样本框、抽样方法、样本量依据','待锁定；须核对文件/年度系列的开发暴露'],['最终样本数与各维阳性数','待实际计数；不得预填427或1268'],['协商前五维κ与观察一致率','待原始独立标签计算'],['分歧数量、裁定者与裁定规则','待真实记录'],['两条方法对裁定参考的逐维指标','待规则/提示词冻结后计算']],[6,10])
h('剩余工作与验收条件')
tab(['事项','当前状态','完成条件'],[['LLM修订重跑','候选v29提示词与结构检查已完成；未有新全量模型运行','按指定GPT-5.6-SOL真实执行并封存4924条，重新算指标及固定子样本稳定性'],['真人与新独立验证','报告格式已备；实施未完成','真人首次独立标签、裁定记录及未用于调参的验证结果齐备'],['投稿信息与终核','目标为《科学决策》；现行要求和文献逐条核验报告另存','核对最新要求、作者信息、匿名与基金，补齐真实验证结果后才能正式提交']],[3.1,5.8,7.1])
p('本轮已经完成：当前规则与旧LLM的部门分布、文件与系列等权、年度系列最新成员、剔除联合文本、逐文件剔除，以及五项操作定义边界的敏感性。原A/C主分析及原图保持冻结；机器口径的新计算已接入表5及结果段。当前没有把新增统计误写成所有原图已经重新绘制。')
h('期刊格式与来源核验')
report=ROOT/'期刊核验.md'
if report.exists():
    p('已核验新官网https://www.qkkxjc.cn/及其2026年适用投稿格式。篇幅范围为9000—15000字，中文摘要约300字、关键词3—8个，采用双向匿名。旧官网模板与现行要求有差异，应以新官网为准；具体字体、图表与参考文献差异见本轮期刊核验报告。')
else:p('现行期刊规范正在核验；以本轮期刊核验报告中的已核实项为准，不因历史模板存在而宣称全部格式已对齐。')
p('所有文献的存在、题名作者、卷期页码/文号和支持范围分别核对。真实存在的文献也可能被错误概括，经典理论不能直接证明本文的部门因果关系。完整逐条核验状态、官方链接与未确认项保存在参考文献核验报告；未确认项不算已完成。κ、P/R/F1、bootstrap、有限总体区间及同时覆盖各有对应依据，保守合并和工作阈值属于本研究选择。')
audit=ROOT/'参考文献核验.json'
if audit.exists():
    x=json.loads(audit.read_text());p('原稿45条的存在及核心著录均可确认，其中44条取得发表原文、出版社信息或政策/会议原始记录；原[22]尚未取得全文，正文已改为述评背景，具体句段计数方法转由已读原[23]支撑。原[36]只作词典编码例证，不援引未读的具体算法。6条新增统计方法依据已核对，原编号与新正文编号对照另存。本轮未发现编造的期刊或理论；这不代表每句论述已获全文证实。原AIFA资料的访问日更新为本轮实际核查日。')
p('期刊预警条款：原45条中32篇期刊论文涉及25种刊，均未命中已取得的2025官方名单，历史2023名单亦无命中。中科院说明年度名单不能累积合并。投稿时仍需确认是否已有2026正式更新，以及编辑部适用年份口径；这项检查不等于所有文章撤稿状态的全量检索。')
p('投稿前还须确认两份访谈及研讨会的引用与匿名授权、语料截止日期和后续政策的适用范围，以及作者署名、单位、基金和通讯信息。2026政策文号的本轮核验记录见文献审计，不把未确认项写成已完成。')
memo=ROOT/'待确认事项.docx';m.save(memo)
print(json.dumps({'main':str(main),'memo':str(memo),'changed_parts':preserve['changed_parts'],'abstract_chars':preserve['abstract_characters']},ensure_ascii=False))
