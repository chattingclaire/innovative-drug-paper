"""Read-only provenance and numerical verification of final manuscript files."""
from pathlib import Path
import json
import hashlib
import re
from datetime import datetime, timezone
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
MAIN = OUT / '论文完整修订稿-待补真人验证.docx'
PENDING = OUT / '待确认事项.docx'
METRICS = ROOT / 'outputs/autocode-v28-review-20261004-111306/extended_metrics.json'
SENSITIVITY = OUT / '敏感性复核.json'
DIMS = ['S', 'E', 'Ec', 'I', 'F']
GROUPS = ['NMPA', 'NHSA', 'NHC', 'GOV', 'ALL']

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    before = {str(p): sha(p) for p in (MAIN, PENDING)}
    doc, pending = Document(MAIN), Document(PENDING)
    statistics = json.loads(METRICS.read_text())
    sensitivity = json.loads(SENSITIVITY.read_text())
    metrics = {(r['scope'], r['method'], r['reference'], r['dimension']): r for r in statistics['records']}
    checks = []

    def check(location, actual, expected):
        checks.append({'location': location, 'actual': actual, 'expected': expected, 'passed': actual == expected})

    # Table 2 is doc.tables[1]; each slash pairs A then C.
    table2 = doc.tables[1]
    for row_index, row in enumerate(table2.rows[1:]):
        dimension = DIMS[row_index // 2]
        method = 'rules_v28' if row_index % 2 == 0 else 'llm_unchanged'
        for col_index, metric_name in enumerate(['kappa', 'precision', 'recall', 'f1'], 2):
            formatted = []
            for ref in ('A', 'C'):
                estimate = metrics[('H427', method, ref, dimension)]['metrics'][metric_name]['estimate']
                formatted.append(f'{estimate * 100:.1f}' if metric_name in ('precision', 'recall') else f'{estimate:.3f}')
            check(f'Table 2 row {row_index+1} {dimension} {method} {metric_name}', row.cells[col_index].text, '/'.join(formatted))

    # The full 1,268-unit kappa table appears in both manuscript and pending list.
    column_keys = [('A_C', 'A_C'), ('rules_v28', 'A'), ('rules_v28', 'C'), ('llm_unchanged', 'A'), ('llm_unchanged', 'C')]
    for label, table in [('Table A1', doc.tables[5]), ('Pending document kappa table', pending.tables[0])]:
        for row_index, dimension in enumerate(DIMS, 1):
            for col_index, (method, ref) in enumerate(column_keys, 1):
                expected = f"{metrics[('current1268', method, ref, dimension)]['metrics']['kappa']['estimate']:.3f}"
                check(f'{label} {dimension} {method}/{ref}', table.rows[row_index].cells[col_index].text, expected)

    # Table 5 ratios retain all units in denominator and show pending separately.
    table5 = doc.tables[4]
    for row_index, row in enumerate(table5.rows[1:]):
        group = GROUPS[row_index // 2]
        method = 'rules_v28_2' if row_index % 2 == 0 else 'llm_gpt56_sol_v13_original'
        values = sensitivity['primary_distributions'][method][group]
        check(f'Table 5 {group} {method} N', row.cells[2].text, str(values['S']['N']))
        for col_index, dimension in enumerate(DIMS, 3):
            value = values[dimension]
            expected = f"{value['positive_proportion']*100:.2f}\n({value['pending_n']/value['N']*100:.2f})"
            check(f'Table 5 {group} {method} {dimension}', row.cells[col_index].text, expected)

    # Verify prose summaries rather than accepting only the tables.
    rules = [r for r in statistics['records'] if r['method'] == 'rules_v28']
    llm = [r for r in statistics['records'] if r['method'] == 'llm_unchanged']
    check('Rule comparison count', len(rules), 20)
    check('Rule kappa target pass count', sum(r['metrics']['kappa']['estimate'] >= .8 for r in rules), 20)
    check('Rule kappa CI lower below .80', sum(r['metrics']['kappa']['ci95']['low'] < .8 for r in rules), 5)
    check('Old LLM kappa target pass count', sum(r['metrics']['kappa']['estimate'] >= .8 for r in llm), 11)
    check('Old LLM I/F below target', sum(r['dimension'] in ('I','F') and r['metrics']['kappa']['estimate'] < .8 for r in llm), 8)
    check('Old LLM H427 E-C below target', metrics[('H427','llm_unchanged','C','E')]['metrics']['kappa']['estimate'] < .8, True)

    prose = '\n'.join(p.text for p in doc.paragraphs)
    pending_prose = '\n'.join(p.text for p in pending.paragraphs)
    required_disclosures = {
        'main_AI_reference_not_human': '现有A/C为agent首次独立判读，不能作为已完成的真人金标准。',
        'main_real_human_results_pending': '实施记录和真人结果尚待补齐，不以现有A/C数值代填。',
        'main_Table_A1_not_human_validation': '该表属于开发复核，不是真人验证结果。',
        'main_v29_no_full_model_results': 'v29仅完成候选提示词和结构检查，尚无新全量模型结果。',
        'main_old_LLM_49_pending': '旧LLM仍保留49条五维待定。',
        'pending_AI_table_not_real_human': 'A/C为agent原始判读，不能改称真人数据。',
        'pending_human_future_only': '以下段落只有在对应步骤真实完成后，才改为过去时并填入真实人数、样本和结果。',
    }
    for key, text in required_disclosures.items():
        source = pending_prose if key.startswith('pending_') else prose
        check(key, text in source, True)

    prose_evidence = {
        str(MAIN): [{'paragraph_index': i, 'text': p.text} for i,p in enumerate(doc.paragraphs)
            if any(term in p.text for term in ('真人', 'agent', 'v29', '旧LLM', '旧 LLM'))],
        str(PENDING): [{'paragraph_index': i, 'text': p.text} for i,p in enumerate(pending.paragraphs)
            if any(term in p.text for term in ('真人', 'agent', 'v29', '旧LLM'))],
    }
    unsupported_completed_human_claims = []
    bad_patterns = [r'两名真人[^。]*已完成', r'真人[^。]*验证结果为', r'真人[^。]*κ为', r'v29[^。]*已完成[^。]*全量(?:运行|重跑)']
    for file_path, content in [(MAIN, prose), (PENDING, pending_prose)]:
        for pattern in bad_patterns:
            for match in re.finditer(pattern, content):
                unsupported_completed_human_claims.append({'file': str(file_path), 'text': match.group(0), 'pattern': pattern})
    check('No unsupported completed human/v29 claims', len(unsupported_completed_human_claims), 0)
    after = {str(p): sha(p) for p in (MAIN, PENDING)}
    check('DOCX files unchanged', before, after)
    failures = [c for c in checks if not c['passed']]
    report = {
        'created_at': datetime.now(timezone.utc).isoformat(),
        'read_only': True, 'needs_correction': bool(failures or unsupported_completed_human_claims),
        'conclusion': 'The existing A/C kappa tables remain explicitly AI-reference development results, human validation is pending, and the old LLM has not been claimed as a completed v29 rerun.',
        'checks_n': len(checks), 'all_checks_passed': not failures,
        'numeric_tables_checked': ['Table 2 H427 kappa/P/R/F1', 'Table A1 current1268 kappa', 'Table 5 full-corpus positive and pending proportions', 'Pending document kappa table'],
        'checks': checks, 'failures': failures,
        'unsupported_completed_human_claims': unsupported_completed_human_claims,
        'provenance_prose_evidence': prose_evidence,
        'source_bindings': {str(p): sha(p) for p in (MAIN, PENDING, METRICS, SENSITIVITY, Path(__file__))},
        'scope_limits': ['Read-only text/table provenance and numerical audit; this does not certify external literature or actual human study implementation.',
            'Visual pagination was independently checked on final v4 page images in separate reports.'],
    }
    target = OUT / '最终真实性与数值核对.json'
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'output': str(target), 'checks_n': len(checks), 'all_passed': not failures,
        'needs_correction': report['needs_correction'], 'failures': failures}, ensure_ascii=False))

if __name__ == '__main__':
    main()
