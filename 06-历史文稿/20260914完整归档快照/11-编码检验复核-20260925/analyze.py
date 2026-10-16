from pathlib import Path
import json, csv, hashlib, re
from collections import Counter
import numpy as np
from openpyxl import load_workbook

BASE=Path('/Users/rainbow/Desktop/战略院/创新药论文-完整归档-20260914')
OUT=Path(__file__).resolve().parent
DIMS=['安全性','有效性','经济性','创新性','公平性']
OLD=BASE/'10-人工编码基准/旧口径-作废'
PATHS={'A':OLD/'编码表-编码员A-已填写（旧口径）.xlsx','C':OLD/'编码表-编码员C（机器编码-Claude）.xlsx'}
def read_sheet(path):
    ws=load_workbook(path,read_only=True,data_only=True).active
    out={}; bad=[]
    for row in ws.iter_rows(min_row=2):
        if row[1].value is None:continue
        sid=int(row[1].value); assert sid not in out
        vals=[]
        for c in row[3:8]:
            v=c.value
            if v in (None,'',0,'0'):vals.append(0)
            elif v in (1,'1','1.0'):vals.append(1)
            else:
                vals.append(np.nan);bad.append({'cell':c.coordinate,'id':sid,'value':v,'sentence':row[2].value})
        out[sid]={'values':vals,'sent':row[2].value,'row':row[0].row,'seq':row[0].value}
    return out,bad
S={}; invalid={}
for k,p in PATHS.items(): S[k],invalid[k]=read_sheet(p)
ids=sorted(S['A']); assert ids==sorted(S['C'])
src=BASE/'03-编码数据与脚本'
def jread(name):return json.loads((src/name).read_text())
sent={x['id']:x for x in jread('all_sentences.json')}
oldllm={x['id']:x for x in jread('llm_all_results.json')}
vd={x['id']:x for x in jread('coding_v3_dict.json')}
vl={x['id']:x for x in jread('coding_v3_llm.json')}
assert len(sent)==3387 and len(vd)==3387 and len(vl)==2965
text_mismatch=[]
for sid in ids:
    assert S['A'][sid]['sent']==S['C'][sid]['sent'],sid
    assert sent[sid]['sent']==vd[sid]['sent'],sid
    if S['A'][sid]['sent']!=sent[sid]['sent']:
        text_mismatch.append({'id':sid,'human_chars':len(S['A'][sid]['sent']),'auto_chars':len(sent[sid]['sent']),'prefix_only':sent[sid]['sent'].startswith(S['A'][sid]['sent'])})
def binary(values):return [int(d in values) for d in DIMS]
arr={k:np.array([S[k][i]['values'] for i in ids],float) for k in ['A','C']}
arr['词典v2']=np.array([binary(sent[i]['keyword_dims']) for i in ids],float)
arr['LLMv2']=np.array([binary(oldllm[i]['dims']) for i in ids],float)
arr['词典v3']=np.array([binary(vd[i]['dims_v3']) for i in ids],float)
arr['LLMv3']=np.array([binary(vl[i]['dims']) if i in vl else [np.nan]*5 for i in ids],float)
deps=np.array([sent[i]['dept'] for i in ids]);kept=np.array([not vd[i]['junk'] for i in ids])
textmatched=np.array([S['A'][i]['sent']==sent[i]['sent'] for i in ids])
def metrics(a,b):
    # a is the declared reference, not adjudicated truth.
    keep=np.isfinite(a)&np.isfinite(b); a=a[keep];b=b[keep];n=len(a)
    tp=int(np.sum((a==1)&(b==1)));fp=int(np.sum((a==0)&(b==1)))
    fn=int(np.sum((a==1)&(b==0)));tn=int(np.sum((a==0)&(b==0)))
    ratio=lambda x,y:x/y if y else None
    po=ratio(tp+tn,n);pa=ratio(tp+fn,n);pb=ratio(tp+fp,n)
    pe=pa*pb+(1-pa)*(1-pb) if n else 1
    return dict(n=n,tp=tp,fp=fp,fn=fn,tn=tn,agreement=po,kappa=(po-pe)/(1-pe) if pe<1 else None,
                precision=ratio(tp,tp+fp),recall=ratio(tp,tp+fn),f1=ratio(2*tp,2*tp+fp+fn),
                positive_agreement=ratio(2*tp,2*tp+fp+fn),negative_agreement=ratio(2*tn,2*tn+fp+fn))
def pair(a,b,bootstrap=False):
    result={d:metrics(a[:,di],b[:,di]) for di,d in enumerate(DIMS)}
    result['合并格']=metrics(a.ravel(),b.ravel())
    f=[result[d]['f1'] for d in DIMS];result['macro_f1']=float(np.mean(f)) if None not in f else None
    complete=np.isfinite(a).all(axis=1)&np.isfinite(b).all(axis=1)
    result['whole_sentence']={'n':int(complete.sum()),'agree':int(np.all(a[complete]==b[complete],axis=1).sum()),'rate':float(np.mean(np.all(a[complete]==b[complete],axis=1)))}
    return result
pairs={}
for k in ['C','词典v2','LLMv2']:pairs['A vs '+k]=pair(arr['A'],arr[k])
for k in ['词典v2','LLMv2']:pairs['C vs '+k]=pair(arr['C'],arr[k])
pairs['词典v2 vs LLMv2']=pair(arr['词典v2'],arr['LLMv2'])
diagnostic={}
for k in ['词典v3','LLMv3']:diagnostic['旧A vs '+k]=pair(arr['A'][kept],arr[k][kept]);diagnostic['旧C vs '+k]=pair(arr['C'][kept],arr[k][kept])
def shares(mask):
    out={}
    # Exclude whole sentences with malformed A codes for consistent denominators.
    common=mask&np.isfinite(arr['A']).all(axis=1)
    for dep in sorted(set(deps)):
        m=common&(deps==dep);out[dep]={'sentences':int(m.sum()),'methods':{}}
        for k,a in arr.items():
            if 'v3' in k and not np.all(kept[m]): continue
            c=np.sum(a[m],axis=0); total=int(c.sum())
            out[dep]['methods'][k]={'counts':[int(x) for x in c],'hits':total,'shares':(c/total).tolist() if total else None,'covered':int(np.any(a[m]==1,axis=1).sum())}
    return out
common_ac=np.isfinite(arr['A'])&(arr['A']==arr['C'])
double_old=(arr['词典v2']==arr['LLMv2']); disagreement={}
for ref in ['A','C']:
    valid=np.isfinite(arr[ref]);bothwrong=valid&double_old&(arr[ref]!=arr['词典v2'])
    disagreement[ref]={'valid_cells':int(valid.sum()),'automethod_agree_cells':int((valid&double_old).sum()),
       'both_disagree_reference':int(bothwrong.sum()),'conditional_rate':float(bothwrong.sum()/(valid&double_old).sum())}
ac_counts={'ac_agreement_cells':int(common_ac.sum()),'ac_disagreement_cells':int((np.isfinite(arr['A'])&~common_ac).sum()),
 'ac1_auto_both0':int((common_ac&(arr['A']==1)&(arr['词典v2']==0)&(arr['LLMv2']==0)).sum()),
 'ac0_either_auto1':int((common_ac&(arr['A']==0)&((arr['词典v2']==1)|(arr['LLMv2']==1))).sum()),
 'ac0_auto_both1':int((common_ac&(arr['A']==0)&(arr['词典v2']==1)&(arr['LLMv2']==1)).sum())}
# Paired sentence-level, department-stratified bootstrap preserves the five correlated labels.
rng=np.random.default_rng(20260925);B=5000;boot={}
groups=[np.flatnonzero(deps==d) for d in sorted(set(deps))]
idx=np.concatenate([rng.choice(g,(B,len(g)),replace=True) for g in groups],axis=1)
def vector_stats(a,b):
    valid=np.isfinite(a)&np.isfinite(b);n=valid.sum(axis=1)
    tp=((a==1)&(b==1)&valid).sum(axis=1);tn=((a==0)&(b==0)&valid).sum(axis=1)
    fp=((a==0)&(b==1)&valid).sum(axis=1);fn=((a==1)&(b==0)&valid).sum(axis=1)
    with np.errstate(invalid='ignore',divide='ignore'):
        po=(tp+tn)/n;pa=(tp+fn)/n;pb=(tp+fp)/n;pe=pa*pb+(1-pa)*(1-pb)
        return {'kappa':(po-pe)/(1-pe),'precision':tp/(tp+fp),'recall':tp/(tp+fn),'f1':2*tp/(2*tp+fp+fn)}
for key in pairs:
    x,y=key.split(' vs ');a=arr[x][idx];b=arr[y][idx]
    v=vector_stats(a,b)
    boot[key]={d:{name:np.nanquantile(vals[:,di],[.025,.975]).tolist() for name,vals in v.items()} for di,d in enumerate(DIMS)}
    boot[key]['macro_f1']=np.nanquantile(np.nanmean(v['f1'],axis=1),[.025,.975]).tolist()
    v2=vector_stats(a.reshape(B,-1,1),b.reshape(B,-1,1))
    boot[key]['合并格']={name:np.nanquantile(vals[:,0],[.025,.975]).tolist() for name,vals in v2.items()}
full={}
for ver in ['v2','v3']:
    allids=sorted(sent) if ver=='v2' else sorted(vl)
    da=np.array([binary(sent[i]['keyword_dims'] if ver=='v2' else vd[i]['dims_v3']) for i in allids],float)
    la=np.array([binary((oldllm if ver=='v2' else vl)[i]['dims']) for i in allids],float)
    full[ver]={'agreement':pair(da,la),'by_dept':{},'leave_one_out':{}}
    for dep in ['NMPA','NHSA','NHC','GOV']:
        dis=[i for i in allids if sent[i]['dept']==dep];files=sorted({sent[i]['file'] for i in dis})
        entry={}
        for method in ['词典','LLM']:
            labels={i:binary((sent[i]['keyword_dims'] if ver=='v2' else vd[i]['dims_v3']) if method=='词典' else (oldllm if ver=='v2' else vl)[i]['dims']) for i in dis}
            counts=np.sum(list(labels.values()),axis=0);fc={f:np.sum([labels[i] for i in dis if sent[i]['file']==f],axis=0) for f in files}
            validfc=[c/c.sum() for c in fc.values() if c.sum()]
            pool=counts/counts.sum();equal=np.mean(validfc,axis=0)
            loo=[]
            for f,c in fc.items():
                remain=counts-c
                if remain.sum():loo.append({'file_removed':f,'shares':(remain/remain.sum()).tolist(),'leader':DIMS[int(remain.argmax())],'innovation_gt_safety_efficacy':bool(remain[3]>remain[0] and remain[3]>remain[1])})
            entry[method]={'sentences':len(dis),'counts':counts.tolist(),'shares_pooled':pool.tolist(),'shares_file_equal':equal.tolist(),'files_total':len(files),'files_with_hits':len(validfc),'covered':sum(any(v) for v in labels.values()),'leave_one_out':loo}
        full[ver]['by_dept'][dep]=entry
new_sample=json.loads((BASE/'10-人工编码基准/sample_v3.json').read_text());newids={r['id'] for r in new_sample}
old_sample=json.loads((OLD/'sample_340.json').read_text());assert {r['id'] for r in old_sample}==set(ids)
manifest=[{'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in list(PATHS.values())+[src/n for n in ['all_sentences.json','llm_all_results.json','coding_v3_dict.json','coding_v3_llm.json']]]
result={'dims':DIMS,'invalid':invalid,'manifest':manifest,'sample_n':len(ids),'sample_department_counts':dict(Counter(deps)),
    'text_mismatch':text_mismatch,
    'matched_text_pairs_338':{key:pair(arr[key.split(' vs ')[0]][textmatched],arr[key.split(' vs ')[1]][textmatched]) for key in pairs},
    'old_sample_retained_v3':int(kept.sum()),'old_new_sample_overlap':len(set(ids)&newids),'v3_sample_n':len(newids),
    'pairs_old339':pairs,'pairs_CI95':boot,'bootstrap':{'B':B,'seed':20260925,'unit':'sentence','strata':'department','interval':'percentile'},
    'shares_old336_complete':shares(textmatched),'shares_retained':shares(kept),
    'cross_version_diagnostic_only':diagnostic,'common_disagreements':disagreement,'ac_consensus_diagnostic_only':ac_counts,'full':full,
    'sensitivity_invalid_as_zero':{k:pair(np.nan_to_num(arr['A'],nan=0),arr[k]) for k in ['C','词典v2','LLMv2']}}
result['nhsa_original173_counts']={k:np.nansum(a[deps=='NHSA'],axis=0).astype(int).tolist() for k,a in arr.items() if 'v3' not in k}
# Targeted version-aware sensitivity; these old labels cannot validate changed constructs.
result['new_sample_overlap_ids']=sorted(newids&set(ids))
result['ac_consensus_diagnostic_only']['both_auto_agree_on_ac_agreements']=int((common_ac&double_old).sum())
result['claim_comparison_ci95']={}
g=np.flatnonzero(deps=='NHSA');claim_idx=rng.choice(g,(B,len(g)),replace=True)
for k in ['A','C','词典v2','LLMv2']:
    s=arr[k][claim_idx]
    result['claim_comparison_ci95'][k]={}
    for di,name in [(0,'创新性减安全性'),(1,'创新性减有效性')]:
        dif=(s[:,:,3]-s[:,:,di]).mean(axis=1)
        result['claim_comparison_ci95'][k][name]={'difference_in_sentence_positive_rate':float(np.mean(arr[k][g,3]-arr[k][g,di])),
            'ci95':np.quantile(dif,[.025,.975]).tolist()}
OUT.joinpath('results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False))
with OUT.joinpath('sentence_comparison.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['序号','句子编号','部门','文件','维度','A','C','词典v2','LLMv2','词典v3','LLMv3','v3剔除原因','句子'])
    for ri,sid in enumerate(ids):
        for di,d in enumerate(DIMS):
            w.writerow([S['A'][sid]['seq'],sid,deps[ri],sent[sid]['file'],d]+[int(arr[k][ri,di]) if np.isfinite(arr[k][ri,di]) else '缺失/非法' for k in arr]+[vd[sid]['junk'],sent[sid]['sent']])
print('samples',result['sample_department_counts'],'v3 retained',int(kept.sum()),'new overlap',result['old_new_sample_overlap'])
for key,p in pairs.items():
    print(key,'po',round(p['合并格']['agreement'],4),'k',round(p['合并格']['kappa'],4),'macroF1',round(p['macro_f1'],4),'whole',p['whole_sentence'])
    print('perdim',[(d,round(p[d]['kappa'],4),round(p[d]['f1'],4)) for d in DIMS])
print('ac',ac_counts,'joint',disagreement)
for dep,x in result['shares_old336_complete'].items():
 print(dep,x)
