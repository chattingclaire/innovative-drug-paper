# -*- coding: utf-8 -*-
"""从清洗后的2965句中按部门分层随机抽取296句（10%），生成两份盲编码表。
口径：《目录调整申报指南》五个维度（申报指南口径 v3）。盲编码表不含词典与LLM判定，避免锚定。"""
import json, random, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20260925
DIMS = ['安全性', '有效性', '经济性', '创新性', '公平性']
BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '03-编码数据与脚本')
data = [x for x in json.load(open(os.path.join(BASE, 'coding_v3_dict.json'), encoding='utf-8')) if not x['junk']]
by = {}
for s in data: by.setdefault(s['dept'], []).append(s)
N = round(len(data) * 0.1)
rng = random.Random(SEED)
sample = []
for dept, lst in sorted(by.items()):
    k = round(N * len(lst) / len(data))
    sample += rng.sample(lst, k)
rng.shuffle(sample)
json.dump([{'id': x['id'], 'dept': x['dept'], 'file': x['file'], 'sent': x['sent']} for x in sample],
          open('sample_v3.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

def sheet(path, coder):
    wb = Workbook(); ws = wb.active; ws.title = '编码表'
    hdr = ['序号', '句子编号', '句子内容'] + DIMS + ['备注']
    ws.append(hdr)
    for c in range(1, len(hdr)+1):
        cell = ws.cell(1, c); cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDDDDD')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for i, x in enumerate(sample, 1):
        ws.append([i, x['id'], x['sent'][:900]] + ['']*len(DIMS) + [''])
    ws.column_dimensions['A'].width = 6; ws.column_dimensions['B'].width = 9
    ws.column_dimensions['C'].width = 96
    for col in 'DEFGH': ws.column_dimensions[col].width = 8
    ws.column_dimensions['I'].width = 20
    dv = DataValidation(type='list', formula1='"1,0"', allow_blank=True,
                        prompt='涉及该维度填 1，不涉及填 0 或留空', showInputMessage=True)
    ws.add_data_validation(dv); dv.add('D2:H%d' % (len(sample)+1))
    for r in range(2, len(sample)+2):
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical='top')
        for c in range(4, 9): ws.cell(r, c).alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'D2'
    wb.save(path)
    print('已生成 %s（编码员%s，%d句）' % (path, coder, len(sample)))

sheet('编码表-编码员A.xlsx', 'A')
sheet('编码表-编码员B.xlsx', 'B')
print('抽样种子 %d，分层：%s' % (SEED, {d: sum(1 for x in sample if x['dept'] == d) for d in sorted(by)}))
