"""显式样本对齐；空白不当0；无裁决表不生成金标准。"""
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook
DIMS=['安全性','有效性','经济性','创新性','公平性']

def binary(v,legacy=False):
    if v is None or isinstance(v,str) and not v.strip():
        if legacy:return 0
        raise ValueError('标签空白')
    if isinstance(v,bool):raise ValueError('布尔值不是明确的0/1')
    if isinstance(v,(int,float)) and v in (0,1):return int(v)
    if isinstance(v,str) and v.strip() in ('0','1','0.0','1.0'):return int(float(v))
    raise ValueError(f'非法标签 {v!r}')

def index(path):
    rows=json.loads(Path(path).read_text());out={}
    if not isinstance(rows,list) or not rows:raise ValueError('记录列表为空或格式错误')
    for r in rows:
        if r['id'] in out:raise ValueError(f'重复编号 {r["id"]}')
        out[r['id']]=r
    return out

def read(path,sample,legacy=False):
    book=load_workbook(path,read_only=True,data_only=False);out={};errors=[]
    try:
        rows=iter(book.active.iter_rows(values_only=True));head=next(rows)
        for j,d in enumerate(DIMS,3):
            if len(head)<=j or d not in str(head[j]):raise ValueError(f'{path}: 第{j+1}列不是{d}')
        for n,row in enumerate(rows,2):
            if all(v is None for v in row):continue
            if len(row)<8:errors.append(f'行{n}缺列');continue
            k=row[1]
            if isinstance(k,str) and k.strip().isdigit():k=int(k.strip())
            if k not in sample:errors.append(f'行{n}未知编号{k!r}');continue
            if k in out:errors.append(f'行{n}重复编号{k}');continue
            if row[2]!=sample[k]['sent']:errors.append(f'编号{k}显示文本不一致')
            vals=[]
            for j,d in enumerate(DIMS,3):
                try:vals.append(binary(row[j],legacy))
                except ValueError as e:errors.append(f'行{n}/{d}: {e}');vals.append(None)
            out[k]=vals
        if set(sample)!=set(out):errors.append(f'缺少{len(set(sample)-set(out))}行')
        if errors:raise ValueError(f'{path}: {len(errors)}项错误；'+'；'.join(errors[:8]))
        return out
    finally:book.close()

def check(a,b):
    if not a or len(a)!=len(b) or any(v not in (0,1) for v in a+b):raise ValueError('标签须非空、等长且二元')

def kappa(a,b):
    check(a,b);n=len(a);po=sum(x==y for x,y in zip(a,b))/n;pa=sum(a)/n;pb=sum(b)/n
    pe=pa*pb+(1-pa)*(1-pb)
    return dict(n=n,kappa=(po-pe)/(1-pe) if pe<1 else None,agreement=po,
                undefined_reason='边际分布退化，期望一致率为1' if pe==1 else None)

def prf(a,b):
    check(a,b);tp=sum(x==1 and y==1 for x,y in zip(a,b));fp=sum(x==1 and y==0 for x,y in zip(a,b));fn=sum(x==0 and y==1 for x,y in zip(a,b));tn=len(a)-tp-fp-fn
    return dict(n=len(a),tp=tp,fp=fp,fn=fn,tn=tn,precision=tp/(tp+fp) if tp+fp else None,
        recall=tp/(tp+fn) if tp+fn else None,f1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None,accuracy=(tp+tn)/len(a))

def predictions(path,sample,key,text_check):
    rows=index(path);out={}
    for i,s in sample.items():
        if i not in rows:raise ValueError(f'{path}: 缺少预测{i}')
        r=rows[i]
        if text_check and r.get('sent')!=s['sent']:raise ValueError(f'{path}: 预测输入文本{i}不一致')
        if r.get('junk'):raise ValueError(f'{path}: 样本{i}在此版本被剔除')
        if key not in r:raise ValueError(f'{path}: 缺少维度字段')
        ds=r[key]
        if not isinstance(ds,list) or len(ds)!=len(set(ds)) or set(ds)-set(DIMS):raise ValueError('维度列表非法')
        out[i]=[int(d in ds) for d in DIMS]
    return out

def main():
    p=argparse.ArgumentParser(description=__doc__)
    for name in ['a','b']:p.add_argument(name)
    for name in ['sample','codebook','output']:p.add_argument('--'+name,required=True)
    for name in ['gold','dictionary','llm']:p.add_argument('--'+name)
    p.add_argument('--legacy-blank-is-zero',action='store_true');args=p.parse_args()
    supplied=[bool(args.gold),bool(args.dictionary),bool(args.llm)]
    if any(supplied) and not all(supplied):p.error('裁决表gold、dictionary、llm须同时提供')
    sample=index(args.sample);ids=list(sample);a=read(args.a,sample,args.legacy_blank_is_zero);b=read(args.b,sample,args.legacy_blank_is_zero)
    flat=lambda x:[v for i in ids for v in x[i]]
    result=dict(sample_n=len(ids),blank_policy='旧表约定留空为0' if args.legacy_blank_is_zero else '空白报错',
        reliability={d:kappa([a[i][j] for i in ids],[b[i][j] for i in ids]) for j,d in enumerate(DIMS)},
        pooled_reliability=kappa(flat(a),flat(b)),
        disagreements=[dict(id=i,dimension=d,a=a[i][j],b=b[i][j]) for i in ids for j,d in enumerate(DIMS) if a[i][j]!=b[i][j]],validity=None,
        provenance_caution='哈希只记录文件版本；编码员实际使用何种手册、是否独立须有真实操作记录支持。')
    if args.gold:
        g=read(args.gold,sample);dic=predictions(args.dictionary,sample,'dims_v3',True);llm=predictions(args.llm,sample,'dims',False)
        result['validity']={name:{d:prf([pred[i][j] for i in ids],[g[i][j] for i in ids]) for j,d in enumerate(DIMS)} for name,pred in [('dictionary',dic),('llm',llm)]}
        agree=sum(dic[i][j]==llm[i][j] for i in ids for j in range(5));wrong=sum(dic[i][j]==llm[i][j] and dic[i][j]!=g[i][j] for i in ids for j in range(5))
        result['common_errors']=dict(automatic_agreement_cells=agree,both_wrong_cells=wrong,conditional_error_rate=wrong/agree if agree else None,all_cell_error_rate=wrong/(5*len(ids)))
        result['llm_input_caution']='当前LLM结果格式不含原文；须另核查完整输入、截断和会话独立性。'
    result['input_files']={k:dict(path=str(Path(v).resolve()),sha256=hashlib.sha256(Path(v).read_bytes()).hexdigest()) for k,v in vars(args).items() if k in ('a','b','sample','codebook','gold','dictionary','llm') and v}
    dest=Path(args.output)
    if dest.exists():raise ValueError('输出文件已存在，请指定新路径保留历史结果')
    dest.parent.mkdir(parents=True,exist_ok=True);dest.write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False))
    print(f'样本{len(ids)}句，分歧{len(result["disagreements"])}格；'+('已计算裁决后效度' if args.gold else '仅报告信度，不生成金标准'))

if __name__=='__main__':main()
